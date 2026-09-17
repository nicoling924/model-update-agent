"""ANATOMY PROBE (owner 2026-09-17): what does the structure turn actually ask,
what does the brain actually answer, and what does the referee keep?

No pipeline change — this only drives pipeline/anatomy.py the way a cold run
drives it: the same prompt builder, the same ONE live call, the same JSON
validator, the same referee. It prints the prompt, the raw answer, what was
TAKEN, what was DROPPED and why, and the balance pair it would build.

Usage: python tools/anatomy_probe.py [<company_dir>] [<target_year>] [<period>]
Default: companies/CX 2025 FY
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl  # noqa: E402

from pipeline import anatomy, llm  # noqa: E402
from pipeline.checks import year_columns  # noqa: E402
from pipeline.discover import discover  # noqa: E402
from pipeline.evaluator import Evaluator  # noqa: E402
from pipeline.spec import extend_axis  # noqa: E402

PROMPT_HEAD = 3000


def model_path(company_dir):
    mdir = Path(company_dir) / "model"
    cands = sorted([p for p in mdir.glob("*.xls[xm]")
                    if not p.name.startswith("~$") and "(pipeline" not in p.name],
                   key=lambda p: p.stat().st_mtime, reverse=True)
    if not cands:
        raise FileNotFoundError(f"no model workbook in {mdir}")
    return cands[0]


def label_of(ws, row):
    for c in range(1, 6):
        v = ws.cell(row, c).value
        if isinstance(v, str) and v.strip() and not v.startswith("="):
            return v.strip()[:48]
    return ""


def cell_of(spec, sheet, row, target_year):
    col = year_columns(spec, sheet).get(str(target_year))
    return f"{sheet}!{col}{row}" if col else f"{sheet}!row {row} (no {target_year} column)"


def value_at(ev, wbv, spec, sheet, row, target_year):
    col = year_columns(spec, sheet).get(str(target_year))
    if not col:
        return "— no target column"
    try:
        v = ev.cell(sheet, f"{col}{row}")
    except Exception as e:                       # noqa: BLE001 — said, not swallowed
        v = None
        note = f" (live eval: {type(e).__name__})"
    else:
        note = ""
    if not isinstance(v, (int, float)):
        try:
            v = wbv[sheet][f"{col}{row}"].value
        except Exception:                        # noqa: BLE001
            v = None
        note += " [cached]"
    if isinstance(v, (int, float)):
        return f"{v:,.2f}{note}"
    return f"— not evaluable offline{note}"


class Capturing:
    """The run's own client, with the raw answer kept so the probe can print it."""

    def __init__(self, inner):
        self.inner = inner
        self.raws = []

    def __getattr__(self, k):
        return getattr(self.inner, k)

    def chat(self, *a, **kw):
        raw = self.inner.chat(*a, **kw)
        self.raws.append(raw)
        return raw

    def json(self, system, user, validate, repair_retries=2, images=None):
        return type(self.inner).json(self, system, user, validate,
                                     repair_retries=repair_retries, images=images)


class Canned:
    """A client that answers with the object already fetched — so the referee runs
    on exactly the answer printed above, with no second call."""

    def __init__(self, obj):
        self.obj = obj

    def json(self, *a, **kw):
        return self.obj


def main(company_dir, target_year, period):
    print(f"=== CONTEXT: {company_dir}, target {period} {target_year} (COLD RUN — no spec read)")
    path = model_path(company_dir)
    print(f"model: {path.name}")
    wb = openpyxl.load_workbook(path)
    wbv = openpyxl.load_workbook(path, data_only=True)
    print(f"sheets: {wb.sheetnames}")
    print(f"_SPEC tab present: {'_SPEC' in wb.sheetnames}  "
          f"spec.yaml present: {(Path(company_dir) / 'spec.yaml').exists()}")

    # the cold path of pipeline/run.py: no spec -> deterministic discovery
    kind = ("1H" if str(period).upper().startswith(("1H", "2H", "H1", "H2"))
            else "Q" if "Q" in str(period).upper() else "FY")
    spec = discover(openpyxl.load_workbook(path), openpyxl.load_workbook(path, data_only=True),
                    target_year=int(target_year), period_kind=kind)
    extend_axis(spec, int(target_year))
    print(f"discovery: {len(spec.get('year_axis') or {})} sheet(s) with a year axis, "
          f"{len(spec.get('check_rows') or [])} check row(s), "
          f"{len(spec.get('key_rows') or [])} key row(s)")
    for sh, ax in (spec.get("year_axis") or {}).items():
        cols = ax.get("columns") or {}
        print(f"  axis {sh}: {len(cols)} year(s); {target_year} -> {cols.get(str(target_year)) or 'NONE'}")
    for c in (spec.get("check_rows") or []):
        print(f"  discovered check {c['sheet']}!{c['row']}  '{label_of(wb[c['sheet']], int(c['row']))}'")
    for k in (spec.get("key_rows") or []):
        print(f"  discovered key '{k.get('name')}' {k['sheet']}!{k['row']}  "
              f"'{label_of(wb[k['sheet']], int(k['row']))}'")
    print(f"anatomy.wanted(spec) = {anatomy.wanted(spec)}")

    # 1. the exact prompt the run would build
    user = (f"This model is being updated to {period} {target_year}.\n\n"
            + anatomy.sheet_reading(wb, spec, int(target_year)))
    sent = user[:60000]
    print("\n=== PROMPT — SYSTEM ({} chars)".format(len(anatomy.SYSTEM)))
    print(anatomy.SYSTEM)
    print(f"\n=== PROMPT — USER: {len(user):,} chars built, {len(sent):,} sent "
          f"(the run truncates at 60,000); first {PROMPT_HEAD} chars:")
    print(sent[:PROMPT_HEAD])
    print(f"... [{len(sent) - PROMPT_HEAD:,} more chars]" if len(sent) > PROMPT_HEAD else "")
    print("=== PROMPT — USER tail (last 600 chars):")
    print(sent[-600:])

    # 2. ONE live call, the run's own client and validator
    if not llm.env_ready():
        print("\n!!! no LLM env (LLM_BASE_URL / LLM_API_KEY / LLM_MODEL) — no call made")
        return 2
    client = Capturing(llm.make_client())
    print(f"\n=== BRAIN: model={client.model} base={client.base_url}")
    print("handshake: " + str(llm.handshake(client.inner)))
    try:
        obj = client.json(anatomy.SYSTEM, sent, anatomy._validate, repair_retries=1)
    except Exception as e:                       # noqa: BLE001
        print(f"!!! the call failed: {e!r}")
        for i, raw in enumerate(client.raws):
            print(f"--- raw attempt {i + 1}:\n{raw}")
        return 3

    # 3. the raw answer
    print(f"\n=== RAW ANSWER ({len(client.raws)} attempt(s))")
    for i, raw in enumerate(client.raws):
        print(f"--- attempt {i + 1} ({len(raw):,} chars):")
        print(raw)

    # 4. the referee, exactly as the run runs it
    print("\n=== REFEREE (pipeline/anatomy.read, same spec, same workbooks)")
    logs = []

    def log(m):
        logs.append(str(m))
        print("  " + str(m))

    before_c = len(spec.get("check_rows") or [])
    before_k = len(spec.get("key_rows") or [])
    took_c, took_k, dropped = anatomy.read(wb, spec, Canned(obj), int(target_year), log,
                                           period=str(period), values_wb=wbv)
    ev = Evaluator(wb)

    print(f"\n=== TAKEN: {took_c} check row(s), {took_k} key row(s) "
          f"(the brain named {len(obj.get('checks') or [])} check(s) and {len(obj.get('keys') or [])} key(s))")
    print("-- checks now on the spec:")
    for i, c in enumerate(spec.get("check_rows") or []):
        src = "discovery" if i < before_c else "the brain"
        print(f"   {cell_of(spec, c['sheet'], c['row'], target_year):28} "
              f"'{label_of(wb[c['sheet']], int(c['row'])):48}' "
              f"= {value_at(ev, wbv, spec, c['sheet'], c['row'], target_year):24} [{src}]")
    print("-- keys now on the spec:")
    for i, k in enumerate(spec.get("key_rows") or []):
        src = "discovery" if i < before_k else (k.get("source") or "the brain")
        print(f"   {str(k.get('name'))[:26]:26} {cell_of(spec, k['sheet'], k['row'], target_year):22} "
              f"'{label_of(wb[k['sheet']], int(k['row'])):48}' "
              f"= {value_at(ev, wbv, spec, k['sheet'], k['row'], target_year):24} [{src}]")
    print(f"-- input sites taken: {len(spec.get('input_sites') or [])}")
    for s in (spec.get("input_sites") or [])[:20]:
        print(f"   {s['sheet']}!{s['row']}  '{label_of(wb[s['sheet']], int(s['row']))}'")

    print(f"\n=== DROPPED: {len(dropped)}")
    for d in dropped:
        print(f"   {d}")

    print("\n=== BALANCE")
    pairs = spec.get("check_pairs") or []
    if pairs:
        for p in pairs:
            print(f"   pair built: {p['sheet']}!{p['a_row']} '{label_of(wb[p['sheet']], int(p['a_row']))}' "
                  f"LESS {p['b_sheet']}!{p['b_row']} '{label_of(wb[p['b_sheet']], int(p['b_row']))}'")
            for year, col in sorted((year_columns(spec, p['sheet']) or {}).items()):
                bcol = (year_columns(spec, p['b_sheet']) or {}).get(year) or col
                try:
                    d = ev.cell(p['sheet'], f"{col}{p['a_row']}") - ev.cell(p['b_sheet'], f"{bcol}{p['b_row']}")
                    print(f"     {year}: {d:,.2f}")
                except Exception as e:           # noqa: BLE001
                    print(f"     {year}: not evaluable here ({type(e).__name__})")
    elif any("BALANCE IS NOT MEASURED" in m or "not measured" in m for m in logs):
        print("   BALANCE NOT MEASURED")
    else:
        print("   no pair built — a check row of the model is held to measure the balance sheet "
              "(see the checks taken above)")
    print("\n=== ANATOMY NOTES")
    for n in (spec.get("anatomy_notes") or []):
        print(f"   {n}")
    print(f"\n=== because: {str(obj.get('because') or '')}")
    return 0


if __name__ == "__main__":
    a = sys.argv[1:]
    sys.exit(main(a[0] if a else "companies/CX",
                  int(a[1]) if len(a) > 1 else 2025,
                  a[2] if len(a) > 2 else "FY"))
