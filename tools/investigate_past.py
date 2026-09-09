"""PRACTICE ON PAST RUNS (owner 2026-09-10): replay the investigator, dry,
over past delivered CLP workbooks against the analyst's pre-update model —
which headline lines it would flag, where each swing traces to, and what
it would call the swing factor. No cell is written.
    python3 tools/investigate_past.py [--fix] <company_dir> <PERIOD> <YEAR> <artifact_dir>...
--fix: run the corrections on a copy in memory and show the lines before/after (nothing on disk changes).
"""
import glob, json, re, sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from openpyxl import load_workbook as load
from pipeline import spec as spec_mod
from pipeline.sensecheck import headline_deltas, suspicious
from pipeline.investigate import trace, unusual_by_history


def _loop_for(cd, wb, spec, art, period, year):
    """A loop over the artifact's own ledger and provenance, with a writer
    whose 'written' and 'flags' mirror the delivered file — the state the
    live final pass would see."""
    from pipeline.ledger import Ledger
    from pipeline.orchestrator import ObjectiveLoop
    from pipeline.writer import Writer
    from pipeline.targets import TargetRow
    from pipeline.checks import prior_column, year_columns
    led = Ledger.from_json(Path(glob.glob(art + f"/**/replay/{period}/ledger.json", recursive=True)[0]).read_text())
    prov = json.load(open(glob.glob(art + f"/**/replay/{period}/provenance.json", recursive=True)[0]))
    served = {(k.rsplit("!", 1)[0], int(k.rsplit("!", 1)[1])): dict(v, homed=True) for k, v in prov.items()}
    targets = []
    for sh in spec["year_axis"]:
        if sh not in wb.sheetnames:
            continue
        pc = prior_column(spec, sh, year); ws = wb[sh]
        for r in range(1, ws.max_row + 1):
            lab = ws.cell(r, 1).value; pv = ws[f"{pc}{r}"].value if pc else None
            if isinstance(lab, str) and lab.strip():
                targets.append(TargetRow(sh, r, lab.strip(), float(pv) if isinstance(pv, (int, float)) else None))
    w = Writer(wb)
    for k in prov:
        sh, r = k.rsplit("!", 1); tc = year_columns(spec, sh).get(str(year))
        if not tc or sh not in wb.sheetnames:
            continue
        w.log["written"].append(f"{sh}!{tc}{r}")
        try:
            if str(wb[sh][f"{tc}{r}"].fill.fgColor.rgb)[-6:] == "FFC7CE":
                w.log["flags"].append(f"{sh}!{tc}{r}")
        except Exception:
            pass
    led.classify_doc_periods([t.prior_value for t in targets if isinstance(t.prior_value, (int, float))], [])
    return ObjectiveLoop(wb, spec, year, led, targets, served, w, None)


def main(argv):
    fix = "--fix" in argv
    argv = [a for a in argv if a != "--fix"]
    company, period, year = argv[0], argv[1], int(argv[2])
    cd = Path(company)
    pre = load(sorted(cd.glob("model-archive/*_pre.xlsx"))[-1])
    for art in argv[3:]:
        models = glob.glob(art + "/**/model/*(pipeline).xlsx", recursive=True)
        krs = glob.glob(art + f"/**/replay/{period}/key_rows.json", recursive=True)
        provs = glob.glob(art + f"/**/replay/{period}/provenance.json", recursive=True)
        if not models or not krs:
            continue
        wb = load(models[0]); spec = spec_mod.load(cd, wb); spec_mod.extend_axis(spec, year)
        kr = json.load(open(krs[0])); spec["key_rows"] = kr if isinstance(kr, list) else kr.get("key_rows", kr)
        prov = json.load(open(provs[0])) if provs else {}
        print(f"=== {Path(art).name} ({Path(models[0]).name})")
        if fix:
            from pipeline.sensecheck import investigate_line
            loop = _loop_for(cd, wb, spec, art, period, year)
            before = suspicious(headline_deltas(wb, pre, spec, year))
            for d in before:
                v, t, _leaf = investigate_line(loop, pre, d, lambda s: None, rerun=None)
                print(f"  {d['name']:22} {d['d0']*100:+6.1f}% / {d['d1']*100:+6.1f}%  {v.upper():7} {t[:230]}")
            after = suspicious(headline_deltas(wb, pre, spec, year))
            print(f"  -> suspicious lines {len(before)} → {len(after)}: {[(x['name'], round(abs(x['d1']-x['d0'])*100)) for x in after]}")
            continue
        for d in suspicious(headline_deltas(wb, pre, spec, year)):
            sh1, c1 = d["ref1"].split("!")
            trail, leaf = trace(wb, pre, sh1, c1)
            path = " → ".join(f"{t[2] or t[1]}({t[3]*100:.0f}%)" for t in trail)
            if leaf is None:
                print(f"  {d['name']:22} {d['d0']*100:+6.1f}% / {d['d1']*100:+6.1f}%  spread: {path}")
                continue
            sh, coord = leaf; r = int(re.sub(r"[A-Z]", "", coord))
            rgb = str(wb[sh][coord].fill.fgColor.rgb or "")[-6:]
            colour = "RED" if rgb == "FFC7CE" else "ORANGE" if rgb == "FFC000" else "plain"
            e = prov.get(f"{sh}!{r}") or {}
            un, mv, band = unusual_by_history(wb, spec, sh, r, year)
            print(f"  {d['name']:22} {d['d0']*100:+6.1f}% / {d['d1']*100:+6.1f}%  {path} → {sh}!{coord} [{colour}] "
                  f"{str(wb[sh].cell(r,1).value)[:24]!r} = {wb[sh][coord].value!r:.40} "
                  f"{'UNUSUAL' if un else 'within history'} src={str(e.get('note') or e.get('line'))[:50]!r}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
