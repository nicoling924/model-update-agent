#!/usr/bin/env python3
"""Score a delivered CX model against the FY2025 answer key.

    python tools/score_cx.py <delivered.xlsx> [--key KEY.json] [--answer ANSWER.xlsx]

The key names every cell of the FY2025 column the analyst actually filled
(Valuation G/AZ, CXMODEL AO/FK, Fleet AD). A delivered cell is CORRECT when it
lands within 0.5 absolute OR 0.01% of the answer, WRONG when it holds something
else, UNFILLED when it is empty or is a formula the delivered file carries no
cached value for (openpyxl writes no cached values: a model saved by the agent
without a recalc scores as unfilled, not as wrong - recalc it in Excel or
LibreOffice before scoring if you want its formulas counted).

Nothing here is CX-specific beyond the default paths; it reads the key.
"""
import argparse
import json
from pathlib import Path

import openpyxl

TESTSET = Path("/Users/lingling/Project M/test-set/Cathay Pacific")
STEM = "0001007416_67_20260805073722"
RED, ORANGE = "FFC7CE", "FFC000"


def isnum(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def close(a, b, abs_tol=0.5, rel_tol=1e-4):
    if not (isnum(a) and isnum(b)):
        return str(a).strip() == str(b).strip()
    return abs(a - b) <= abs_tol or abs(a - b) <= abs(b) * rel_tol


def fill_of(cell):
    """red / orange / plain - the flag the house conventions put on the cell."""
    try:
        rgb = cell.fill.start_color.rgb
    except Exception:
        return "plain"
    if not isinstance(rgb, str):
        return "plain"
    rgb = rgb[-6:].upper()
    return {RED: "red", ORANGE: "orange"}.get(rgb, "plain")


def answer_values(key, answer_path):
    """value per cell: the key's own, the ANSWER workbook's cache where the key
    has none (13 cells carried no cached value when the key was built)."""
    vals = {}
    wb = openpyxl.load_workbook(answer_path, data_only=True) if answer_path else None
    for sh, d in key["sheets"].items():
        for ref, rec in d["cells"].items():
            v = rec["value"]
            if v is None and wb is not None and sh in wb.sheetnames:
                v = wb[sh][ref].value
            vals[(sh, ref)] = (v, rec["label"])
    return vals


def score(delivered, key_path, answer_path, show=60):
    key = json.load(open(key_path))
    want = answer_values(key, answer_path)
    wbf = openpyxl.load_workbook(delivered)              # formulas + formats
    wbv = openpyxl.load_workbook(delivered, data_only=True)   # cached values

    rows, per_sheet = [], {}
    for (sh, ref), (ans, label) in want.items():
        if sh not in wbf.sheetnames:
            got, kind, flag = None, "unfilled", "plain"
        else:
            cf, cv = wbf[sh][ref], wbv[sh][ref]
            flag = fill_of(cf)
            raw, cached = cf.value, cv.value
            got = cached
            if got is None and not (isinstance(raw, str) and raw.startswith("=")):
                got = raw                      # a hardcode needs no recalc
            if got is None or (isinstance(got, str) and not got.strip()):
                kind = "unfilled"
            elif ans is None:
                kind = "unscorable"            # no answer value to compare to
            else:
                kind = "correct" if close(got, ans) else "wrong"
        rows.append((sh, ref, label, got, ans, kind, flag))
        s = per_sheet.setdefault(sh, {"correct": 0, "wrong": 0, "unfilled": 0,
                                      "unscorable": 0})
        s[kind] += 1

    tot = {"correct": 0, "wrong": 0, "unfilled": 0, "unscorable": 0}
    for s in per_sheet.values():
        for k in tot:
            tot[k] += s[k]
    n = sum(tot.values())

    print(f"scored {Path(delivered).name}  against  {Path(key_path).name}"
          f"  ({key['target_year']} column)")
    print(f"\n{n} key cells: {tot['correct']} correct, {tot['wrong']} wrong, "
          f"{tot['unfilled']} unfilled, {tot['unscorable']} unscorable"
          f"   -> {100.0 * tot['correct'] / n:.1f}% correct")
    print("\nper sheet:")
    for sh, s in per_sheet.items():
        m = sum(s.values())
        print(f"  {sh:<10} {m:>4} cells   correct {s['correct']:>4}  "
              f"wrong {s['wrong']:>4}  unfilled {s['unfilled']:>4}  "
              f"unscorable {s['unscorable']:>3}")

    # the 0.5 absolute band makes every small-magnitude cell (ratios, EPS, DPS)
    # a free pass - say how many of the correct ones only cleared that band.
    strict = sum(1 for sh, ref, lab, got, ans, kind, flag in rows
                 if kind == "correct" and isnum(ans) and isnum(got)
                 and not close(got, ans, abs_tol=0.0))
    print(f"\n{strict} of the correct cells cleared only the 0.5 absolute band "
          f"(small-magnitude ratios), not the 0.01% band.")

    wrong = [r for r in rows if r[5] == "wrong"]
    if wrong:
        print(f"\nwrong cells ({len(wrong)}"
              f"{', first %d shown' % show if len(wrong) > show else ''}):")
        print(f"  {'cell':<16}{'fill':<8}{'delivered':>20}{'answer':>20}   label")
        for sh, ref, label, got, ans, kind, flag in wrong[:show]:
            print(f"  {sh + '!' + ref:<16}{flag:<8}{_fmt(got):>20}{_fmt(ans):>20}"
                  f"   {label[:44]}")
    return tot


def _fmt(v):
    if isnum(v):
        return f"{v:,.4f}".rstrip("0").rstrip(".")
    return str(v)[:20]


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("delivered")
    ap.add_argument("--key", default=str(TESTSET / f"{STEM}__KEY.json"))
    ap.add_argument("--answer", default=str(TESTSET / f"{STEM}__ANSWER.xlsx"))
    ap.add_argument("--show", type=int, default=60)
    a = ap.parse_args()
    score(a.delivered, a.key, a.answer if Path(a.answer).exists() else None, a.show)


if __name__ == "__main__":
    main()
