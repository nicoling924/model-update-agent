"""Score a delivered workbook's INPUT cells (hardcodes + embedded literals) in
the actual column against a reference workbook, the owner's three counts."""
import sys, re, json
sys.path.insert(0, ".")
from openpyxl import load_workbook
from pipeline.evaluator import Evaluator
from pipeline.spec import read_spec_tab
from pipeline.checks import year_columns
from openpyxl.utils import column_index_from_string as ci
deliv, ref, year = sys.argv[1], sys.argv[2], int(sys.argv[3])
wb = load_workbook(deliv); rb = load_workbook(ref)
spec = read_spec_tab(wb); ev = Evaluator(wb); rev = Evaluator(rb)
LIT = re.compile(r"(?<![A-Z$!:'])(?<![A-Za-z0-9_.])[-+]?\d+(?:\.\d+)?(?![0-9.]*[A-Za-z!$])")
def fill(c):
    try: return (c.fill.fgColor.rgb or "")[-6:] if c.fill.fill_type == "solid" else ""
    except Exception: return ""
from pipeline.composites import literals_of
def is_input(v):
    """A typed number, or a formula carrying an embedded hardcode as the agent's own constants law defines it."""
    if isinstance(v, bool): return False
    if isinstance(v, (int, float)): return True
    if isinstance(v, str) and v.startswith("="):
        try: return bool(literals_of(v))
        except Exception: return False
    return False
tot = ok = 0; wrong = []; colours = {"FFC7CE": "red", "FFC000": "orange"}
per_sheet = {}
for sh in spec.get("year_axis") or {}:
    if sh not in wb.sheetnames or sh not in rb.sheetnames: continue
    tc = year_columns(spec, sh).get(str(year))
    if not tc: continue
    ws, rs = wb[sh], rb[sh]
    for r in range(1, ws.max_row + 1):
        c = ws[f"{tc}{r}"]
        if not is_input(c.value): continue
        try: dv = ev.cell(sh, f"{tc}{r}")
        except Exception: dv = None
        try: rv = rev.cell(sh, f"{tc}{r}")
        except Exception: rv = None
        if not isinstance(rv, (int, float)) or isinstance(rv, bool):
            continue                      # the reference holds nothing to judge against
        tot += 1; per_sheet.setdefault(sh, [0, 0])[0] += 1
        tol = max(0.5, abs(rv) * 0.005)
        if isinstance(dv, (int, float)) and abs(dv - rv) <= tol:
            ok += 1; per_sheet[sh][1] += 1
        else:
            wrong.append((sh, f"{tc}{r}", str(ws.cell(r, 1).value)[:28], dv, rv, colours.get(fill(c), "none")))
print(f"reference: {ref.split('/')[-1]}")
print(f"1. filled input cells judged: {tot}; correct: {ok} ({100*ok/tot if tot else 0:.0f}%)")
n_red = sum(1 for w in wrong if w[5] == "red"); n_or = sum(1 for w in wrong if w[5] == "orange"); n_none = sum(1 for w in wrong if w[5] == "none")
print(f"2. wrong: {len(wrong)} — red {n_red}, orange {n_or}")
print(f"3. wrong and NOT highlighted: {n_none}")
print("per sheet (judged, correct):", per_sheet)
for w in sorted(wrong, key=lambda x: (x[5] != "none", x[0], x[1])):
    sh, co, lab, dv, rv, col = w
    print(f"   {col:6s} {sh}!{co:6s} {lab:28s} agent={dv if not isinstance(dv,float) else round(dv,2)!s:>14} ref={round(rv,2):>14}")
