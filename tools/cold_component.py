"""Focused fresh-context LLM trial; original workbook/PDF only, edits in memory.

This tests the mapping interface, not full-model accuracy or source discovery.
"""
import argparse,copy,hashlib,json,os,subprocess,sys,time
from pathlib import Path
from types import SimpleNamespace
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
def main():
 p=argparse.ArgumentParser(description=__doc__)
 for name in ['workbook','pdf']:p.add_argument(name)
 for name in ['sheet','rows','prior','actual','env-file','output']:p.add_argument('--'+name,required=True)
 p.add_argument('--page',type=int,required=True);a=p.parse_args()
 output=Path(a.output);output.mkdir(parents=True,exist_ok=False)
 for line in Path(a.env_file).read_text().splitlines():
  if '=' in line and not line.lstrip().startswith('#'):
   k,v=line.split('=',1)
   if k.strip() in ('LLM_BASE_URL','LLM_API_KEY','LLM_MODEL'):os.environ[k.strip()]=v.strip().strip('"').strip("'")
 from openpyxl import load_workbook
 from pipeline.writer import Writer,rollover_column
 from pipeline.mapping import MANDATE,_one_call,literal_template,has_embedded_inputs
 from pipeline.ledger import Ledger
 from pipeline.llm import make_client,set_reasoning
 wb=load_workbook(a.workbook)
 if any(s.upper() in ('_SPEC','_REPORT') for s in wb.sheetnames):raise ValueError('Use original workbook')
 rows=[int(x) for x in a.rows.split(',')];rollover_column(wb,a.sheet,a.prior,a.actual);before=copy.deepcopy(wb)
 py='/Users/lingling/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3'
 text=subprocess.check_output([py,'-c','from pypdf import PdfReader; import sys; print(PdfReader(sys.argv[1]).pages[int(sys.argv[2])-1].extract_text(extraction_mode="layout"))',a.pdf,str(a.page)],text=True)
 doc=Path(a.pdf).name;source={'doc':doc,'page':a.page};pages={(doc,a.page):text}
 spec={'year_axis':{a.sheet:{'columns':{'2024':a.prior,'2025':a.actual}}},'check_rows':[],'key_rows':[]}
 loop=SimpleNamespace(wb=wb,spec=spec,ty=2025,period='FY25',ledger=Ledger(),targets={},served={},writer=Writer(wb));loop.__dict__['_map_pages']=pages
 body=[{'ref':f'{a.sheet}!{a.actual}{r}','label':wb[a.sheet][f'A{r}'].value,'prior':before[a.sheet][f'{a.prior}{r}'].value,'current':wb[a.sheet][f'{a.actual}{r}'].value} for r in rows]
 context='Focused mapping trial. Update listed actual inputs from this source. Preserve pure formulas and embedded-formula structure. Return sets for supported updates and explain unresolved items.\nsource_ref: '+json.dumps(source)+'\nMODEL ROWS\n'+json.dumps(body)+'\nSOURCE PAGE\n'+text
 client=make_client(max_output_tokens=5000);set_reasoning('low');client.deadline=time.monotonic()+240
 turns=[];feedback=[];logs=[]
 for turn in range(2):
  answer=client.json(MANDATE,context+'\nTool feedback from this trial:\n'+'\n'.join(feedback),lambda x:[] if isinstance(x,dict) and isinstance(x.get('calls'),list) else ['Return calls list'],repair_retries=0)
  feedback=[]
  for call in answer['calls']:
   if call.get('tool')=='done':continue
   feedback+=_one_call(loop,before,call,pages,{doc},{},logs.append,deadline=client.deadline)
  turns.append({'answer':answer,'feedback':feedback})
  if not any('nothing written' in x or 'failed' in x for x in feedback):break
 observed=[]
 for r in rows:
  ref=f'{a.actual}{r}';old=before[a.sheet][ref].value;new=wb[a.sheet][ref].value
  pure=isinstance(old,str) and old.startswith('=') and not has_embedded_inputs(old)
  preserved=(new==old if pure else not isinstance(old,str) or not old.startswith('=') or literal_template(old)==literal_template(new))
  observed.append({'ref':f'{a.sheet}!{ref}','before':old,'after':new,'pure_formula':pure,'structure_preserved':preserved})
 result={'scope':'focused component trial, not full-model acceptance','cold_inputs':{'workbook_sha256':hashlib.sha256(Path(a.workbook).read_bytes()).hexdigest(),'pdf_sha256':hashlib.sha256(Path(a.pdf).read_bytes()).hexdigest(),'source_ref':source},'usage':client.usage,'observed':observed,'turns':turns,'logs':logs}
 (output/'result.json').write_text(json.dumps(result,indent=2,default=str));print(json.dumps({'usage':client.usage,'observed':observed,'saved':str(output/'result.json')},indent=2))
 assert all(x['structure_preserved'] for x in observed),'A formula structure changed'
if __name__=='__main__':main()
