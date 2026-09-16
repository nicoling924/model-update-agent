"""DRY TEST (owner 2026-09-16): can the brain, shown the MODEL instead of a card,
find what a run got wrong? No pipeline change — one prompt, one answer.

Usage: python tools/drytest_review.py <artifact_dir> [<year>]
  <artifact_dir> holds model/<...> (pipeline).xlsx, model-archive/*_pre.xlsx,
  replay/<PERIOD>/writes.json, key_panel.json, key_rows.json.
Prints the context it built, the brain's answer, and a scorecard against the
known faults of run 35066977462 (HK Sales!AI16, Final!AI95/AI97, Final!99 +223)."""
import glob
import json
import os
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from agent.evaluator import Evaluator  # noqa: E402

HIST = ["AE", "AF", "AG", "AH"]          # 2021..2024 on this model
ACT, FC = "AI", ["AJ", "AK", "AL", "AM", "AN"]
HEADLINE = [(7, "Total revenue"), (10, "EBITDA"), (15, "Net operating income"), (20, "EBIT"),
            (27, "Net profit (reported)"), (31, "Net profit (recurring)"), (35, "EPS"),
            (57, "Cash and equivalents"), (75, "Total assets"), (96, "Total shareholders' equity"),
            (98, "Total liabilities and equity"), (121, "Operating cash flow"),
            (129, "Investing cash flow"), (145, "Free cash flow")]


def num(x):
    try:
        return round(float(x), 2)
    except (TypeError, ValueError):
        return None


def fmt(x):
    return "" if x is None else (f"{x:,.2f}" if abs(x) < 100 and x != int(x) else f"{x:,.0f}")


def label_of(ws, r):
    for c in ("A", "B", "C", "D"):
        v = ws[f"{c}{r}"].value
        if isinstance(v, str) and v.strip():
            return v.strip()[:40]
    return f"row {r}"


def fill_of(cell):
    try:
        rgb = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None
    except Exception:
        rgb = None
    return {"FFFFC7CE": "red", "FFFFC000": "orange"}.get(rgb or "", "plain")


def build_context(art, period):
    delivered = next(p for p in glob.glob(f"{art}/model/*.xlsx") if "pipeline" in p)
    pre = glob.glob(f"{art}/model-archive/*_pre.xlsx")[0]
    rp = f"{art}/replay/{period}"
    wb, wb0 = openpyxl.load_workbook(delivered), openpyxl.load_workbook(pre)
    ev, ev0 = Evaluator(wb), Evaluator(wb0)
    kp = json.load(open(f"{rp}/key_panel.json"))
    krows = json.load(open(f"{rp}/key_rows.json"))
    writes = json.load(open(f"{rp}/writes.json"))
    L = []
    L.append("## 1. OBJECTIVES, measured by code")
    L.append("Balance check Final!99 (must be 0) by year: " + ", ".join(
        f"{c}={fmt(num(ev.cell('Final', f'{c}99')))}" for c in HIST + [ACT] + FC))
    L.append("ROAFNA check ROAFNA!31: " + ", ".join(f"{c}={fmt(num(ev.cell('ROAFNA', f'{c}31')))}" for c in HIST + [ACT]))
    L.append("Keys (model now vs the print):")
    for k in krows:
        v = num(ev.cell(k["sheet"], f"{ACT}{k['row']}"))
        p = kp.get(k["name"], {}).get("print")
        gap = None if (v is None or p is None) else round(v - p, 2)
        L.append(f"  {k['name']:24} {k['sheet']}!{ACT}{k['row']:<4} model {fmt(v):>12}  print {fmt(p):>12}  gap {fmt(gap)}")
    L.append("Cash Final!57 next years (must not be negative): " + ", ".join(f"{c}={fmt(num(ev.cell('Final', f'{c}57')))}" for c in [ACT] + FC))
    L.append("Total assets Final!75 next years: " + ", ".join(f"{c}={fmt(num(ev.cell('Final', f'{c}75')))}" for c in [ACT] + FC[:2]))
    L.append("")
    L.append("## 2. HEADLINE LINES (history 2021-2024 | the analyst's pre-update estimate for 2025 | now 2025 | 2026 forecast before -> now)")
    for r, name in HEADLINE:
        hist = [fmt(num(ev0.cell("Final", f"{c}{r}"))) for c in HIST]
        est, now = fmt(num(ev0.cell("Final", f"{ACT}{r}"))), fmt(num(ev.cell("Final", f"{ACT}{r}")))
        f0, f1 = fmt(num(ev0.cell("Final", f"AJ{r}"))), fmt(num(ev.cell("Final", f"AJ{r}")))
        L.append(f"  Final!{r:<4} {name:28} {' | '.join(hist)} | est {est} | now {now} | 2026 {f0} -> {f1}")
    L.append("")
    L.append("## 3. EVERY CELL THE RUN WROTE IN THE ACTUAL COLUMN (sorted by the move against the cell's own history, largest first)")
    rows = []
    first = {}                      # one entry per cell: the FIRST 'old' (the analyst's) and the final value
    for w in writes:
        sh, co = w.get("sheet"), w.get("coord", "")
        if co.startswith(ACT) and sh in wb.sheetnames and (sh, co) not in first:
            first[(sh, co)] = w
    for (sh, co), w in first.items():
        r = int(co[len(ACT):])
        if r <= 3:
            continue
        ws = wb[sh]
        old, new = num(w.get("old")), num(ev.cell(sh, co))
        if new is None:
            continue
        hist = [num(ev0.cell(sh, f"{c}{r}")) for c in HIST]
        hist = [h for h in hist if h is not None]
        lo, hi = (min(hist), max(hist)) if hist else (None, None)
        if hist and (hi - lo) > 0:
            span = hi - lo
            out = max(lo - new, new - hi, 0) / max(abs(hi), abs(lo), 1)
        elif hist:
            out = abs(new - hi) / max(abs(hi), 1)
        else:
            out = -1.0                 # no history to judge against — listed last, not first
        cell = ws[co]
        note = (cell.comment.text if cell.comment else (w.get("before") or ["", ""])[1] or "") or ""
        rows.append((out, sh, co, label_of(ws, r), old, new, hist, fill_of(cell), " ".join(str(note).split())[:160]))
    rows.sort(key=lambda t: -t[0])
    for out, sh, co, lab, old, new, hist, flag, note in rows[:260]:
        L.append(f"  {sh}!{co:<5} {lab:40} was {fmt(old):>11} -> now {fmt(new):>11} | history {', '.join(fmt(h) for h in hist):32} | {flag:6} | {note}")
    L.append(f"  ({len(rows)} written cells in total)")
    L.append("")
    L.append("## 4. TRACE — for each headline line whose 2026 forecast moved more than 10% vs the pre-update book, "
             "the typed inputs that carry the move (share of the swing when that input's pre-update content is put back)")
    from pipeline.investigate import swing_leaves
    for r, name in HEADLINE:
        f0, f1 = num(ev0.cell("Final", f"AJ{r}")), num(ev.cell("Final", f"AJ{r}"))
        if f0 is None or f1 is None or abs(f0) < 1 or abs(f1 - f0) / abs(f0) < 0.10:
            continue
        try:
            leaves = swing_leaves(wb, wb0, "Final", f"AJ{r}", budget_s=60)[:6]
        except Exception as e:
            L.append(f"  Final!AJ{r} {name}: trace failed ({e!r})")
            continue
        parts = []
        for (sh, co), share in leaves:
            parts.append(f"{sh}!{co} '{label_of(wb[sh], int(''.join(ch for ch in co if ch.isdigit())))}' "
                         f"{fmt(num(ev0.cell(sh, co)))} -> {fmt(num(ev.cell(sh, co)))} ({share:+.0%})")
        L.append(f"  Final!AJ{r} {name} {fmt(f0)} -> {fmt(f1)}: " + "; ".join(parts))
    return "\n".join(L)


SYSTEM = ("You are the equity analyst reviewing your own model update before delivery. The model must "
          "(1) balance every year, (2) tie every key number to the print, (3) roll forward sanely — no negative "
          "cash or assets, no headline line moving out of its history without a printed reason. A wrong number "
          "is a failure; so is a wrong number you could have found and left. You are shown the model's objectives "
          "measured by code, the headline lines, and every cell the run wrote with its move against its own "
          "history. Reason like an analyst: when an objective is off, find the INPUT that caused it — look first "
          "at what moved against its history — and say what to change, to what, and why. Compensating errors "
          "come in pairs; name both. Answer in JSON only.")

ASK = ("\n\nReturn JSON: {\"findings\": [{\"cell\": \"Sheet!AI16\", \"problem\": \"...\", \"change_to\": number or "
       "null, \"because\": \"...\"}], \"objectives\": {\"balance\": \"holds | broken because ...\", "
       "\"keys\": \"...\", \"rollforward\": \"...\"}}. List the findings in order of impact. Be specific: cells, "
       "values, and the reason an analyst would accept.")


def main():
    art = sys.argv[1]
    period = os.path.basename(glob.glob(f"{art}/replay/*/writes.json")[0].rsplit("/", 2)[-2]) if len(sys.argv) < 3 else sys.argv[2]
    ctx = build_context(art, period)
    print("=" * 30, "CONTEXT", "=" * 30)
    print(ctx)
    print(f"\n[context: {len(ctx):,} chars ≈ {len(ctx)//4:,} tokens]")
    from pipeline.llm import make_client, handshake
    client = make_client(temperature=0.1, max_output_tokens=6000)
    handshake(client)

    def _val(o):
        return [] if isinstance(o, dict) and isinstance(o.get("findings"), list) else ["findings list missing"]
    ans = client.json(SYSTEM, ctx + ASK, _val, repair_retries=1)
    print("=" * 30, "BRAIN", "=" * 30)
    print(json.dumps(ans, indent=1, ensure_ascii=False))
    text = json.dumps(ans)
    print("=" * 30, "SCORECARD (run 35066977462 faults)", "=" * 30)
    for name, needles in [("fuel clause HK Sales!AI16 (44.3 -> 1.00)", ["HK Sales!AI16", "Fuel Clause", "fuel clause"]),
                          ("RE/MI pair Final!AI95 / Final!AI97", ["AI95", "AI97", "retained", "Retained", "minority", "Minority", "perpetual"]),
                          ("balance +223", ["223"])]:
        print(f"  {'FOUND ' if any(n in text for n in needles) else 'MISSED'}  {name}")


if __name__ == "__main__":
    main()
