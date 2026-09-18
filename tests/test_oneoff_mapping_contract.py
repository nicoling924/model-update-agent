"""The mapper's semantic event decision reaches forecast-intent preservation."""
import copy
import runpy
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from pipeline.mapping import _apply
from pipeline.freeze import hold_oneoff_forecasts
from pipeline.teachings import oneoff_watch
from pipeline.writer import Writer
from pipeline.orchestrator import ObjectiveLoop
from pipeline.ledger import Ledger
import openpyxl


def test_mapped_event_preserves_zero_future_but_leaves_stock_roll_live():
    fixture = runpy.run_path(str(Path(__file__).with_name('test_pipeline_museum.py')))
    loop, pages, _ = fixture['_map_model']()
    ws = loop.wb['Final']
    loop.spec['year_axis']['Final']['columns'].update({'2026':'D','2027':'E'})
    ws['C3'] = 0; ws['D3'] = '=C3'; ws['E3'] = '=D3'
    ws['C11'] = 100; ws['D11'] = '=C11+D3'; ws['E11'] = '=D11+E3'
    before = copy.deepcopy(loop.wb)
    _apply(loop, [{'sheet':'Final','coord':'C3','printed':460,'doc':'ar.pdf','page':23,
                  'line':'Other gains, net 460 420','one_off':True,
                  'one_off_reason':'Discrete disposal gain; future event forecasts were zero.'}],
           pages, {'ar.pdf'}, lambda *_: None)
    # A rollover may materialize the event as a number; the original zero still owns intent.
    ws['E3'] = 460
    event = loop.writer.log['oneoff_inputs']['Final!3']
    hold_oneoff_forecasts(loop.wb,before,loop.spec,2025,'Final',3,loop.writer,event['because'])
    assert ws['C3'].value == 460
    assert ws['D3'].value == 0 and ws['E3'].value == 0
    assert ws['D11'].value == '=C11+D3' and ws['E11'].value == '=D11+E3'


def _unmapped_event_loop():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Final'
    ws['A3'] = 'Discrete event'
    ws['B3'] = 0
    ws['C3'] = 460
    ws['D3'] = '=C3'
    spec = {'year_axis': {'Final': {'columns': {
        '2024': 'B', '2025': 'C', '2026': 'D'}}}}
    writer = Writer(wb)
    loop = ObjectiveLoop(wb, spec, 2025, Ledger(), [], {}, writer, None)
    return loop


def test_unmapped_event_becomes_review_candidate_without_auto_hold():
    loop = _unmapped_event_loop()
    assert oneoff_watch(loop.wb, loop.spec, 2025, loop.writer) == 1
    candidate = loop.writer.log['oneoff_candidates']['Final!C3']
    assert candidate['forecast'] == ['Final!D3']
    assert candidate['label'] == 'Discrete event'
    assert candidate['actual_value'] == 460 and candidate['prior_value'] == 0
    assert 'Final!C3' not in loop.writer.log.get('oneoff_inputs', {})
    state = loop._state_block()
    assert 'ONE-OFF CANDIDATES' in state
    assert 'Final!C3' in state


def test_event_classification_false_keeps_links_and_true_reuses_preservation_path():
    loop = _unmapped_event_loop()
    oneoff_watch(loop.wb, loop.spec, 2025, loop.writer)
    false = loop.t_classify_event({'candidate': 'Final!C3', 'one_off': False,
                                   'reason': 'Recurring item; no discrete event evidence.'})
    assert 'RECURRING' in false
    assert 'Final!C3' not in loop.writer.log.get('oneoff_inputs', {})
    assert loop.wb['Final']['D3'].value == '=C3'

    loop2 = _unmapped_event_loop()
    oneoff_watch(loop2.wb, loop2.spec, 2025, loop2.writer)
    true = loop2.t_classify_event({'candidate': 'Final!C3', 'one_off': True,
                                   'reason': 'Discrete disposal event; future forecast was zero.'})
    assert 'ONE-OFF' in true
    event = loop2.writer.log['oneoff_inputs']['Final!C3']
    assert event['row'] == 3 and event['sheet'] == 'Final'
    before = openpyxl.Workbook()
    before.active.title = 'Final'
    before['Final']['B3'] = 0
    before['Final']['C3'] = 0
    before['Final']['D3'] = 0
    hold_oneoff_forecasts(loop2.wb, before, loop2.spec, 2025, 'Final', 3,
                          loop2.writer, event['because'])
    assert loop2.wb['Final']['D3'].value == 0


def test_preview_finds_event_when_prior_was_nonzero_and_excludes_output_rows():
    from pipeline.teachings import oneoff_watch
    pre = openpyxl.Workbook(); pre.active.title = 'Final'
    pws = pre['Final']
    pws['A3'] = 'Capacity event'; pws['B3'] = 5; pws['C3'] = 0; pws['D3'] = 0
    pws['A4'] = 'Accounting output'; pws['B4'] = 5; pws['C4'] = 0; pws['D4'] = 0
    wb = openpyxl.Workbook(); wb.active.title = 'Final'
    ws = wb['Final']
    ws['A3'] = 'Capacity event'; ws['B3'] = 5; ws['C3'] = 460; ws['D3'] = '=C3'
    ws['A4'] = 'Accounting output'; ws['B4'] = 5; ws['C4'] = 460; ws['D4'] = '=C3'
    spec = {'year_axis': {'Final': {'columns': {
        '2024': 'B', '2025': 'C', '2026': 'D'}}}}
    writer = Writer(wb)
    before = (ws['D3'].value, ws['D4'].value)
    n = oneoff_watch(wb, spec, 2025, writer, pre_wb=pre)
    assert n == 1
    assert list(writer.log['oneoff_candidates']) == ['Final!C3']
    assert writer.log['oneoff_candidates']['Final!C3']['forecast'] == ['Final!D3']
    assert (ws['D3'].value, ws['D4'].value) == before

def test_production_review_routes_event_classification_and_applies_hold():
    from pipeline.review import build_context, _one_call
    loop = _unmapped_event_loop()
    pre = copy.deepcopy(loop.wb)
    pre['Final']['C3'] = 0
    oneoff_watch(loop.wb, loop.spec, 2025, loop.writer, pre_wb=pre)
    context = build_context(loop, pre, {}, None)
    assert 'Final!C3' in context and 'classify_event' in context
    def repair(_tag):
        for event in loop.writer.log.get('oneoff_inputs', {}).values():
            hold_oneoff_forecasts(loop.wb, pre, loop.spec, 2025,
                                  event['sheet'], event['row'], loop.writer, event['because'])
    lines, result = _one_call(loop, pre,
        {'tool': 'classify_event', 'candidate': 'Final!C3', 'one_off': True,
         'reason': 'Discrete event; original future forecast was zero'},
        {}, None, lambda *a: None, repair, lambda: (True, [], {}), None, {})
    assert lines[0].startswith('CLASSIFIED ONE-OFF')
    assert loop.wb['Final']['D3'].value == 0 and result[0]


if __name__ == '__main__':
    test_mapped_event_preserves_zero_future_but_leaves_stock_roll_live()
    test_unmapped_event_becomes_review_candidate_without_auto_hold()
    test_event_classification_false_keeps_links_and_true_reuses_preservation_path()
    test_production_review_routes_event_classification_and_applies_hold()
    print('PASS mapped and unmapped event forecast contracts')
