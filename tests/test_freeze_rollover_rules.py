import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import openpyxl

from pipeline.freeze import apply_freezes, plan_freezes
from pipeline.writer import Writer


def _books():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A5"] = "Growth assumption"
    ws["A6"] = "Zero input"
    ws["A7"] = "Nonzero input"
    ws["A8"] = "Absent pre-update"
    ws["V8"] = 0
    ws["V9"] = "=0"
    ws["U5"] = 0.30
    ws["V5"] = "=U5"
    ws["W5"] = "=V5"
    ws["V6"] = 0
    ws["W6"] = "=V6+1"
    ws["V7"] = 5
    for ref in ("V5", "W5"):
        ws[ref].number_format = "0%"
    values = openpyxl.Workbook()
    values.active.title = "S"
    values["S"]["V5"] = 0.30
    values["S"]["W5"] = 0.30
    values["S"]["V6"] = 0
    values["S"]["W6"] = 1
    values["S"]["V7"] = 5
    values["S"]["V9"] = 0
    return wb, values


def test_only_first_forecast_growth_is_frozen_from_pre_update_value():
    wb, values = _books()
    plans = plan_freezes(wb, values, ["S"], 21, horizon=2)
    assert [(p["coord"], p["value"]) for p in plans if p.get("reason") != "zero_forecast_hold"] == [("V5", 0.30)]
    assert wb["S"]["W5"].value == "=V5"


def test_plan_freezes_keeps_growth_rule_but_does_not_promote_zero_rows():
    wb, values = _books()
    plans = plan_freezes(wb, values, ["S"], 21, horizon=2,
                         check_rows=[{"sheet": "S", "row": 9}])
    zero = [p for p in plans if p.get("reason") == "zero_forecast_hold"]
    assert zero == []
    assert all(p["coord"] != "V7" for p in plans)
    assert all(p["coord"] != "W6" for p in plans)
    assert all(p["coord"] != "V8" for p in plans)


def test_explicit_oneoff_hold_requires_independent_input_structure():
    from pipeline.freeze import hold_oneoff_forecasts
    wb, pre = _books()
    ws, pws = wb["S"], pre["S"]
    ws["A11"], pws["A11"] = "One-off event", "One-off event"
    ws["T11"], pws["T11"] = 2, 2
    ws["U11"], pws["U11"] = 2, 2
    ws["V11"], pws["V11"] = 0, 0
    ws["W11"], pws["W11"] = 0, 0
    ws["A12"], pws["A12"] = "Stock balance", "Stock balance"
    ws["T12"], pws["T12"] = "=T2+T3", "=T2+T3"
    ws["U12"], pws["U12"] = "=U2+U3", "=U2+U3"
    ws["V12"], pws["V12"] = 0, 0
    ws["W12"], pws["W12"] = 0, 0
    ws["A13"], pws["A13"] = "Missing", "Missing"
    ws["V13"], pws["V13"] = 0, 0
    pws["V13"] = None
    ws["A14"], pws["A14"] = "Cycle", "Cycle"
    ws["T14"], pws["T14"] = 1, 1
    ws["U14"], pws["U14"] = 1, 1
    ws["V14"], pws["V14"] = "=W14", "=W14"
    ws["W14"], pws["W14"] = "=V14", "=V14"
    ws["A15"], pws["A15"] = "Event carry", "Event carry"
    ws["T15"], pws["T15"] = 1, 1
    ws["U15"], pws["U15"] = 1, 1
    ws["V15"], pws["V15"] = "=U15-1", "=U15-1"
    ws["W15"], pws["W15"] = "=V15", "=V15"
    spec = {"year_axis": {"S": {"columns": {"2024": "T", "2025": "U", "2026": "V", "2027": "W"}}}}
    writer = Writer(wb); writer.plugs_allowed = True
    lines = hold_oneoff_forecasts(wb, pre, spec, 2025, "S", 11, writer,
                                  "explicit one-off event")
    for held_row in (12, 13, 14, 15):
        lines += hold_oneoff_forecasts(wb, pre, spec, 2025, "S", held_row,
                                       writer, "explicit one-off event")
    assert [p for p in lines if "V11" in p or "W11" in p]
    assert ws["V11"].value == 0 and ws["W11"].value == 0
    assert ws["V12"].value == 0 and ws["V12"].fill.fgColor.rgb.endswith("BDD7EE") is False
    assert ws["V13"].value == 0 and ws["V14"].value == "=W14"
    assert ws["V15"].value == 0 and ws["W15"].value == 0
    assert ws["V11"].comment is not None
    assert "explicit one-off classification" in ws["V11"].comment.text
    again = hold_oneoff_forecasts(wb, pre, spec, 2025, "S", 11, writer,
                                  "explicit one-off event")
    assert again == []


def test_oneoff_formula_proof_rejects_names_functions_and_cross_row_outputs():
    from pipeline.freeze import hold_oneoff_forecasts
    wb, pre = _books()
    ws, pws = wb["S"], pre["S"]
    for row, formula in ((16, "=NamedThing+1"), (17, "=SUM(U17)"),
                         (18, "=U18+U2"), (19, "=Other!U19")):
        ws[f"U{row}"] = pws[f"U{row}"] = 1
        ws[f"V{row}"] = pws[f"V{row}"] = formula
        ws[f"W{row}"] = pws[f"W{row}"] = "=V%d" % row
    spec = {"year_axis": {"S": {"columns": {"2024": "T", "2025": "U",
                                               "2026": "V", "2027": "W"}}}}
    writer = Writer(wb); writer.plugs_allowed = True
    for row in (16, 17, 18, 19):
        assert hold_oneoff_forecasts(wb, pre, spec, 2025, "S", row, writer,
                                     "classified event") == []
        assert ws[f"V{row}"].value == pws[f"V{row}"].value


def test_oneoff_same_sheet_qualified_event_carry_is_allowed():
    from pipeline.freeze import hold_oneoff_forecasts
    wb, pre = _books()
    ws, pws = wb["S"], pre["S"]
    ws["T20"] = pws["T20"] = 1
    ws["U20"] = pws["U20"] = 1
    ws["V20"] = pws["V20"] = "=S!U20-1"
    ws["W20"] = pws["W20"] = "=S!V20"
    spec = {"year_axis": {"S": {"columns": {"2024": "T", "2025": "U",
                                               "2026": "V", "2027": "W"}}}}
    writer = Writer(wb); writer.plugs_allowed = True
    lines = hold_oneoff_forecasts(wb, pre, spec, 2025, "S", 20, writer,
                                  "classified event")
    assert any("V20" in line for line in lines)
    assert any("W20" in line for line in lines)


def test_oneoff_hold_requires_brain_classification():
    from pipeline.freeze import hold_oneoff_forecasts
    wb, pre = _books(); wb["S"]["T11"] = 1; pre["S"]["T11"] = 1
    wb["S"]["U11"] = 1; pre["S"]["U11"] = 1
    wb["S"]["V11"] = 0; pre["S"]["V11"] = 0
    writer = Writer(wb); writer.plugs_allowed = True
    assert hold_oneoff_forecasts(
        wb, pre,
        {"year_axis": {"S": {"columns": {"2024": "T", "2025": "U", "2026": "V"}}}},
        2025, "S", 11, writer, "") == []
    assert wb["S"]["V11"].value == 0


def test_check_rows_are_protected_and_authorized_formula_holds_use_over_formula():
    wb, values = _books()
    plans = plan_freezes(wb, values, ["S"], 21, check_rows=[{"sheet": "S", "row": 9}])
    assert all(p["coord"] != "V9" for p in plans)
    writer = Writer(wb)
    applied = apply_freezes(wb, plans, writer=writer)
    assert wb["S"]["V5"].value == 0.30
    assert wb["S"]["V6"].value == 0
    assert wb["S"]["V9"].value == "=0"
    assert any("V5" in line for line in applied)


def test_freeze_survives_later_writes_and_uses_the_actual_next_axis_column():
    wb, values = _books()
    wb["S"]["X5"] = "=U5"; wb["S"]["X5"].number_format = "0%"
    values["S"]["X5"] = .3
    plans = plan_freezes(wb, values, ["S"], 21, forecast_col=24)
    assert [p["coord"] for p in plans] == ["X5"]
    writer = Writer(wb); apply_freezes(wb, plans, writer)
    assert not writer.write("S","X5",.8,trusted=True,force_lock=True)
    assert wb["S"]["X5"].value == .3

if __name__ == "__main__":
    tests = [(name, fn) for name, fn in globals().copy().items() if name.startswith("test_") and callable(fn)]
    for name, fn in tests:
        fn()
        print("PASS", name)
    print(f"{len(tests)} tests passed")
