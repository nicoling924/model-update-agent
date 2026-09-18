import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
"""Focused contract tests for the compact _REPORT page."""

import openpyxl

from pipeline.reportpage import build


def _book():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A2"], ws["T2"], ws["U2"], ws["V2"] = "year", 2024, 2025, 2026
    ws["A4"], ws["T4"], ws["U4"], ws["V4"] = "Revenue", 100, 120, 130
    return wb


def _pre():
    wb = _book()
    ws = wb["Model"]
    ws["T4"], ws["U4"], ws["V4"] = 100, 110, 121
    return wb


def test_production_report_contains_only_changed_block():
    wb, pre = _book(), _pre()
    spec = {"year_axis": {"Model": {"columns": {
        "2024": "T", "2025": "U", "2026": "V"}}},
            "key_rows": [], "check_rows": []}
    build(wb, pre, spec, 2025, "FY25", {"units": "model units"},
          log=lambda *_a, **_k: None)
    values = [str(c.value) for row in wb["_REPORT"].iter_rows()
              for c in row if c.value is not None]
    text = "|".join(values)
    assert "WHAT'S CHANGED — new vs old" in text
    assert "NEW — after the update" in text
    assert "OLD — before the update" in text
    assert "Mini P&L" not in text
    assert "Key number snapshot" not in text
    assert "FLAGGED RED" not in text
    assert "=Model!U4" in text and "=Model!T4" in text
    assert "Look here" not in text
    assert "Updated to" not in text
    assert "Documents received" not in text

if __name__ == "__main__":
    tests = [(name, fn) for name, fn in globals().copy().items() if name.startswith("test_") and callable(fn)]
    for name, fn in tests:
        fn()
        print("PASS", name)
    print(f"{len(tests)} tests passed")
