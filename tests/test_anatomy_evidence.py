"""A numeric reading cannot redefine an input as a zero-balance objective."""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import unittest
from types import SimpleNamespace
from openpyxl import Workbook
from pipeline.writer import Writer
from pipeline.ledger import Ledger
from pipeline.mapping import _t_anatomy
from pipeline.anatomy import check_identity


class AnatomyEvidence(unittest.TestCase):
    def model(self):
        wb = Workbook(); ws = wb.active; ws.title = "Model"
        for col in "BCDE":
            for row, value in [(2,90),(3,60),(4,30),(7,0)]:
                ws[f"{col}{row}"] = value
            ws[f"{col}5"] = f"={col}2-{col}3-{col}4"
            ws[f"{col}6"] = f"={col}2*1000"
        spec = {"year_axis":{"Model":{"columns":{"2022":"B","2023":"C","2024":"D","2025":"E"}}},
                "check_rows":[], "key_rows":[]}
        return SimpleNamespace(wb=wb,spec=spec,ty=2025,writer=Writer(wb),ledger=Ledger(),targets={},served={})

    def test_only_reconciliation_evidence_can_promote_a_check(self):
        loop = self.model()
        _t_anatomy(loop, {"checks":[
            {"ref":"Model!5","identity":"Assets less liabilities and equity"},
            {"ref":"Model!2","identity":"Assets"},
            {"ref":"Model!6","identity":"An emissions quantity"},
            {"ref":"Model!7","identity":"A zero forecast input"}]}, lambda *a:None)
        self.assertEqual([c['row'] for c in loop.spec['check_rows']], [5])

    def test_bare_reference_does_not_supply_identity_evidence(self):
        loop = self.model()
        _t_anatomy(loop, {"checks":["Model!5"]}, lambda *a:None)
        self.assertEqual(loop.spec['check_rows'], [])

    def test_model_constants_do_not_disqualify_a_proven_ratio_identity(self):
        loop = self.model(); loop.wb['Model']['E8'] = '=E2/(E3+E4)-1'
        self.assertTrue(check_identity(loop.wb,'Model','E8','Assets divided by financing equals one',{8})[0])

    def test_operating_metrics_cannot_become_financial_headlines(self):
        loop = self.model()
        _t_anatomy(loop, {"keys":[{"name":"SO2","ref":"Model!6"}]}, lambda *a:None)
        self.assertEqual(loop.spec['key_rows'], [])


if __name__ == '__main__': unittest.main()
