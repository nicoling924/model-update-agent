"""Independent pins for the source -> proposal -> write -> objective contract."""
import copy
import importlib.util
from pathlib import Path
import sys
import unittest
from unittest.mock import patch
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
spec = importlib.util.spec_from_file_location("museum", Path(__file__).with_name("test_pipeline_museum.py"))
museum = importlib.util.module_from_spec(spec)
spec.loader.exec_module(museum)
from pipeline.mapping import _apply, is_own_arithmetic, input_rows
from pipeline.review import _one_call, _metrics
from pipeline.consequence import snapshot, restore

class Contract(unittest.TestCase):
    def test_key_repair_review_and_acceptance_share_the_printed_tie(self):
        from pipeline.keytie import key_state, matches_print
        from pipeline.review import _broken
        loop, _, _ = museum._map_model()
        loop.spec["key_rows"] = [{"name": "eps", "sheet": "Final", "row": 2}]
        loop.wb["Final"]["C2"] = 0.94
        panel = {"eps": {"print": 1.15}}
        self.assertFalse(key_state(loop.wb, loop.spec, 2025, None, panel)[0][-1])
        self.assertTrue(any(o[0] == "key" for o in _broken(loop, {}, panel, None)))
        for got, want in ((float("nan"), 1), (1, float("inf")), (True, 1), (-1.15, 1.15)):
            self.assertFalse(matches_print(got, want))
        loop.wb["Final"]["C2"] = 1.150001
        self.assertTrue(key_state(loop.wb, loop.spec, 2025, None, panel)[0][-1])
        self.assertFalse(any(o[0] == "key" for o in _broken(loop, {}, panel, None)))

    def test_a_repair_invalidates_only_its_cells_proof_and_rollback_restores_it(self):
        loop, _, _ = museum._map_model()
        w = loop.writer
        w.served = loop.served
        w.actual_cols = {"Final": "C"}
        loop.served[("Final", 2)] = {"value": 76061, "conf": 4, "source": {"doc": "ar.pdf"}}
        before = copy.deepcopy(loop.served)
        loop.wb["Final"]["D2"] = 80000
        w.write("Final", "D2", 81000, trusted=True)
        self.assertEqual(loop.served, before)
        snap = snapshot(loop)
        w.write("Final", "C2", 90000, trusted=True)
        self.assertNotIn(("Final", 2), loop.served)
        loop._map_cells = 20
        restore(loop, snap)
        self.assertEqual(loop.served, before)
        self.assertEqual(loop.wb["Final"]["C2"].value, 76061)
        self.assertEqual(loop._map_cells, 0)
        snap = snapshot(loop)
        w.write("Final", "C2", "=80000+10", trusted=True)
        restore(loop, snap)
        self.assertEqual(loop.wb["Final"]["C2"].value, 76061)
        self.assertEqual(loop.served, before)

    def test_acceptance_does_not_inherit_the_legacy_whole_unit_key_tolerance(self):
        from pipeline.acceptance import assess
        loop, _, _ = museum._map_model()
        for got, printed, expected in ((0.94, 1.15, "fail"), (4.6476, 4.14, "fail"),
                                       (4.14001, 4.14, "pass"), (-4.14, 4.14, "fail")):
            with patch("pipeline.acceptance.key_state", return_value=[("eps", "Final!C2", got, printed, True)]):
                result = assess(loop, None)
            self.assertEqual(result["keys"]["checks"][0]["status"], expected)

    def test_cross_sheet_and_accounting_rolls_keep_their_structure(self):
        loop, pages, census = museum._map_model()
        raw = loop.wb.create_sheet("Raw")
        raw["U3"] = 460
        ws = loop.wb["Final"]
        ws["C3"] = "=Raw!U3"
        ws["B4"], ws["C4"] = "=B2+B3", "=B4+C2-C3"
        self.assertTrue(is_own_arithmetic(loop.wb,"Final","C3","C","B"))
        self.assertTrue(is_own_arithmetic(loop.wb,"Final","C4","C","B"))
        out = _apply(loop,[{"sheet":"Final","coord":"C3","value":999}],pages,{"ar.pdf"},lambda *x:None)
        self.assertEqual(ws["C3"].value,"=Raw!U3")
        self.assertIn("own arithmetic",str(out))

    def test_disclosed_amount_has_one_units_contract_for_mapping_and_review(self):
        loop, pages, census = museum._map_model()
        ws=loop.wb["Final"]
        ws["B2"],ws["C2"]=100,100
        pages[("ar.pdf",23)]="Revenue 120,000,000 100,000,000\n"
        loop._map_scales={("ar.pdf",23):1e6}
        loop.page_text=pages
        entry={"sheet":"Final","coord":"C2","ref":"Final!C2","doc":"ar.pdf","page":23,
               "printed":120000000,"line":"Revenue 120,000,000 100,000,000","because":"revenue, same definition"}
        _apply(loop,[entry],pages,{"ar.pdf"},lambda *x:None)
        self.assertEqual(ws["C2"].value,120)
        self.assertEqual(loop.served[("Final",2)]["value"],120)
        pages[("ar.pdf",23)]="Revenue 125,000,000 100,000,000\n"
        entry.update(printed=125000000,line="Revenue 125,000,000 100,000,000")
        out,_=_one_call(loop,copy.deepcopy(loop.wb),{"tool":"set","sets":[entry]}, {},None,lambda *x:None,
                        lambda *x:None,lambda:(True,[],{}),None,{})
        self.assertEqual(ws["C2"].value,125,str(out))
        self.assertEqual(loop.served[("Final",2)]["value"],125)
        self.assertEqual(len(loop.writer.log["change_records"]),2)
        self.assertNotIn("Final!C2",loop.writer.log["flags"])

    def test_an_unwritable_partner_rolls_back_values_and_evidence(self):
        loop,pages,_=museum._map_model();loop.page_text=pages
        before=copy.deepcopy(loop.served)
        out,_=_one_call(loop,copy.deepcopy(loop.wb),{"tool":"set","sets":[
            {"ref":"Final!C3","value":460,"because":"proposed correction"},
            {"ref":"Final!C10","value":999,"because":"cannot replace subtotal"}]},
            {},None,lambda *x:None,lambda *x:None,lambda:(True,[],{}),None,{})
        self.assertEqual(loop.wb["Final"]["C3"].value,420)
        self.assertEqual(loop.served,before)
        self.assertEqual(loop.writer.log["change_records"],[])
        self.assertEqual(loop.__dict__["_map_written"],{})
        self.assertIn("none of it was kept",str(out))

    def test_distinct_objectives_on_same_cell_cannot_erase_each_other(self):
        loop,pages,_=museum._map_model()
        loop.wb["Final"]["C2"]=120
        with patch("pipeline.keytie.key_state",return_value=[("total assets","Final!C2",120,150,False)]), \
             patch("pipeline.review.sanity_rows",return_value={"total assets":("Final",2)}):
            rows=_metrics(loop,{},None)
        self.assertEqual(rows[("Final!C2","key","total assets")][3],-30)
        self.assertEqual(rows[("Final!C2","sanity","total assets")][3],0)
        with patch("pipeline.keytie.key_state",return_value=[("total assets","Final!C2",120,None,False)]):
            rows=_metrics(loop,{},None)
        self.assertIsNone(rows[("Final!C2","key","total assets")][3])

    def test_document_identity_is_used_when_reading_pages(self):
        loop,pages,_=museum._map_model();pages[("old.pdf",23)]="Old disclosure"
        loop.page_text=pages
        out,_=_one_call(loop,loop.wb,{"tool":"page","doc":"old.pdf","n":23},{},None,
                        lambda *x:None,None,None,None,{})
        self.assertEqual(out,["--- old.pdf p23 ---\nOld disclosure"])

    def test_embedded_actual_amounts_can_change_without_changing_the_formula(self):
        loop,pages,census=museum._map_model()
        ws=loop.wb["Final"]
        ws["B9"],ws["C9"]="=76061-20000", "=76061-20000"
        self.assertIn(("Final","C9",9),input_rows(loop,census))
        _apply(loop,[{"sheet":"Final","coord":"C9","formula":"=88018-21000",
                     "because":"same composition with current disclosed amounts"}],pages,{"ar.pdf"},lambda *x:None)
        self.assertEqual(ws["C9"].value,"=88018-21000")
        _apply(loop,[{"sheet":"Final","coord":"C9","formula":"=C2+C3",
                     "because":"attempt to replace the model logic"}],pages,{"ar.pdf"},lambda *x:None,correction=True)
        self.assertEqual(ws["C9"].value,"=88018-21000")

    def test_repair_and_measurement_use_the_same_definition_of_balanced(self):
        from pipeline.forecast_balance import last_resort_plug
        from pipeline.checks import CHECK_TOL
        from pipeline.evaluator import Evaluator
        from pipeline.writer import Writer
        import openpyxl
        for gap in (CHECK_TOL/2, CHECK_TOL*1.5, CHECK_TOL*2, CHECK_TOL*5):
            wb=openpyxl.Workbook(); ws=wb.active; ws.title="S"
            ws["A8"]="Cash flow statement"; ws["A9"]="Others"
            ws["B6"]="=0";ws["C6"]=f"={gap}+C9"
            ws["B9"]=0;ws["C9"]=0;ws["C3"]=10000
            writer=Writer(wb);writer.plugs_allowed=True
            changes=last_resort_plug(wb,writer,lambda:lambda sh,co:Evaluator(wb).cell(sh,co),
                                    "S",6,["C"],3,lambda *a:None)
            self.assertLessEqual(abs(Evaluator(wb).cell("S","C6")),CHECK_TOL)
            self.assertEqual(bool(changes),gap>CHECK_TOL)

    def test_snapshot_restores_nested_evidence(self):
        loop,_,_=museum._map_model()
        loop.served[("Final",2)]={"value":100,"source":{"doc":"original.pdf"}}
        snap=snapshot(loop)
        loop.served[("Final",2)]["source"]["doc"]="changed.pdf"
        restore(loop,snap)
        self.assertEqual(loop.served[("Final",2)]["source"]["doc"],"original.pdf")

    def test_input_rows_below_old_scan_limit_are_not_lost(self):
        loop,_,census=museum._map_model()
        ws=loop.wb["Final"];ws["A900"]="New segment";ws["B900"]=50;ws["C900"]="=B900*1.1"
        self.assertIn(("Final","C900",900),input_rows(loop,census))

if __name__=="__main__": unittest.main()
