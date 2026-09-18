"""Boundary tests for deterministic forecast balance repairs.

These tests deliberately exercise the paths that pass ``trusted`` or probe a
cell directly: formulas must remain formulas, while only typed numeric input
cells may be changed by a plug.
"""

import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import openpyxl

from pipeline.evaluator import Evaluator
from pipeline.forecast_balance import last_resort_plug, place_flow
from pipeline.writer import Writer


def _writer(wb):
    out = Writer(wb)
    out.plugs_allowed = True
    return out


def test_place_flow_refuses_formula_target_and_keeps_structure():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A8"] = "Cash flow statement"
    ws["A9"] = "Others"
    ws["B4"] = 100
    ws["C4"] = 120
    ws["C9"] = "=76061-20000"  # embedded inputs: updateable elsewhere, not a plug
    before = ws["C9"].value

    value, error = place_flow(wb, _writer(wb), "S", 4, 9, "C", "B")

    assert value is None
    assert "typed CF input" in error
    assert ws["C9"].value == before


def test_last_resort_plug_changes_only_typed_numeric_input():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A8"] = "Cash flow statement"
    ws["A9"] = "Others"
    ws["B1"] = 0
    ws["C1"] = "=10-C9"
    ws["B9"] = 0
    ws["C9"] = 0
    ws["C3"] = 1000
    writer = _writer(wb)

    changes = last_resort_plug(
        wb, writer, lambda: lambda sh, co: Evaluator(wb).cell(sh, co),
        "S", 1, ["C"], 3, lambda *_: None,
    )

    assert changes
    assert isinstance(ws["C9"].value, (int, float))
    assert abs(Evaluator(wb).cell("S", "C1")) <= 0.01


def test_last_resort_plug_withholds_embedded_formula_catchall():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "S"
    ws["A8"] = "Cash flow statement"; ws["A9"] = "Others"
    ws["B1"] = 0; ws["C1"] = "=10-C9"; ws["D1"] = "=10-D9"
    ws["B9"] = 0; ws["C9"] = 0; ws["D9"] = "=2+3"
    ws["C3"] = ws["D3"] = 1000
    messages = []
    last_resort_plug(wb, _writer(wb), lambda: lambda sh, co: Evaluator(wb).cell(sh,co),
                    "S", 1, ["C","D"], 3, messages.append)
    assert ws["C9"].value == 10
    assert ws["D9"].value == "=2+3"
    assert any("typed numeric input" in m for m in messages)

def test_orange_colour_does_not_authorize_plugging_a_pure_formula():
    from openpyxl.styles import PatternFill
    from pipeline.orchestrator import ObjectiveLoop
    from pipeline.ledger import Ledger
    from unittest.mock import patch
    wb=openpyxl.Workbook();ws=wb.active;ws.title='S'
    ws['B1']=0;ws['C1']='=10-C9';ws['C8']=2;ws['C9']='=C8'
    spec={'year_axis':{'S':{'columns':{'2024':'B','2025':'C'}}},'check_rows':[{'sheet':'S','row':1,'expect':0}],'key_rows':[]}
    writer=_writer(wb);writer.log['written'].append('S!C9')
    ws['C9'].fill=PatternFill('solid',fgColor='FFC000')
    loop=ObjectiveLoop(wb,spec,2025,Ledger(),[],{},writer,None)
    with patch.object(loop,'t_diagnose_balance',return_value='No evidence candidates'):
        answer=loop.t_plug_residual({'check':'S!1','into':'S!C9','why':'test'})
    assert ws['C9'].value=='=C8'
    assert 'not a numeric input' in answer

def test_residual_cannot_erase_mapped_claim_even_when_red_or_forced():
    from copy import deepcopy
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'S'
    ws['B9'] = 80; ws['C9'] = 100; ws['D9'] = 0
    writer = _writer(wb)
    writer.actual_cols = {'S': 'C'}
    writer.served = {('S', 9): {'value': 100, 'homed': True, 'home': ('S', 'C9'),
                               'conf': 3, 'flag': 'red', 'note': 'mapping needs definition review'}}
    writer.flag('S', 'C9', 'red', 'Review mapping scope')
    before = deepcopy(writer.served)
    assert not writer.write('S', 'C9', 90, kind='plug', trusted=True, force_lock=True)
    assert ws['C9'].value == 100 and writer.served == before
    assert writer.residual_refusal('S', 'D9') is None
    # Normal evidence correction remains possible; only unsupported residual edits are refused.
    assert writer.write('S', 'C9', 105, trusted=True)
    assert ('S', 9) not in writer.served


def test_terminal_residual_uses_same_home_contract():
    from pipeline.orchestrator import ObjectiveLoop
    from pipeline.ledger import Ledger
    from unittest.mock import patch
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'S'
    ws['B9'] = 80; ws['C9'] = 100; ws['C1'] = '=110-C9'
    spec = {'year_axis': {'S': {'columns': {'2024': 'B', '2025': 'C'}}},
            'check_rows': [{'sheet': 'S', 'row': 1, 'expect': 0}], 'key_rows': []}
    writer = _writer(wb)
    writer.served = {('S', 9): {'value': 100, 'homed': True, 'home': ('S', 'C9'),
                               'conf': 3, 'flag': 'red'}}
    loop = ObjectiveLoop(wb, spec, 2025, Ledger(), [], writer.served, writer, None)
    with patch.object(loop, 't_diagnose_balance', return_value='No evidence candidates'):
        result = loop.t_plug_residual({'check': 'S!1', 'into': 'S!C9', 'why': 'residual'})
    assert result.startswith('REFUSED:') and ws['C9'].value == 100
    assert Evaluator(wb).cell('S', 'C1') == 10


if __name__ == "__main__":
    import unittest
    suite=unittest.TestSuite(unittest.FunctionTestCase(f) for name,f in list(globals().items()) if name.startswith("test_") and callable(f))
    result=unittest.TextTestRunner().run(suite)
    raise SystemExit(not result.wasSuccessful())
