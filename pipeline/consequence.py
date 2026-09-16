"""THE CONSEQUENCE CARD (owner 2026-09-15: "if I move this number, will it
change my key numbers? if so how should I solve it? how do I make sure the
model meets the balancing requirement — these should all be thought by the
AI"). Goal-based, not scripted: code measures the objectives — the balance
checks in every year, the keys against the print — and when one is broken
it puts the consequence in front of the brain: what is off, which inputs
the run moved into it (the swing census, own flags first, each with its
evidence), and the ways to resolve it — take a mover back to what the
analyst had, absorb the residual in a mover as a traceable back-out, plug
the model's own residual row (the last resort), or leave it open as a
question for the analyst with the gap named. The brain picks; code applies
the pick through the writer's laws, re-measures, and takes a pick back that
made things worse. This replaces the bulk take-back of every serve and the
automatic plug ladder as the deciders; they remain the executors of a pick.
"""
import re
import time

from .checks import CHECK_TOL, forecast_columns, prior_column, year_columns
from .evaluator import Evaluator


def snapshot(loop):
    """What a trial may disturb: the write journal mark, the provenance
    (served), the locks, the sense rulings. Previews and take-backs restore
    all four (audit 2026-09-15: a preview un-served and unlocked proven
    movers; a taken-back pick left its ruling in the memory)."""
    w = loop.writer
    return (len(w.log.get("writes_all", [])), dict(loop.served or {}), set(getattr(w, "locked", ())),
            dict(w.log.get("rulings", {}) or {}), list(w.log.get("plugs", []) or []))


def restore(loop, snap):
    """Unwind every write since the snapshot through the writer (value AND
    look), then put back what the run believed: served, locks, rulings."""
    w = loop.writer
    mark, served0, locked0, rulings0, plugs0 = snap
    journal = w.log.get("style_journal", [])[mark:]
    for sh_w, co_w, old_w, _new in reversed(list(w.log.get("writes_all", []))[mark:]):
        style = next((j for j in journal if (j[0], j[1]) == (sh_w, co_w)), (sh_w, co_w, "", None, False))
        w.take_back(sh_w, co_w, old_w, style)
    if isinstance(loop.served, dict):
        loop.served.clear(); loop.served.update(served0)
    if hasattr(w, "locked"):
        w.locked.clear(); w.locked.update(locked0)
    if "rulings" in w.log:
        w.log["rulings"].clear(); w.log["rulings"].update(rulings0)
    if "plugs" in w.log:
        w.log["plugs"][:] = plugs0


def _fault(loop, text):
    """An objective the measure could not take is SAID, never swallowed
    (run 34935869107: the key objectives vanished from the ending without
    a word). Kept on the loop; run_ending logs each new one once."""
    faults = loop.__dict__.setdefault("objective_faults", [])
    if text not in faults:
        faults.append(text)


def broken_objectives(loop, keys_before, key_panel, panel_path, quiet=False):
    """-> [(kind, sheet, coord, amount, text)] largest first.
    kind: check (actual year) | forecast-check | key (rule 2 / off the print).
    quiet: a trial measure (plugs lifted) says nothing about cells it cannot evaluate."""
    wb, spec, ty = loop.wb, loop.spec, int(loop.ty)
    out = []
    ev = Evaluator(wb)
    say = (lambda t: None) if quiet else (lambda t: _fault(loop, t))
    for c in (spec.get("check_rows") or []):
        sheet, row = c.get("sheet"), int(c.get("row"))
        if sheet not in wb.sheetnames:
            continue
        expect = float(c.get("expect", 0))
        tcol = year_columns(spec, sheet).get(str(ty))
        cols = [(tcol, "check")] if tcol else []
        cols += [(fc, "forecast-check") for fc in (forecast_columns(spec, sheet, ty) or [])]
        for col, kind in cols:
            try:
                v = ev.cell(sheet, f"{col}{row}")
            except Exception as e:  # noqa: BLE001
                say(f"check {sheet}!{col}{row} cannot be evaluated: {e!r}")
                continue
            if isinstance(v, (int, float)) and abs(v - expect) > CHECK_TOL:
                out.append((kind, sheet, f"{col}{row}", float(v - expect),
                            f"balance check {sheet}!{col}{row} computes {v:,.1f} (should be {expect:,.0f})"))
    from .keytie import key_violations, key_state
    try:
        for nm, ref, then, now in key_violations(wb, spec, ty, loop.ledger, panel_path, keys_before, panel=key_panel):
            sh, co = ref.split("!", 1)
            amt = (now if isinstance(now, (int, float)) else 0.0) - then
            out.append(("key", sh, co, float(amt), f"key '{nm}' at {ref} was proven-printed {then:,.1f}, now {now if now is None else f'{now:,.1f}'} — printed nowhere"))
    except Exception as e:  # noqa: BLE001
        _fault(loop, f"key violations unavailable: {e!r}")
    try:
        # key_state speaks five fields (name, ref, value, print, tied); run
        # 34935869107 unpacked four, raised on every round, and the keys
        # were never an objective — the fault is now said, never swallowed.
        # A key the tie already JUDGED (confirmed by the bridge, or a
        # definition question for the analyst) is not re-raised every round.
        judged = set(loop.writer.log.get("key_verdicts", {}) or {})
        for nm, ref, got, want, ok in key_state(wb, spec, ty, panel_path, panel=key_panel):
            if ok or not isinstance(got, (int, float)) or nm in judged:
                continue
            sh, co = ref.split("!", 1)
            if any(o[1] == sh and o[2] == co for o in out):
                continue
            out.append(("key", sh, co, float(got - want), f"key '{nm}' at {ref} computes {got:,.1f} vs printed {want:,.1f}"))
    except Exception as e:  # noqa: BLE001
        _fault(loop, f"key objectives unavailable: {e!r}")
    try:
        out += sanity_breaks(loop)
    except Exception as e:  # noqa: BLE001
        _fault(loop, f"sanity objectives unavailable: {e!r}")
    out.sort(key=lambda o: -abs(o[3]))
    return out


def sanity_rows(loop):
    """The rows that must never go negative: cash (the report's own role) and
    total assets (a key row). -> {name: (sheet, row)}"""
    out = {}
    try:
        from .reportpage import resolve_rows
        rows, _primary = resolve_rows(loop.wb, loop.spec, int(loop.ty), getattr(loop, "period", None) or "FY")
        if rows.get("cash"):
            out["cash"] = (rows["cash"][0], int(rows["cash"][1]))
    except Exception as e:  # noqa: BLE001
        _fault(loop, f"cash row not resolved for the sanity objective: {e!r}")
    for kk in (loop.spec.get("key_rows") or []):
        nm = str(kk.get("name") or "").lower()
        if nm == "cash" and "cash" not in out:
            out["cash"] = (kk["sheet"], int(kk["row"]))
        if "total assets" in nm:
            out["total assets"] = (kk["sheet"], int(kk["row"]))
    return out


def _sanity_scan(loop):
    """Every cash / total-assets cell from the actual period on, in period order.
    -> [(name, sheet, coord, period_label, value)]"""
    out = []
    ev = Evaluator(loop.wb)
    ty = int(loop.ty)
    for nm, (sh, r) in sanity_rows(loop).items():
        if sh not in loop.wb.sheetnames:
            continue
        yc = year_columns(loop.spec, sh)
        cols = [(yc.get(str(ty)), str(ty))] + sorted(((c, y) for y, c in yc.items() if str(y).isdigit() and int(y) > ty), key=lambda x: int(x[1]))
        for col, year in cols:
            if not col:
                continue
            try:
                v = ev.cell(sh, f"{col}{r}")
            except Exception:
                continue
            if isinstance(v, (int, float)):
                out.append((nm, sh, f"{col}{r}", year, float(v)))
    return out


def sanity_breaks(loop, horizon=2):
    """Cash or total assets negative in the actual period or the next two
    periods (owner 2026-09-15: a normal roll-forward may go negative later
    on the analyst's own assumptions — that is flagged, not fixed).
    -> [("sanity", sheet, coord, value, text)]"""
    out = []
    seen = {}
    for nm, sh, coord, year, v in _sanity_scan(loop):
        k = seen.get(nm, 0)
        seen[nm] = k + 1
        if k > horizon:
            continue
        if v < -0.5:
            out.append(("sanity", sh, coord, v, f"{nm} goes negative in {year}: {v:,.0f} at {sh}!{coord}"))
    return out


def sanity_mass(loop, horizon=2):
    """The sanity objective's size, in the same money as the balance mass:
    how deep cash or total assets is under water in the actual period and
    the next two (run 35043265913: a 2028 balance plug drove 2026 cash to
    -966 and the objective measure never saw it, so the pick stood). An
    objective the measure ignores is an objective a pick may break for
    free."""
    try:
        return sum(abs(o[3]) for o in sanity_breaks(loop, horizon))
    except Exception as e:  # noqa: BLE001
        _fault(loop, f"sanity mass unavailable: {e!r}")
        return 0.0


def sanity_watch(loop, horizon=2):
    """Negatives beyond the horizon: for the analyst, never fixed. -> [text]"""
    out, seen = [], {}
    for nm, sh, coord, year, v in _sanity_scan(loop):
        k = seen.get(nm, 0)
        seen[nm] = k + 1
        if k > horizon and v < -0.5:
            out.append(f"{nm} goes negative in {year} ({v:,.0f} at {sh}!{coord}) — beyond the next two periods: your assumptions, your call")
    return out


def measure(loop, keys_before, key_panel, panel_path, lift_plugs=True):
    """One reading of the objectives — balance mass across every year, the
    keys off the print, cash/assets negatives — with the plug rows lifted for
    the balance and the keys (owner 2026-09-15: a plug masks a consequence).

    THE SANITY OBJECTIVE IS READ AS THE VERDICT READS IT — plugs live
    (reviewer 2026-09-16): check_mass, which takes the pick back, measures
    the model as it stands, so a preview taken with the plugs lifted showed
    the brain a cash break the verdict did not see, or hid one it did. One
    measurement for both. -> dict"""
    wb, writer = loop.wb, loop.writer
    lifted = []
    if lift_plugs:
        for ref in list(writer.log.get("plugs") or []):
            try:
                sh, co = ref.split("!", 1)
                if sh in wb.sheetnames:
                    lifted.append((sh, co, wb[sh][co].value))
                    wb[sh][co].value = 0
            except Exception:
                pass
    try:
        objs = broken_objectives(loop, keys_before, key_panel, panel_path, quiet=True)
        balance = sum(abs(o[3]) for o in objs if o[0] in ("check", "forecast-check"))
        keys_off = [(re.search(r"key '(.+?)' at ", o[4]).group(1) if re.search(r"key '(.+?)' at ", o[4]) else o[2]) for o in objs if o[0] == "key"]
    finally:
        for sh, co, v in lifted:
            wb[sh][co].value = v
    san = sanity_breaks(loop)              # the model as the verdict will measure it
    return {"balance": balance, "keys_off": keys_off, "sanity": san}


def describe(m0, m1):
    """The consequence of a pick in one phrase: balance, keys, cash/assets."""
    bits = [f"balance off {m0['balance']:,.0f} → {m1['balance']:,.0f}"]
    if m1["keys_off"] != m0["keys_off"]:
        bits.append("keys off: " + (", ".join(m1["keys_off"]) or "none"))
    elif m1["keys_off"]:
        bits.append(f"keys still off: {', '.join(m1['keys_off'])}")
    before = {(x[1], x[2]) for x in m0["sanity"]}
    new_san = [o for o in m1["sanity"] if (o[1], o[2]) not in before]
    if new_san:
        bits.append("✗ " + "; ".join(o[4] for o in new_san[:2]))
    elif m0["sanity"] and not m1["sanity"]:
        bits.append("cash/assets back to positive")
    elif not m1["sanity"]:
        bits.append("cash/assets positive")
    return " | ".join(bits)


def _plug_here(loop, sheet, coord, log):
    """The last resort, in the period that is broken: the terminal ladder
    for an actual-period check, the forecast residual row of THAT year for
    a forecast one (the ladder reads only the actual period's checks, so a
    forecast break answered 'plug' was left to whichever forecast year the
    repair suite plugged next)."""
    from .forecast_balance import plug_period
    if plug_period(loop, sheet, coord, log):
        return
    from .orchestrator import terminal_ladder
    terminal_ladder(loop, log)


def apply_pick(loop, pre_wb, pick, movers, amount, text):
    """Apply one consequence-card pick through the writer. -> True/False (written)."""
    wb, spec, ty, writer = loop.wb, loop.spec, int(loop.ty), loop.writer
    if pick.startswith("derive:"):
        if pick == "derive:via":
            from .investigate import named_ref
            why_txt = str(getattr(loop, "last_why", "") or "")
            refs = [f"{a}!{b}" for a, b in named_ref(why_txt, wb)]
            cell_ref, via_ref = (refs[0], refs[1]) if len(refs) >= 2 else ("", "")
        else:
            i = int(pick.split(":")[1]) - 1
            (sh, c), _share = movers[i]
            from .investigate import derivations
            d_ = derivations(loop, sh, c)[:1]
            cell_ref, via_ref = f"{sh}!{c}", (d_[0][0] if d_ else "")
        res = str(loop.t_derive({"cell": cell_ref, "via": via_ref, "why": f"consequence card for {text[:60]}"}))
        return res.startswith("DERIVED")
    i = int(pick.split(":")[1]) - 1
    (sh, c), _share = movers[i]
    r = int(re.sub(r"[A-Z]", "", c))
    pcol = prior_column(spec, sh, ty)
    if pick.startswith("revert:"):
        from openpyxl.utils import column_index_from_string as _ci
        pre_content = pre_wb[sh].cell(r, _ci(re.sub(r"\d", "", c))).value
        back = pre_content if pre_content is not None else 0.0
        ok = writer.write(sh, c, back, prior_coord=f"{pcol}{r}" if pcol else None, trusted=True, force_lock=True, flag="red",
                          note=f"Taken back by the brain's judgment: {text}; this input was moved by the run and is put back to what you had. Please check.")
        if ok:
            loop.served.pop((sh, r), None)
        return bool(ok)
    try:
        cur = Evaluator(wb).cell(sh, c)
    except Exception:
        cur = None
    held = wb[sh][c].value
    body = held[1:] if isinstance(held, str) and held.startswith("=") else (f"{cur:g}" if isinstance(cur, (int, float)) else None)
    if not body:
        return False
    ok_b = writer.write(sh, c, f"=({body})-({amount:.6g})", prior_coord=f"{pcol}{r}" if pcol else None, trusted=True, flag="orange",
                        note=f"Backed out by the brain's judgment so {text} closes: absorbed {amount:+,.1f} here. True up when disclosed.")
    if ok_b and isinstance(loop.served.get((sh, r)), dict):
        loop.served[(sh, r)] = dict(loop.served[(sh, r)], conf=3, flag="orange", note="backed out by the brain's judgment")
    return bool(ok_b)


def _colour(wb, writer, sheet, coord):
    ref = f"{sheet}!{coord}"
    from .writer import _fill_rgb
    rgb = _fill_rgb(wb[sheet][coord])
    if rgb == "FFC7CE":
        return "red"
    if rgb == "FFC000":
        return "orange"
    return "blue" if rgb == "BDD7EE" else ("red" if ref in writer.log.get("flags", []) and not rgb else "plain")


def _evidence(loop, sheet, coord):
    r = int(re.sub(r"[A-Z]", "", coord))
    e = (loop.served or {}).get((sheet, r))
    if isinstance(e, dict):
        return f"served from '{str(e.get('line') or '')[:36]}' p{e.get('page')} (conf {e.get('conf')})"
    note = None
    try:
        c = loop.wb[sheet][coord].comment
        note = c.text if c else None
    except Exception:
        pass
    return (str(note)[:60] if note else "no evidence recorded")


def build_card(loop, pre_wb, obj, movers, named, keys_before=None, key_panel=None, panel_path=None,
               tree_budget_s=None):
    kind, sheet, coord, amount, text = obj
    wb, spec, ty = loop.wb, loop.spec, int(loop.ty)
    ev = Evaluator(wb)
    lines = [f"CARD CONSEQUENCE {sheet}!{coord}", f"  OBJECTIVE BROKEN: {text} — off by {amount:+,.1f}"]
    if kind == "sanity":
        lines.append("  A negative cash or asset balance is usually caused by a wrong input elsewhere in the model, not by the last "
                     "pick — look at the movers below for the wrong one; if the model's own assumptions cause it, answer question.")
    if named:
        if any(str(lab).startswith("≈ ") for _v, lab, _w in named):
            lines.append("  the gap is ABOUT the size of a printed figure not yet in the model (not equal — judge whether this is where it belongs): "
                         + "; ".join(f"{v:,.0f} {lab} ({where})" for v, lab, where in named))
        else:
            lines.append("  the gap equals a printed figure not yet in the model: " + "; ".join(f"{v:,.0f} = '{lab}' ({where})" for v, lab, where in named))
    lines.append("  the inputs the run moved into this line, by their share of the move (your own flags first):")
    options = {}
    for i, ((sh, c), share) in enumerate(movers):
        r = int(re.sub(r"[A-Z]", "", c))
        label = str(wb[sh].cell(r, 1).value or "")[:30]
        try:
            now = ev.cell(sh, c)
        except Exception:
            now = None
        from .execreport import _pre_val
        from openpyxl.utils import column_index_from_string as _ci
        pre = _pre_val(pre_wb, sh, r, _ci(re.sub(r"\d", "", c)))
        col = _colour(wb, loop.writer, sh, c)
        from .rollover import input_is_proven as _iip
        proven_m = col != "red" and _iip(loop.served, sh, c, wb[sh][c].value, loop.writer.log.get("flags", []), wb)
        lines.append(f"    [{i + 1}] {sh}!{c} '{label}': now {now if now is None else f'{now:,.2f}'} (was {pre if pre is None else f'{pre:,.2f}'}), "
                     f"{col}{', PROVEN from the print (each figure tied) — not a place to absorb a gap' if proven_m else ''}, "
                     f"{_evidence(loop, sh, c)}, carries {abs(share) * 100:.0f}% of the move")
        try:
            from .workqueue import row_context_short, cell_story
            lines.append(f"          [{row_context_short(loop, sh, re.sub(r'[0-9]', '', c), r)}]")
            _st = cell_story(loop, sh, re.sub(r'[0-9]', '', c), r)
            if _st:
                lines.append("        " + _st)
        except Exception:
            pass
        options[f"revert:{i + 1}"] = f"put {sh}!{c} back to what the analyst had ({pre if pre is None else f'{pre:,.2f}'}) — red, taken back by your judgment"
        if proven_m:
            continue                      # a proven figure is never backed out or derived away
        options[f"backout:{i + 1}"] = f"absorb the residual in {sh}!{c} as a traceable formula — orange"
        try:
            from .investigate import derivations
            for u_ref, u_lab, implied, target, why in derivations(loop, sh, c)[:1]:
                options[f"derive:{i + 1}"] = f"set {sh}!{c} to {implied:,.2f}, the value that makes {u_ref} '{u_lab}' equal its proven {target:,.2f} ({why}) — orange with that proof"
        except Exception:
            pass
    # THE WHOLE INPUT TREE, NOT ONLY WHAT MOVED (run 34993405014: the two
    # cells that were short of the print — Final!AI14, put back to a hardcode
    # at the same value, and Final!AI16, never written — moved 0 against the
    # analyst's baseline, so the census could not show them and the brain was
    # asked to choose among PROVEN movers only; it answered 'question' twice).
    moved = {(sh, c) for (sh, c), _s in movers}
    tree = []
    # the walk spends what is LEFT of the round's own census budget, never a
    # second budget of its own (reviewer 2026-09-16: a card could cost the
    # census's 60 s plus 20 s again)
    _tb = 20.0 if tree_budget_s is None else float(tree_budget_s)
    try:
        from .investigate import input_leaves
        cols = {sh: year_columns(spec, sh).get(str(ty)) for sh in wb.sheetnames}
        tree = ([x for x in input_leaves(wb, sheet, coord, cols=cols, budget_s=_tb) if x not in moved]
                if _tb > 0 else [])
        if _tb <= 0:
            lines.append("  (the census used this round's whole budget — this line's other inputs are not listed)")
    except Exception as ex:  # noqa: BLE001
        lines.append(f"  (this line's input tree could not be walked: {type(ex).__name__}: {str(ex)[:80]})")
    if tree:
        lines.append("  every OTHER actual-period input this line is built from — these did not move against your "
                     "baseline, which is exactly what an input the run left short looks like (name one in 'why' to derive it):")
        from .execreport import _pre_val
        from openpyxl.utils import column_index_from_string as _ci2
        from .rollover import input_is_proven as _iip2
        for sh, c in tree:
            r2 = int(re.sub(r"[A-Z]", "", c))
            lab2 = str(wb[sh].cell(r2, 1).value or "")[:30]
            now2 = wb[sh][c].value
            pre2 = _pre_val(pre_wb, sh, r2, _ci2(re.sub(r"\d", "", c)))
            col2 = _colour(wb, loop.writer, sh, c)
            served2 = (loop.served or {}).get((sh, r2))
            same = isinstance(now2, (int, float)) and isinstance(pre2, (int, float)) and abs(now2 - pre2) < 1e-9
            if served2 and col2 != "red" and _iip2(loop.served, sh, c, now2, loop.writer.log.get("flags", []), wb):
                status = "PROVEN from the print"
            elif served2 and same:
                status = "written back to your own figure by the run — the update left it where it was"
            elif served2:
                status = "written by the run"
            elif same:
                status = "UNTOUCHED — still your own estimate, never marked to the print"
            else:
                status = "moved by the model's own formulas"
            lines.append(f"    {sh}!{c} '{lab2}': now {now2 if not isinstance(now2, (int, float)) else f'{now2:,.2f}'} "
                         f"(your estimate {pre2 if not isinstance(pre2, (int, float)) else f'{pre2:,.2f}'}), {col2}, {status}; "
                         f"{_evidence(loop, sh, c)}")
    if kind == "key":
        # A GAP SHARED BY SEVERAL KEYS SITS ABOVE THEM (r1212: the card for
        # recurring net profit never said reported net profit was off by the
        # same -2,315, and the brain derived the one-off line to close it)
        try:
            from .keytie import key_state
            same_gap = [(nm, ref, got - want) for nm, ref, got, want, ok in key_state(wb, spec, ty, panel_path, panel=key_panel)
                        if not ok and isinstance(got, (int, float)) and ref != f"{sheet}!{coord}"
                        and abs((got - want) - amount) <= max(0.5, abs(amount) * 0.02)]
            if same_gap:
                lines.append("  the SAME gap is open on: "
                             + "; ".join(f"{nm} at {ref} (off {d:+,.1f})" for nm, ref, d in same_gap)
                             + " — a gap that sits on several keys at once comes from a line they ALL consume, "
                               "not from anything under this key alone")
        except Exception as ex:  # noqa: BLE001
            lines.append(f"  (the other keys could not be measured: {type(ex).__name__}: {str(ex)[:80]})")
    options["derive:via"] = ("derive an input yourself — a mover above or one of the untouched inputs: say in 'why' which "
                             "input cell and via which proven formula cell (e.g. 'ROAFNA!AI71 via ROAFNA!AI64'); code solves and verifies")
    options["plug"] = "plug the model's own residual row — the last resort, orange, reported"
    options["question"] = "leave it open, red: a definition question for the analyst (say what)"
    # THE CONSEQUENCE OF EACH WAY (owner 2026-09-15: "the brain has to know
    # whether it will cause cash to go negative, assets to go negative,
    # whether the model will be unbalanced, or affect the key numbers"):
    # every way is tried, measured with the plug rows lifted, and taken back
    previews = {}
    if keys_before is not None:
        writer = loop.writer
        m0 = measure(loop, keys_before, key_panel, panel_path)
        for k in list(options):
            if not (k.startswith("revert:") or k.startswith("backout:") or (k.startswith("derive:") and k != "derive:via")):
                continue
            snap = snapshot(loop)
            try:
                ok = apply_pick(loop, pre_wb, k, movers, amount, text)
                previews[k] = describe(m0, measure(loop, keys_before, key_panel, panel_path)) if ok else "cannot be written"
            except Exception as ex:  # noqa: BLE001
                previews[k] = f"could not be tried ({type(ex).__name__}: {str(ex)[:60]})"
            finally:
                restore(loop, snap)
    lines.append("  the ways to resolve it, each with what it does to the objectives (plugs lifted while measuring):")
    for k, v in options.items():
        lines.append(f"    {k}: {v}" + (f"  → {previews[k]}" if k in previews else ""))
    lines.append("  Think like the analyst: which input is wrong, or is the model's definition different from the print? "
                 "A plug is only for a gap nothing explains. answers: " + ", ".join(options))
    return "\n".join(lines), options


def resolve_objectives(loop, pre_wb, log, ask, gate_once, repair_round, check_mass, keys_before, key_panel, panel_path,
                       deadline_s=600.0, max_rounds=10):
    """The goal-based ending. -> (ok, failures, card) from the last gate_once, or None when no brain answered."""
    from .investigate import swing_leaves
    from .keytie import name_gap
    wb, spec, ty, writer = loop.wb, loop.spec, int(loop.ty), loop.writer
    t0 = time.monotonic()
    asked = set()
    result = None
    answered_any = False
    for _round in range(max_rounds):
        if time.monotonic() - t0 > deadline_s:
            log("[consequence] clock: the objectives that remain are written up")
            break
        objs = [o for o in broken_objectives(loop, keys_before, key_panel, panel_path) if (o[1], o[2]) not in asked]
        if not objs:
            break
        obj = objs[0]
        kind, sheet, coord, amount, text = obj
        flagged = {tuple(ref.split("!")) for ref in writer.log.get("flags", [])}
        census = swing_leaves(wb, pre_wb, sheet, coord, flagged=flagged, budget_s=60)
        movers = census[:6]
        named, _rest = name_gap(loop.ledger, amount, served=loop.served) if kind == "key" else ([], abs(amount))
        card, options = build_card(loop, pre_wb, obj, movers, named)
        log(f"[queue] card CONSEQUENCE {sheet}!{coord}: " + " | ".join(ln.strip() for ln in card.splitlines()[1:4 + len(movers)])[:900])
        pick = ask(card, options, "__auto__")
        if pick not in options:
            if not answered_any:
                return None                       # no brain here: the deterministic floor path decides
            asked.add((sheet, coord))
            continue
        answered_any = True
        log(f"[queue] CONSEQUENCE {sheet}!{coord} -> {pick}")
        mass0 = check_mass()
        mark = len(writer.log.get("writes_all", []))
        applied = True
        if pick == "question":
            writer.flag_ref(f"{sheet}!{coord}", "red", f"OPEN, by the brain's judgment: {text} — a question for the analyst"
                            + ("; the gap equals " + "; ".join(f"{v:,.0f} '{lab}' ({where})" for v, lab, where in named) if named else ""))
            asked.add((sheet, coord))
            continue
        if pick == "plug":
            _plug_here(loop, sheet, coord, log)
        else:
            i = int(pick.split(":")[1]) - 1
            (sh, c), _share = movers[i]
            r = int(re.sub(r"[A-Z]", "", c))
            pcol = prior_column(spec, sh, ty)
            if pick.startswith("revert:"):
                from .execreport import _pre_val
                from openpyxl.utils import column_index_from_string as _ci
                pre_content = pre_wb[sh].cell(r, _ci(re.sub(r"\d", "", c))).value
                back = pre_content if pre_content is not None else 0.0
                applied = writer.write(sh, c, back, prior_coord=f"{pcol}{r}" if pcol else None, trusted=True, force_lock=True, flag="red",
                                       note=f"Taken back by the brain's judgment: {text}; this input was moved by the run and is put back to what you had. Please check.")
                if applied:
                    loop.served.pop((sh, r), None)
            else:
                try:
                    cur = Evaluator(wb).cell(sh, c)
                except Exception:
                    cur = None
                held = wb[sh][c].value
                body = held[1:] if isinstance(held, str) and held.startswith("=") else (f"{cur:g}" if isinstance(cur, (int, float)) else None)
                if body is None:
                    applied = False
                else:
                    applied = writer.write(sh, c, f"=({body})-({amount:.6g})", prior_coord=f"{pcol}{r}" if pcol else None, trusted=True, flag="orange",
                                           note=f"Backed out by the brain's judgment so {text} closes: absorbed {amount:+,.1f} here. True up when disclosed.")
        if not applied:
            log(f"[consequence] the pick {pick} could not be written (the writer's law refused) — asked again without it")
            asked.add((sheet, coord))
            continue
        repair_round("consequence")
        result = gate_once()
        mass1 = check_mass()
        if mass1 > mass0 + 1.0:
            journal = writer.log.get("style_journal", [])[mark:]
            for sh_w, co_w, old_w, _new in reversed(list(writer.log.get("writes_all", []))[mark:]):
                style = next((j for j in journal if (j[0], j[1]) == (sh_w, co_w)), (sh_w, co_w, "", None, False))
                writer.take_back(sh_w, co_w, old_w, style)
            repair_round("consequence take-back")
            result = gate_once()
            log(f"[consequence] {pick} made the objectives worse ({mass0:,.0f} -> {mass1:,.0f}); taken back")
            asked.add((sheet, coord))
        else:
            log(f"[consequence] {sheet}!{coord}: {pick} -> objectives mass {mass0:,.0f} -> {mass1:,.0f}")
    return result


def run_ending(loop, pre_wb, log, ask, gate_once, repair_round, check_mass, keys_before, key_panel, panel_path,
               deadline_s, hold_zero=None, brain=True, max_rounds=14):
    """THE ENDING (owner 2026-09-15: "a loop at the end that identifies the
    issues, fixes them, makes sure the key numbers and the balance sheet
    balance, and if there is another problem fixes it again — and stops
    itself"). One loop, nothing stacked:
      measure the objectives — the balance in every year, the keys against
      the print, the headline lines out of line on the rollover — take the
      biggest break, put it to the brain with the movers and the ways to
      resolve it, apply the pick through the writer's laws, re-run the
      repairs the model itself defines, measure again. EVERY round is
      verified the same way: a new failure or a worse total is taken back.
    Stops when the objectives hold, when every remaining break has been
    judged (a question for the analyst, or a pick that did not help), or at
    the clock. With no brain (a floor) the model's own executors act once
    per break, matched to its kind. -> (ok, failures, card) from the last gate."""
    from .investigate import swing_leaves
    from .keytie import name_gap
    from .sensecheck import headline_deltas, suspicious, investigate_line, reason_text, _sense_row
    wb, spec, ty, writer = loop.wb, loop.spec, int(loop.ty), loop.writer
    t0 = time.monotonic()
    asked = set()
    lines = writer.log.setdefault("ending", [])
    loop.census_budget_s = 90 if (brain and ask is not None) else 10
    if hold_zero:
        hold_zero()
    repair_round("first")
    result = gate_once()
    order = {"check": 0, "key": 1, "sanity": 2, "forecast-check": 3, "sense": 4}
    said_faults = set()

    def _say_faults():
        for f in loop.__dict__.get("objective_faults", []):
            if f not in said_faults:
                said_faults.add(f)
                lines.append(f"objective not measured: {f}")
                log(f"[ending] FAULT {f}")

    judged = set()       # (kind, sheet, row): a question, no answer, or a pick that moved nothing — the break is judged
    tried = {}           # (sheet, coord) -> the picks taken back; a break is dealt again without them while its last pick moved the mass

    def _inherited(res):
        out = set()
        card_ = res[2] if isinstance(res, tuple) and len(res) > 2 and isinstance(res[2], dict) else {}
        for ln in card_.get("inherited_breaks", []) or []:
            m_ = re.match(r"^(.*)!r(\d+) \((\d{4})\)", str(ln).split(": ")[0])
            if m_:
                col_ = year_columns(spec, m_.group(1)).get(m_.group(3))
                if col_:
                    out.add((m_.group(1), f"{col_}{m_.group(2)}"))
        return out

    def _breaks():
        raw = broken_objectives(loop, keys_before, key_panel, panel_path)
        inherited = _inherited(result)
        out, seen_rows = [], set()
        for o in raw:
            kind, sheet, coord, amount, text = o
            if (sheet, coord) in inherited:
                continue                    # the analyst's own pre-update break: reported by the gate, not worked here
            if kind == "forecast-check":
                row = re.sub(r"[A-Z]", "", coord)
                if (sheet, row) in seen_rows:
                    continue                # the forecast years of one check row are ONE objective
                seen_rows.add((sheet, row))
            out.append((order[kind], abs(amount), o))
        for d in suspicious(headline_deltas(wb, pre_wb, spec, ty)):
            sh1, c1 = d["ref1"].split("!")
            out.append((order["sense"], abs(d["d1"] - d["d0"]), ("sense", sh1, c1, d["d1"] - d["d0"], reason_text(d), d)))
        out.sort(key=lambda x: (x[0], -x[1]))
        return [o for _p, _a, o in out if (o[0], o[1], re.sub(r"[A-Z]", "", o[2])) not in judged]

    def _finish(res):
        """The closing measure, whatever the exit: the headline lines as they
        stand NOW (the report's Look-here reads the ending's rows, never a
        checkpoint verdict the ending has since overtaken), the negatives
        beyond the horizon as a watch, the count of what remains."""
        rows_ = [s_ for s_ in writer.log.get("sense_rows", []) if isinstance(s_, dict)]
        done = {s_.get("name") for s_ in rows_ if s_.get("stage") == "ending"}
        earlier = {}
        for s_ in rows_:                       # the checkpoint's trail (its leaf link, its verdict text) travels into the closing row
            if s_.get("stage") != "ending":
                earlier[s_.get("name")] = s_
        try:
            deltas = headline_deltas(wb, pre_wb, spec, ty)
            bad = {d_["name"] for d_ in suspicious(deltas)}
            for d_ in deltas:
                if d_.get("name") in done:
                    continue
                e_ = earlier.get(d_.get("name")) or {}
                leaf_ = tuple(str(e_["leaf"]).split("!", 1)) if e_.get("leaf") and "!" in str(e_["leaf"]) else None
                if d_["name"] in bad:
                    _sense_row(writer, d_, "red", (e_.get("text") or "") + " — still out of line after the ending; your ruling", leaf_, "ending")
                else:
                    _sense_row(writer, d_, e_.get("verdict") if e_.get("verdict") in ("fixed", "genuine", "unusual") else "inline",
                               e_.get("text") or "", leaf_, "ending")
        except Exception as e_:  # noqa: BLE001
            _fault(loop, f"closing sense rows not written: {e_!r}")
        try:
            for w in sanity_watch(loop):
                lines.append("WATCH " + w)
                log("[ending] watch: " + w)
                m_ = re.search(r"at ([^)]+!\S+)\)", w)
                if m_:
                    sh_w, co_w = m_.group(1).split("!", 1)
                    writer.watch(sh_w, co_w, "Sense check: " + w)     # a forecast cell is watch-listed, never painted
        except Exception as e_:  # noqa: BLE001
            _fault(loop, f"the watch was not written: {e_!r}")
        _say_faults()
        inh = _inherited(res)
        left = [o for o in broken_objectives(loop, keys_before, key_panel, panel_path) if (o[1], o[2]) not in inh]
        lines.append(f"ended: {len(left)} objective(s) still broken" if left else "ended: every objective holds")
        log(f"[ending] {lines[-1]}")
        return res

    if result[0] and not broken_objectives(loop, keys_before, key_panel, panel_path) \
            and not suspicious(headline_deltas(wb, pre_wb, spec, ty)):
        lines.append("objectives held at the first measure")
        return _finish(result)

    def _fails(res):
        """The broken objectives as a set — the same measure the loop works to, not the gate's strings."""
        return {(o[1], o[2]) for o in broken_objectives(loop, keys_before, key_panel, panel_path)}

    def _gaps():
        return {d["name"]: abs(d["d1"] - d["d0"]) for d in headline_deltas(wb, pre_wb, spec, ty)}

    def _amount_of(sheet, coord):
        for o in broken_objectives(loop, keys_before, key_panel, panel_path):
            if (o[1], o[2]) == (sheet, coord):
                return abs(o[3])
        return 0.0

    def _take_back(snap, why):
        restore(loop, snap)
        repair_round("ending take-back")
        log(f"[ending] {why}; taken back")
        return gate_once()

    unanswered_checks = []
    for rnd in range(max_rounds):
        elapsed = time.monotonic() - t0
        if elapsed > deadline_s:
            lines.append("the clock ended the loop; what remains is written up")
            log("[ending] clock: the objectives that remain are written up")
            break
        _say_faults()
        breaks = _breaks()
        if not breaks:
            break
        obj = breaks[0]
        kind, sheet, coord, amount = obj[0], obj[1], obj[2], obj[3]
        akey = (kind, sheet, re.sub(r"[A-Z]", "", coord))
        mass0 = check_mass()
        fails0 = _fails(result)
        gaps0 = _gaps()
        snap = snapshot(loop)
        left_s = deadline_s - elapsed
        log(f"[ending] round {rnd + 1}: {kind} {sheet}!{coord} off {amount:+,.2f} (objectives mass {mass0:,.0f}; {len(breaks)} break(s) open)")
        pick = "__auto__"
        sense_note = None
        try:
            if kind == "sense":
                d = obj[5]
                fails_pre = _fails(gate_once())

                def _rerun():
                    """Did THIS pick open a check? (the gate's overall verdict also
                    carries every break that was open before the pick)"""
                    repair_round("ending")
                    return not (_fails(gate_once()) - fails_pre)
                verdict, text, leaf = investigate_line(loop, pre_wb, d, log, rerun=_rerun)
                sense_note = (d, verdict, text, leaf)
                pick = verdict
            else:
                options, movers, named = {}, [], []
                if ask is not None:
                    flagged = {tuple(ref.split("!")) for ref in writer.log.get("flags", [])}
                    _cen_budget = max(10.0, min(60.0, left_s - 60.0))
                    _cen_t0 = time.monotonic()
                    movers = swing_leaves(wb, pre_wb, sheet, coord, flagged=flagged,
                                          budget_s=_cen_budget)[:6]
                    named, _rest = name_gap(loop.ledger, amount, served=loop.served, near=(kind != "key"))
                    # previews cost a measure per way; near the clock the card goes without them
                    card, options = build_card(loop, pre_wb, obj[:5], movers, named,
                                               keys_before if left_s > 120 else None, key_panel, panel_path,
                                               tree_budget_s=max(0.0, _cen_budget - (time.monotonic() - _cen_t0)))
                    for t_ in tried.get((sheet, coord), ()):
                        if t_ in options:
                            options.pop(t_)
                            card += f"\n  (already tried and taken back: {t_})"
                    log(f"[queue] card CONSEQUENCE {sheet}!{coord}: " + " | ".join(ln.strip() for ln in card.splitlines()[1:4 + len(movers)])[:900])
                    loop.last_ask_error = None
                    pick = ask(card, options, "__auto__")
                if pick not in options:
                    if brain:
                        # THE BRAIN WAS ASKED AND GAVE NOTHING (audit 2026-09-15: a dead
                        # or out-of-clock brain fell through to the automatic executors
                        # and placed the keys the tie had only proposed) — a key or a
                        # negative balance stays open, red, with the reason; a balance
                        # check waits for the last resort at the very end (a plug, orange)
                        why_ = getattr(loop, "last_ask_error", None) or f"answer '{pick}' is not one of the ways"
                        log(f"[ending] {kind} {sheet}!{coord}: no answer from the brain ({str(why_)[:120]}) — "
                            + ("left for the last resort" if kind in ("check", "forecast-check") else "left open, red"))
                        if kind in ("check", "forecast-check"):
                            unanswered_checks.append((sheet, coord))
                        else:
                            writer.flag_ref(f"{sheet}!{coord}", "red", f"OPEN: {obj[4]} — the brain gave no answer ({str(why_)[:160]})")
                        lines.append(f"NO ANSWER {sheet}!{coord}: {obj[4]} ({str(why_)[:120]})")
                        judged.add(akey)
                        continue
                    # no brain (a floor): the model's own executor for this kind of break, once
                    pick = "__auto__"
                    if kind == "check":
                        from .orchestrator import terminal_ladder
                        terminal_ladder(loop, log)
                    elif kind == "key":
                        from .keytie import key_tie as _kt
                        _kt(wb, spec, ty, writer, panel_path, log, ledger=loop.ledger, panel=key_panel, absorbers="any")
                    # a forecast-year check: the repairs (forecast plugs) re-solve below
                elif pick == "question":
                    log(f"[queue] CONSEQUENCE {sheet}!{coord} -> question")
                    writer.flag_ref(f"{sheet}!{coord}", "red", f"OPEN, by the brain's judgment: {obj[4]} — a question for the analyst"
                                    + ("; the gap equals " + "; ".join(f"{v:,.0f} '{lab}' ({where})" for v, lab, where in named) if named else ""))
                    lines.append(f"QUESTION {sheet}!{coord}: {obj[4]}")
                    judged.add(akey)
                    continue
                elif pick == "plug":
                    log(f"[queue] CONSEQUENCE {sheet}!{coord} -> plug")
                    _plug_here(loop, sheet, coord, log)
                else:
                    log(f"[queue] CONSEQUENCE {sheet}!{coord} -> {pick}")
                    applied = apply_pick(loop, pre_wb, pick, movers, amount, obj[4])
                    if not applied:
                        log(f"[ending] the pick {pick} could not be written (the writer's law refused)")
                        tried.setdefault((sheet, coord), []).append(pick)
                        if len(tried[(sheet, coord)]) >= len(options):
                            judged.add(akey)
                        continue
            # the same verification for every kind of pick
            if hold_zero:
                hold_zero()
            repair_round("ending")
            result = gate_once()
            mass1 = check_mass()
            objs1 = broken_objectives(loop, keys_before, key_panel, panel_path)
            # A pick is judged on the WHOLE objective mass (owner 2026-09-16: cash
            # counts), so a pick that closes more than the cash it opens stands and
            # the negative it leaves becomes the next objective (owner 2026-09-15:
            # a correct input can turn cash negative because another input is
            # wrong — the loop goes and finds that one). A cash break is not
            # counted TWICE: it moves the mass, never the opened-checks list.
            san1 = {(o[1], o[2]) for o in objs1 if o[0] == "sanity"}
            new_fails = {f for f in (_fails(result) - fails0) if f not in san1}
            new_san = [o for o in objs1 if o[0] == "sanity" and (o[1], o[2]) not in fails0]
            gaps1 = _gaps()
            from .sensecheck import SENSE_GAP
            # the objectives in the owner's order: balance first, keys second, the
            # swing lines third — a widened swing line vetoes a sense pick, never
            # a fix that closes the balance or a key (the loop returns to the line)
            widened = [nm for nm, g in gaps1.items() if g > gaps0.get(nm, 0.0) + SENSE_GAP and g > SENSE_GAP] if kind == "sense" else []
            # a fix that did not fix: a pick aimed at a check or key must close most of it
            left_amt = next((abs(o[3]) for o in objs1 if (o[1], o[2]) == (sheet, coord)), 0.0)
            half_done = kind in ("check", "forecast-check", "key") and pick != "__auto__" and left_amt > 0.5 * abs(amount)
            if mass1 > mass0 + CHECK_TOL or new_fails or widened or half_done:
                reason = ("did not close the break" if half_done and not (new_fails or widened or mass1 > mass0 + CHECK_TOL)
                          else f"made the objectives worse ({mass0:,.0f} -> {mass1:,.0f})")
                reason += (f"; opened {sorted(new_fails)[:2]}" if new_fails else "") + (f"; widened {widened[:2]}" if widened else "")
                result = _take_back(snap, f"{kind} {sheet}!{coord}: {pick} {reason}")
                tried.setdefault((sheet, coord), []).append(pick)
                if kind == "sense" or pick == "__auto__" or not [k_ for k_ in options if k_ not in tried[(sheet, coord)] and k_ not in ("question", "plug")]:
                    judged.add(akey)          # no untried way is left (a plug is never dealt as the retry)
                if sense_note is not None:
                    d, verdict, text, leaf = sense_note
                    _sense_row(writer, d, "red", f"the pick {verdict} was taken back ({reason}) — your ruling", leaf, "ending")
            else:
                log(f"[ending] {sheet}!{coord}: {pick} -> objectives mass {mass0:,.0f} -> {mass1:,.0f}")
                if new_san:
                    # the pick did not make the objectives worse overall and it
                    # stands; the negative it leaves is the next objective, not
                    # a veto (owner 2026-09-15)
                    log(f"[ending] after {pick} on {sheet}!{coord}: {new_san[0][4]} — kept; that is the next objective")
                    lines.append(f"after {pick} on {sheet}!{coord}: {new_san[0][4]} — investigated next")
                if sense_note is not None:
                    d, verdict, text, leaf = sense_note
                    _sense_row(writer, d, verdict, text, leaf, "ending")
                    lines.append(("RESOLVED " if verdict == "fixed" else "") + reason_text(d) + " | " + text)
                    judged.add(akey)
                elif mass1 >= mass0 - CHECK_TOL:
                    judged.add(akey)            # nothing moved: judged, not dealt again
        except Exception as e_round:  # noqa: BLE001
            # a fault mid-round never leaves a half-applied pick in the model
            _fault(loop, f"round {rnd + 1} on {kind} {sheet}!{coord} failed: {e_round!r}")
            result = _take_back(snap, f"{kind} {sheet}!{coord}: {pick} failed with {type(e_round).__name__}")
            judged.add(akey)
    if unanswered_checks:
        # THE LAST RESORT (owner 2026-09-15: "plugs only as the very last resort at the end"):
        # the brain answered nothing on these balance checks; the model's own plug
        # ladder closes what it can, orange, reported — after everything else
        from .orchestrator import terminal_ladder
        log(f"[ending] last resort: the brain answered nothing on {len(unanswered_checks)} check(s) — the model's own plug ladder, orange")
        lines.append(f"LAST RESORT: {len(unanswered_checks)} check(s) the brain did not answer went to the plug ladder")
        try:
            terminal_ladder(loop, log)
            if hold_zero:
                hold_zero()
            repair_round("ending last resort")
            result = gate_once()
        except Exception as e_lr:  # noqa: BLE001
            _fault(loop, f"the last resort failed: {e_lr!r}")
    return _finish(result if result is not None else gate_once())
