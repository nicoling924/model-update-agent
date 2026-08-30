"""The reclassification recipe (owner rulings 2026-08-30).

A reclassification changes the CUT, not the total. The law for a
segment block whose total is known but whose pieces no longer match:

1. MAP only segments untouched by the change — the disclosure line must
   match on BOTH the (normalised) name AND the prior-year figure. The
   join already enforces exactly this, so a reclassified segment simply
   arrives STALE.
2. Every stale segment is BACKED OUT at the overall total's growth:
   =prior * (total_new / total_old), an orange, traceable formula.
3. The residual lands in ONE plug row: the analyst's own designed
   residual when the block has one (structure is the analyst's — never
   add a second plug); otherwise the SMALLEST stale segment becomes
   =total - the other pieces.
4. A plug that comes out ugly (sign flip vs its prior, or moved more
   than 50%) escalates to RED — an ugly plug usually means a mapped
   segment is wrong, and the analyst must rule.

Blocks are discovered generically: a contiguous run of rows carrying
prior-year values, closed by a row labelled Total/合计 whose target-year
cell is alive (wired or served). Works for any sheet laid out this way.
"""
import re

from openpyxl.comments import Comment

_TOTAL = re.compile(r"^\s*(total|合计|總計|总计)\s*$", re.IGNORECASE)


def _label(ws, row, max_col=6):
    for c in range(1, max_col + 1):
        v = ws.cell(row=row, column=c).value
        if isinstance(v, str) and v.strip():
            return v.strip()
    return ""


def find_blocks(ws, prior_col, target_col, max_row=250):
    """[{rows: [r...], total_row: r}] — component rows + their total."""
    blocks, run = [], []
    for r in range(3, max_row + 1):
        pv = ws[f"{prior_col}{r}"].value
        lab = _label(ws, r)
        alive = pv not in (None, "")
        if alive and _TOTAL.match(lab or ""):
            tv = ws[f"{target_col}{r}"].value
            if run and tv not in (None, ""):
                blocks.append({"rows": list(run), "total_row": r})
            run = []
        elif alive:
            run.append(r)
        else:
            run = []
    return blocks


def designed_plug(ws, block, target_col):
    """The analyst's own residual row: a component whose target formula
    references the block's total cell."""
    tref = f"{target_col}{block['total_row']}"
    for r in block["rows"]:
        f = ws[f"{target_col}{r}"].value
        if isinstance(f, str) and f.startswith("=") \
                and re.search(rf"{target_col}\$?{block['total_row']}\b", f):
            return r
    return None


def reclass_sweep(wb, sheets, year_cols, prior_cols, writer, log,
                  evaluate=None):
    """Apply the recipe to every block holding stale (red-flagged, typed)
    segment inputs. Returns the number of cells backed out."""
    n = 0
    for sheet in sheets:
        tcol, pcol = year_cols.get(sheet), prior_cols.get(sheet)
        if not tcol or not pcol or sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        for block in find_blocks(ws, pcol, tcol):
            trow = block["total_row"]
            stale = [r for r in block["rows"]
                     if f"{sheet}!{tcol}{r}" in writer.log["flags"]
                     and isinstance(ws[f"{tcol}{r}"].value, (int, float))]
            if not stale:
                continue
            plug = designed_plug(ws, block, tcol)
            own_plug = None
            if plug is None:
                own_plug = min(
                    stale, key=lambda r: abs(ws[f"{pcol}{r}"].value or 0))
            for r in stale:
                if r == own_plug:
                    others = [k for k in block["rows"] if k != r]
                    f = "=%s%d-%s" % (tcol, trow,
                                      "-".join(f"{tcol}{k}" for k in others))
                    note = ("reclassified block plug: total less the other "
                            "segments (smallest segment carries the "
                            "residual) — true up from the segment note")
                else:
                    f = "=%s%d*%s$%d/%s$%d" % (pcol, r, tcol, trow,
                                               pcol, trow)
                    note = ("reclassified segment — held at the total's "
                            "growth rate; true up from the segment note")
                if writer.write(sheet, f"{tcol}{r}", f,
                                prior_coord=f"{pcol}{r}",
                                flag="orange", note=note):
                    n += 1
            # rule 4: an ugly plug escalates to red
            if own_plug is not None and evaluate is not None:
                try:
                    got = evaluate(sheet, f"{tcol}{own_plug}")
                    pv = ws[f"{pcol}{own_plug}"].value
                    if isinstance(got, (int, float)) \
                            and isinstance(pv, (int, float)) and pv:
                        if got * pv < 0 or abs(got / pv - 1) > 0.5:
                            cell = ws[f"{tcol}{own_plug}"]
                            cell.fill = writer.fills["red"]
                            cell.comment = Comment(
                                "PLUG LOOKS ABNORMAL (%.1f vs prior %.1f) "
                                "— an ugly plug usually means a mapped "
                                "segment is wrong. Analyst ruling needed."
                                % (got, pv), "Model Update Agent")
                except Exception:
                    pass
            log(f"[run] reclass sweep {sheet} rows {stale} -> total-growth "
                f"back-out (plug row {plug or own_plug})")
    return n


def flag_embedded_hardcodes(wb, sheets, year_cols, writer, log):
    """Studio-kernel parity (mindmap: key drivers): a target-year FORMULA
    carrying a numeric constant >= 1 smuggles last year's number into
    this year invisibly (Driver wind: =16602.97-J11). Flag red + note;
    never rewrite the analyst's formula."""
    n = 0
    for sheet in sheets:
        tcol = year_cols.get(sheet)
        if not tcol or sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        for r in range(3, min(ws.max_row, 250) + 1):
            f = ws[f"{tcol}{r}"].value
            if not (isinstance(f, str) and f.startswith("=")):
                continue
            bare = re.sub(r"'[^']*'!|[A-Za-z_][A-Za-z0-9_]*!|"
                          r"\$?[A-Z]{1,3}\$?[0-9]{1,5}", "", f)
            # a smuggled prior is a PRECISE figure (16602.97), not a
            # convention constant (365 days, 12 months, 100, powers of
            # ten for unit conversion)
            def _smuggled(x):
                v = float(x)
                if "." in x and v >= 1.0:
                    return True
                return (v >= 500 and v not in (1000.0, 10000.0, 100000.0,
                                               1000000.0, 8760.0))
            consts = [float(x) for x in
                      re.findall(r"(?<![\w.])(\d+(?:\.\d+)?)", bare)
                      if _smuggled(x)]
            ref = f"{sheet}!{tcol}{r}"
            if consts and ref not in writer.log["flags"]:
                cell = ws[f"{tcol}{r}"]
                cell.fill = writer.fills["red"]
                cell.comment = Comment(
                    "EMBEDDED HARDCODE (key driver): this formula carries "
                    "the constant(s) %s from a prior period — confirm they "
                    "still hold for the new period."
                    % ", ".join(f"{c:g}" for c in consts[:3]),
                    "Model Update Agent")
                writer.log["flags"].append(ref)
                n += 1
    if n:
        log(f"[run] {n} embedded hardcodes flagged as key drivers (red)")
    return n


# ---- the pure planner (restored — an overwrite clobbered it;
# the museum caught the loss). List-based law, no workbook: the
# agent/closer uses it wherever the sweep's block discovery does
# not apply. ----
WILD_MOVE = 0.5      # plug vs its own prior: beyond this is a red flag


def plan_backout(segments, total_actual, total_prior):
    """segments: [{"name": str, "prior": float, "actual": float|None}]
    — "actual" is the disclosed figure for segments the agent MAPPED
    (name + prior both matched); None for everything reclassified.

    Returns {"rows": [...], "growth": g, "plugName": str} where each row
    is {"name", "kind": mapped|growth|plug, "value", "formula", "flag",
    "note"}. Formulas are Excel-style over the OTHER rows by name — the
    caller substitutes real cell refs.
    """
    if not isinstance(total_actual, (int, float)) or not segments:
        raise ValueError("total_actual and segments are required")
    growth = (total_actual / total_prior
              if isinstance(total_prior, (int, float)) and total_prior
              else None)
    unmapped = [s for s in segments if s.get("actual") is None]
    rows = []
    if growth is None and unmapped:
        raise ValueError("total_prior is required to grow the back-outs")
    plug = (min(unmapped, key=lambda s: abs(s.get("prior") or 0))
            if unmapped else None)
    for s in segments:
        if s.get("actual") is not None:
            rows.append({"name": s["name"], "kind": "mapped",
                         "value": float(s["actual"]), "formula": "",
                         "flag": "", "note": ""})
            continue
        if s is plug:
            continue                      # placed last, needs the others
        val = round(float(s["prior"]) * growth, 1)
        rows.append({"name": s["name"], "kind": "growth", "value": val,
                     "formula": "=%s*%s" % (s["prior"], round(growth, 6)),
                     "flag": "orange",
                     "note": "held at the total's growth (%+.1f%%), "
                             "awaiting true-up" % ((growth - 1) * 100)})
    if plug is not None:
        others = sum(r["value"] for r in rows)
        val = round(float(total_actual) - others, 1)
        flag, note = "orange", "= total − other segments (the plug)"
        prior = plug.get("prior")
        if val < 0 or (isinstance(prior, (int, float)) and prior and
                       abs(val / prior - 1) > WILD_MOVE):
            flag = "red"
            note = ("plug = total − other segments came out %s vs prior "
                    "%.1f — check the MAPPED segments before trusting "
                    "any of this split" % (round(val, 1), prior))
        rows.append({"name": plug["name"], "kind": "plug", "value": val,
                     "formula": "=total-(%s)" % "+".join(
                         r["name"] for r in rows),
                     "flag": flag, "note": note})
    ssum = round(sum(r["value"] for r in rows), 1)
    if abs(ssum - total_actual) > 0.5:
        raise AssertionError("backout does not tie: %s vs %s"
                             % (ssum, total_actual))
    return {"rows": rows, "growth": growth,
            "plugName": plug["name"] if plug else ""}
