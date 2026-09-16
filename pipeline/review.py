"""THE REVIEW (owner 2026-09-16: "the card is an issue — there is no
thinking in preparing the cards, so the brain thinks from wrong choices or
choices without enough context"). The brain reviews the MODEL, not a card:
code measures the objectives, lays out the headline lines and every cell
the run wrote with its move against that cell's own history and its
evidence, and answers the tools the brain calls. The brain decides what to
look at and what to change; code applies, verifies the evidence, measures
the consequence and logs the turn. The only decision code makes is the very
last one: at the clock, any actual-year check still off goes to the
terminal ladder, orange, so the model is delivered balanced.
"""
import json
import re
import time

from .checks import CHECK_TOL, forecast_columns, prior_column, year_columns
from .consequence import _fault, _plug_here, restore, sanity_rows, snapshot
from .evaluator import Evaluator

MANDATE = """You are the equity analyst reviewing your own model update before delivery.

The model must (1) balance every year, (2) tie every key number to the print, (3) roll forward
sanely — no negative cash or assets, no headline line moving out of its history without a printed
reason. A wrong number is a failure; so is a number you could have fixed and left.

When something is off, find the INPUT that caused it — look at what moved against its history
first — and fix it with evidence. Compensating errors come in pairs; name both and set them
together. Say `done` only when every objective holds or you have named why it cannot.

You answer in JSON only: {"thinking": "...", "calls": [ ... ]}. Any number of calls per turn,
executed in the order you write them; their answers come back in the next turn. The calls:
  {"tool":"show","ref":"Sheet!AI16"}                  the row's inputs, what uses it, its history, printed lines on file
  {"tool":"find","q":"46.3"}  or {"q":"Fuel Cost"}    printed lines carrying that number or naming that label, with page
  {"tool":"try","sets":[{"ref":"Sheet!AI16","value":46.3}]}    preview on a snapshot: the objectives before and after, then restored
  {"tool":"set","sets":[{"ref":"Sheet!AI16","value":46.3}],"because":"p243 'Fuel Cost Adjustment 46.3 46.3'"}
  {"tool":"restore","ref":"Sheet!AI16","why":"..."}   put back what the analyst had, red
  {"tool":"plug","check":"Sheet!AI99","why":"..."}    the model's own residual site takes the gap; orange, reported
  {"tool":"done","objectives":{"balance":"holds | cannot be closed because ...","keys":"...","rollforward":"..."}}
A `sets` list with two entries is applied and measured as ONE change — that is how a compensating
pair is resolved. A `set` lands plain only when its evidence ties (the quote is on the page and the
comparative ties the model's prior), or when `because` states arithmetic code can verify; otherwise
it lands red for the analyst — it is never refused for want of evidence."""


# ── reading the model ────────────────────────────────────────────────────

def _num(x):
    try:
        return round(float(x), 2)
    except (TypeError, ValueError):
        return None


def _fmt(x):
    if x is None:
        return ""
    return f"{x:,.2f}" if abs(x) < 100 else f"{x:,.0f}"


def _label(ws, r):
    for c in ("A", "B", "C", "D"):
        v = ws[f"{c}{r}"].value
        if isinstance(v, str) and v.strip():
            return v.strip()[:40]
    return f"row {r}"


def _row_of(coord):
    return int(re.sub(r"[A-Z$]", "", coord))


def _col_of(coord):
    return re.sub(r"[0-9$]", "", coord)


def _act_col(loop, sheet):
    return year_columns(loop.spec, sheet).get(str(int(loop.ty)))


def _hist_cols(loop, sheet, n=4):
    yc = year_columns(loop.spec, sheet)
    yrs = sorted(int(y) for y in yc if str(y).isdigit() and int(y) < int(loop.ty))
    return [yc[str(y)] for y in yrs[-n:]]


def _fill(cell):
    try:
        rgb = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None
    except Exception:  # noqa: BLE001
        rgb = None
    return {"FFFFC7CE": "red", "FFFFC000": "orange"}.get(rgb or "", "plain")


def _value_index(loop):
    """{rounded model-units value: [(item, scale)]} over the ledger, built once —
    the comparative tie of 260 written cells costs one pass, not 260."""
    idx = loop.__dict__.get("_review_vindex")
    if idx is not None:
        return idx
    from .writegate import _SCALES, _nums, _sourceable
    idx = {}
    for it in getattr(loop.ledger, "items", []) or []:
        if not _sourceable(it):
            continue
        for n in _nums(it):
            if not isinstance(n, (int, float)) or n == 0:
                continue
            for f in _SCALES:
                v = abs(n) / f
                idx.setdefault(round(v, 1), []).append((it, f))
                idx.setdefault(float(round(v)), []).append((it, f))   # statement rounding: 396.2 IS the printed 396
    loop.__dict__["_review_vindex"] = idx
    return idx


def _hits(loop, value):
    """Printed rows carrying this value at the tie tolerance the write gate
    uses — a lookup that only matched to 0.1 called the rounded print a miss."""
    if not isinstance(value, (int, float)):
        return []
    idx, seen, out = _value_index(loop), set(), []
    for k in (round(abs(value), 1), float(round(abs(value)))):
        for it, f in idx.get(k, ()):
            if (id(it), f) not in seen:
                seen.add((id(it), f))
                out.append((it, f))
    return out


def _quote(it):
    return " ".join(str(getattr(it, "source_line", "") or getattr(it, "label", "")).split())


def _evidence_line(loop, sheet, coord, value, prior):
    """THE PRINTED LINE, NOT THE CELL'S NOTE (owner 2026-09-16: a note is what
    the run said about itself). -> the quoted line with its page and whether the
    comparative ties the model's prior, or 'no printed line'. A TYING row wins
    over the first row that merely carries the number."""
    from .writegate import ties_prior
    first = None
    for it, scale in _hits(loop, value)[:12]:
        tie = ties_prior(it, scale, prior, value) if isinstance(prior, (int, float)) else False
        if tie:
            return f"p{getattr(it, 'page', '?')} '{_quote(it)[:90]}' — comparative ties the model's prior {_fmt(prior)}"
        first = first or it
    if first is not None:
        return (f"p{getattr(first, 'page', '?')} '{_quote(first)[:90]}' — no row carrying this value has a "
                f"comparative that ties the model's prior {_fmt(prior) if prior is not None else '(none)'}")
    e = (loop.served or {}).get((sheet, _row_of(coord)))
    if isinstance(e, dict) and e.get("line"):
        return f"p{e.get('page')} '{str(e.get('line'))[:70]}' — the value is not printed on any page on file"
    return "no printed line"


def _metrics(loop, key_panel, panel_path):
    """ONE READING of every objective, as numbers: {ref: (what, value, kind, off)}
    — the balance checks in every year, each key against its print, cash and
    total assets from the actual period on, the headline lines against the
    analyst's own book. `off` is how far this objective is from where it must
    be (0 = it holds). What the context prints, what `try` diffs, what `done`
    is held to and what the last resort reads — the card era measured the same
    objectives twice, in two places, and the two disagreed."""
    wb, spec, ty = loop.wb, loop.spec, int(loop.ty)
    ev = Evaluator(wb)
    out = {}
    for c in (spec.get("check_rows") or []):
        sheet, row = c.get("sheet"), int(c.get("row"))
        if sheet not in wb.sheetnames:
            continue
        expect = float(c.get("expect", 0))
        act = _act_col(loop, sheet)
        for year, col in sorted(year_columns(spec, sheet).items(), key=lambda kv: str(kv[0])):
            if str(year).isdigit() and int(year) < ty:
                continue                      # the years before the actual are the analyst's own history
            try:
                v = ev.cell(sheet, f"{col}{row}")
            except Exception as e:  # noqa: BLE001
                _fault(loop, f"check {sheet}!{col}{row} cannot be evaluated: {e!r}")
                continue
            if isinstance(v, (int, float)):
                out[f"{sheet}!{col}{row}"] = (f"balance check {year} (must be {expect:,.0f})", float(v),
                                              "check" if col == act else "forecast-check", float(v) - expect)
    try:
        from .keytie import key_state
        judged = set(loop.writer.log.get("key_verdicts", {}) or {})
        for nm, ref, got, want, ok in key_state(wb, spec, ty, panel_path, panel=key_panel):
            if not isinstance(got, (int, float)):
                continue
            off = (float(got) - want) if isinstance(want, (int, float)) and not ok and nm not in judged else 0.0
            out[ref] = (f"key '{nm}' vs print {_fmt(_num(want))}" + ("" if ok else "  ← OFF THE PRINT"), float(got), "key", off)
    except Exception as e:  # noqa: BLE001
        _fault(loop, f"key objectives unavailable: {e!r}")
    try:
        for nm, (sh, r) in sanity_rows(loop).items():
            if sh not in wb.sheetnames:
                continue
            yc = year_columns(spec, sh)
            cols = [(yc.get(str(ty)), str(ty))] + sorted(((c, y) for y, c in yc.items() if str(y).isdigit() and int(y) > ty),
                                                         key=lambda x: int(x[1]))
            for i, (col, year) in enumerate(cols):
                if not col:
                    continue
                try:
                    v = ev.cell(sh, f"{col}{r}")
                except Exception:  # noqa: BLE001
                    continue
                if not isinstance(v, (int, float)):
                    continue
                # a normal roll-forward may go negative later on the analyst's own
                # assumptions (owner 2026-09-15): past the horizon it is a watch
                kind = "sanity" if i <= 2 else "watch"
                out[f"{sh}!{col}{r}"] = (f"{nm} {year} (must not be negative)", float(v), kind,
                                         float(v) if (v < -0.5 and kind == "sanity") else 0.0)
    except Exception as e:  # noqa: BLE001
        _fault(loop, f"sanity objectives unavailable: {e!r}")
    pre = loop.__dict__.get("_review_pre")
    try:
        from .sensecheck import headline_deltas, reason_text, suspicious
        for d in (suspicious(headline_deltas(wb, pre, spec, ty)) if pre is not None else ()):
            sh, co = str(d["ref1"]).split("!", 1)
            out[f"{sh}!{co}"] = (f"roll-forward: {reason_text(d)[:80]}  ← OUT OF LINE",
                                 float(d["d1"]), "sense", float(d["d1"] - d["d0"]))
    except Exception as e:  # noqa: BLE001
        _fault(loop, f"the roll-forward objective could not be measured: {e!r}")
    return out


def _key_panel_of(loop):
    return getattr(loop, "key_panel", None)


def _broken(loop, keys_before, key_panel, panel_path):
    """The objectives that are OFF, from that same reading, largest first.
    -> [(kind, sheet, coord, off, text)]"""
    out = []
    for ref, (what, v, kind, off) in _metrics(loop, key_panel, panel_path).items():
        if abs(off) > CHECK_TOL:
            sh, co = ref.split("!", 1)
            out.append((kind, sh, co, float(off), f"{what} at {ref} reads {_fmt(v)}"))
    try:
        # RULE 2: a key this run PROVED from the print and then moved off it is
        # broken even when the panel has no print to compare against
        from .keytie import key_violations
        seen = {(o[1], o[2]) for o in out}
        for nm, ref, then, now in key_violations(loop.wb, loop.spec, int(loop.ty), loop.ledger, panel_path,
                                                 keys_before, panel=key_panel):
            sh, co = ref.split("!", 1)
            if (sh, co) not in seen:
                out.append(("key", sh, co, float((now if isinstance(now, (int, float)) else 0.0) - then),
                            f"key '{nm}' at {ref} was proven-printed {then:,.1f}, now "
                            f"{now if now is None else f'{now:,.1f}'} — printed nowhere"))
    except Exception as e:  # noqa: BLE001
        _fault(loop, f"key violations unavailable: {e!r}")
    out.sort(key=lambda o: -abs(o[3]))
    return out


def _metrics_text(m):
    return [f"  {ref:<20} {what:<46} {_fmt(v):>14}" for ref, (what, v, _k, _o) in m.items()]


def _metric_diff(m0, m1):
    """What a change did to the objectives, line by line — only what moved."""
    out = []
    for ref, (what, v1, _k, _o) in m1.items():
        v0 = m0.get(ref, (None, None))[1]
        if v0 is None or abs(v1 - v0) > CHECK_TOL:
            out.append(f"    {ref} {what}: {_fmt(v0)} → {_fmt(v1)}")
    for ref in m0:
        if ref not in m1:
            out.append(f"    {ref} {m0[ref][0]}: {_fmt(m0[ref][1])} → (no longer measurable)")
    return out or ["    nothing moved"]


def _headline_rows(loop):
    """The report table's own P&L / BS / CF roles — the model's structure, not a list."""
    try:
        from .reportpage import resolve_rows
        rows, _primary = resolve_rows(loop.wb, loop.spec, int(loop.ty), getattr(loop, "period", None) or "FY")
        return [(role, v[0], int(v[1])) for role, v in rows.items() if v and v[0] in loop.wb.sheetnames]
    except Exception as e:  # noqa: BLE001
        _fault(loop, f"headline lines not resolved: {e!r}")
        return []


def build_context(loop, pre_wb, key_panel, panel_path, notes=(), answers=(), cap=260, trace_budget_s=40.0):
    """One reading of the model for the brain: the objectives measured, the
    headline lines against the analyst's own book, every written actual-column
    cell with its move against its history and its printed evidence, and the
    trace of the forecast lines that moved."""
    wb, spec, ty = loop.wb, loop.spec, int(loop.ty)
    loop.__dict__["_review_pre"] = pre_wb        # the analyst's own book: the roll-forward objective is measured against it
    ev, ev0 = Evaluator(wb), Evaluator(pre_wb)
    L = ["## 1. OBJECTIVES, measured by code"]
    L += _metrics_text(_metrics(loop, key_panel, panel_path))
    L.append("")
    L.append("## 2. HEADLINE LINES (history | the analyst's pre-update estimate for the actual year | now | next year before → now)")
    for role, sh, r in _headline_rows(loop):
        ac, fcs = _act_col(loop, sh), forecast_columns(spec, sh, ty) or []
        if not ac:
            continue
        hist = " | ".join(_fmt(_num(ev0.cell(sh, f"{c}{r}"))) for c in _hist_cols(loop, sh))
        est, now = _fmt(_num(ev0.cell(sh, f"{ac}{r}"))), _fmt(_num(ev.cell(sh, f"{ac}{r}")))
        nx = (f" | next {_fmt(_num(ev0.cell(sh, f'{fcs[0]}{r}')))} → {_fmt(_num(ev.cell(sh, f'{fcs[0]}{r}')))}") if fcs else ""
        L.append(f"  {sh}!{ac}{r:<5} {role:26} {hist} | est {est} | now {now}{nx}")
    L.append("")
    L.append("## 3. EVERY CELL THE RUN WROTE IN THE ACTUAL COLUMN (largest move against the cell's own history first; "
             "a cell with no history to judge against is listed last)")
    first = {}
    for sh_w, co_w, old_w, _new in (loop.writer.log.get("writes_all", []) or []):
        if sh_w in wb.sheetnames and _col_of(co_w) == _act_col(loop, sh_w) and (sh_w, co_w) not in first:
            first[(sh_w, co_w)] = old_w
    rows = []
    for (sh, co), old in first.items():
        r = _row_of(co)
        try:
            new = _num(ev.cell(sh, co))
        except Exception as e:  # noqa: BLE001
            L.append(f"  {sh}!{co} — WRITTEN BY THE RUN AND WILL NOT EVALUATE: {type(e).__name__}: {str(e)[:70]}")
            continue
        if new is None:
            L.append(f"  {sh}!{co} '{_label(wb[sh], r)}' — written by the run, reads as {wb[sh][co].value!r} (not a number)")
            continue
        hist = [h for h in (_num(ev0.cell(sh, f"{c}{r}")) for c in _hist_cols(loop, sh)) if h is not None]
        lo, hi = (min(hist), max(hist)) if hist else (None, None)
        if hist and hi - lo > 0:
            out = max(lo - new, new - hi, 0) / max(abs(hi), abs(lo), 1)
        elif hist:
            out = abs(new - hi) / max(abs(hi), 1)
        else:
            out = -1.0
        pcol = prior_column(spec, sh, ty)
        prior = _num(ev0.cell(sh, f"{pcol}{r}")) if pcol else None
        rows.append((out, sh, co, _label(wb[sh], r), _num(old), new, hist, _fill(wb[sh][co]),
                     _evidence_line(loop, sh, co, new, prior)))
    rows.sort(key=lambda t: -t[0])
    for _o, sh, co, lab, old, new, hist, flag, evid in rows[:cap]:
        L.append(f"  {sh}!{co:<6} {lab:40} was {_fmt(old):>11} → now {_fmt(new):>11} | "
                 f"history {', '.join(_fmt(h) for h in hist):32} | {flag:6} | {evid}")
    L.append(f"  ({len(rows)} written cells in total{'; the rest are within their history' if len(rows) > cap else ''})")
    L.append("")
    L.append("## 4. TRACE — the headline lines whose next-year forecast moved more than 10% against the analyst's book, "
             "and the typed inputs that carry the move")
    # the trace is the expensive half of the context: it is walked once per
    # state of the model (the write journal's mark), not once per turn, and the
    # whole section shares ONE budget — a turn is a turn, not a census
    mark, cached = loop.__dict__.get("_review_trace", (None, None))
    if mark == len(loop.writer.log.get("writes_all", []) or []):
        L += cached
    else:
        trace, t0 = [], time.monotonic()
        try:
            from .investigate import swing_leaves
            for role, sh, r in _headline_rows(loop):
                fcs = forecast_columns(spec, sh, ty) or []
                left = trace_budget_s - (time.monotonic() - t0)
                if not fcs or left <= 0:
                    if left <= 0:
                        trace.append("  (the trace used its budget — the lines below it are not walked)")
                        break
                    continue
                f0, f1 = _num(ev0.cell(sh, f"{fcs[0]}{r}")), _num(ev.cell(sh, f"{fcs[0]}{r}"))
                if f0 is None or f1 is None or abs(f0) < 1 or abs(f1 - f0) / abs(f0) < 0.10:
                    continue
                leaves = swing_leaves(wb, pre_wb, sh, f"{fcs[0]}{r}", budget_s=left)[:6]
                parts = [f"{s2}!{c2} '{_label(wb[s2], _row_of(c2))}' {_fmt(_num(ev0.cell(s2, c2)))} → {_fmt(_num(ev.cell(s2, c2)))}"
                         for (s2, c2), _share in leaves]
                trace.append(f"  {sh}!{fcs[0]}{r} {role} {_fmt(f0)} → {_fmt(f1)}: " + "; ".join(parts))
        except Exception as e:  # noqa: BLE001
            trace.append(f"  (the trace could not be walked: {type(e).__name__}: {str(e)[:80]})")
        loop.__dict__["_review_trace"] = (len(loop.writer.log.get("writes_all", []) or []), trace)
        L += trace
    if notes:
        L.append("")
        L.append("## 5. WHAT THE EARLY LOOK FOUND (the checkpoint sense check, before the writes)")
        L += [f"  {n}" for n in list(notes)[:30]]
    if answers:
        L.append("")
        L.append("## 6. YOUR LAST TURN")
        L += list(answers)
    return "\n".join(L)


# ── the tools ────────────────────────────────────────────────────────────

def _parse_ref(loop, ref):
    sh, _, co = str(ref).partition("!")
    if sh not in loop.wb.sheetnames:
        return None, None, f"no sheet '{sh}' in this workbook"
    co = co.replace("$", "").upper()
    if not re.fullmatch(r"[A-Z]{1,3}\d+", co):
        # 'Sheet!16' — the row alone, in the actual column
        if co.isdigit() and _act_col(loop, sh):
            co = f"{_act_col(loop, sh)}{co}"
        else:
            return None, None, f"'{ref}' is not a cell"
    return sh, co, None


def t_show(loop, pre_wb, ref):
    sh, co, err = _parse_ref(loop, ref)
    if err:
        return [f"show {ref}: {err}"]
    r, wb = _row_of(co), loop.wb
    ev, ev0 = Evaluator(wb), Evaluator(pre_wb)
    out = [f"show {sh}!{co} '{_label(wb[sh], r)}'"]
    out.append("    history " + ", ".join(f"{c}={_fmt(_num(ev0.cell(sh, f'{c}{r}')))}" for c in _hist_cols(loop, sh))
               + f" | analyst's estimate {_fmt(_num(ev0.cell(sh, co)))} | now {_fmt(_num(ev.cell(sh, co)))}")
    held = wb[sh][co].value
    out.append(f"    the cell holds: {str(held)[:120]}")
    try:
        from .investigate import _refs
        if isinstance(held, str) and held.startswith("="):
            ins = _refs(held, sh, wb)[:12]
            out.append("    its inputs: " + "; ".join(f"{a}!{b} '{_label(wb[a], _row_of(b))}' = {_fmt(_num(ev.cell(a, b)))}"
                                                      for a, b in ins))
    except Exception as e:  # noqa: BLE001
        out.append(f"    (its inputs could not be read: {e!r})")
    users, t0 = [], time.monotonic()
    for s2 in wb.sheetnames:
        if time.monotonic() - t0 > 6.0 or len(users) >= 10:
            break
        for row in wb[s2].iter_rows():
            if time.monotonic() - t0 > 6.0:
                break                      # one large sheet used to run the whole walk past its own clock
            for c2 in row:
                v = c2.value
                if isinstance(v, str) and v.startswith("=") and re.search(rf"(?<![A-Z0-9]){co}(?![0-9])", v) \
                        and (s2 == sh or re.search(rf"'?{re.escape(sh)}'?!", v)):
                    users.append(f"{s2}!{c2.coordinate} '{_label(wb[s2], c2.row)}'")
                    if len(users) >= 10:
                        break
            if len(users) >= 10:
                break
    out.append("    used by: " + ("; ".join(users) if users else "nothing found in a 6 s walk"))
    cand = []
    for v in ({_num(ev.cell(sh, co)), _num(ev0.cell(sh, co))} | {_num(ev0.cell(sh, f"{c}{r}")) for c in _hist_cols(loop, sh)}):
        if not isinstance(v, (int, float)) or v == 0:
            continue
        for it, _f in _value_index(loop).get(round(abs(v), 1), [])[:3]:
            cand.append(f"{_fmt(v)} on p{getattr(it, 'page', '?')}: "
                        + " ".join(str(getattr(it, "source_line", "") or getattr(it, "label", "")).split())[:80])
    out.append("    printed lines carrying this row's figures: " + ("; ".join(cand[:8]) if cand else "none on file"))
    return out


def t_find(loop, q):
    q = str(q).strip()
    out = [f"find {q!r}"]
    hits = []
    try:
        v = float(q.replace(",", ""))
    except ValueError:
        v = None
    if v is not None:
        hits = [it for it, _f in _hits(loop, v)]
    else:
        ql = q.lower()
        hits = [it for it in (getattr(loop.ledger, "items", []) or [])
                if ql in str(getattr(it, "label", "") or "").lower()
                or ql in str(getattr(it, "source_line", "") or "").lower()]
    for it in hits[:15]:
        out.append(f"    {getattr(it, 'doc', '?')} p{getattr(it, 'page', '?')}: "
                   + " ".join(str(getattr(it, "source_line", "") or getattr(it, "label", "")).split())[:110])
    if not hits:
        out.append("    nothing on file")
    elif len(hits) > 15:
        out.append(f"    ({len(hits)} lines match; the first 15 are shown — narrow the query)")
    return out


_ARITH = re.compile(r"[-+]?\d[\d,]*\.?\d*(?:\s*[-+*/]\s*[-+]?\d[\d,]*\.?\d*)+")


def _evidence_verdict(loop, sets, because):
    """Does this change land plain, and is its MAGNITUDE proven? The quote is on
    a page and that row's comparative ties the model's prior (proven: the world
    band steps aside), or the brain states arithmetic code re-computes (plain,
    but the band still polices it — code verified the sum, not the operands).
    -> (plain: bool, proven: bool, why: str)"""
    from .writegate import ties_prior
    because = str(because or "")
    proven, arith_only = True, False
    for s_ in sets:
        sh, co, v = s_["sheet"], s_["coord"], s_["value"]
        if not isinstance(v, (int, float)):
            return False, False, "the value is not a number code can tie"
        pcol = prior_column(loop.spec, sh, int(loop.ty))
        prior = _num(Evaluator(loop.wb).cell(sh, f"{pcol}{_row_of(co)}")) if pcol else None
        words = re.findall(r"[A-Za-z一-鿿]{4,}", because)
        quoted = False
        for it, scale in _hits(loop, v)[:20]:
            line = _quote(it)
            named = any(w.lower() in line.lower() for w in words) or str(getattr(it, "page", "")) in because
            if named and ties_prior(it, scale, prior, v):
                quoted = True
                break
        if quoted:
            continue
        # EVERY arithmetic phrase in the reason is tried, not the first one a
        # regex happens to find ('pp. 12-14: 396.2 + 36' matched '12-14' and a
        # genuine derivation landed red)
        if any(_arith_ties(m, v) for m in _ARITH.finditer(because)):
            proven, arith_only = False, True
            continue
        return False, False, f"{sh}!{co}: no quoted page line whose comparative ties the model's prior, and no arithmetic to verify"
    return True, proven, ("the arithmetic re-computes (the operands are the brain's, so the world band still polices it)"
                          if arith_only else "the evidence ties")


def _arith_ties(m, v):
    try:
        got = float(eval(m.group(0).replace(",", ""), {"__builtins__": {}}, {}))  # noqa: S307
    except Exception:  # noqa: BLE001
        return False
    return abs(got - v) <= max(0.05, abs(v) * 1e-4)


def _sets_text(sets):
    return ", ".join("{0}!{1}={2}".format(s["sheet"], s["coord"], _fmt(_num(s["value"]))) for s in sets)


def _apply_sets(loop, sets, plain, proven, note, log):
    """One change — a pair is ONE write batch: all of it lands or none of it
    does (a half-applied compensating pair is worse than neither half)."""
    w, applied, refused = loop.writer, [], []
    snap = snapshot(loop)
    try:
        for s_ in sets:
            sh, co, v = s_["sheet"], s_["coord"], s_["value"]
            pcol = prior_column(loop.spec, sh, int(loop.ty))
            ok = w.write(sh, co, v, prior_coord=f"{pcol}{_row_of(co)}" if pcol else None,
                         trusted=bool(proven), force_lock=True, flag=None if plain else "red", note=note)
            (applied if ok else refused).append(f"{sh}!{co}")
            if not ok:
                continue
            if not plain:
                loop.served.pop((sh, _row_of(co)), None)
            back = Evaluator(loop.wb).cell(sh, co)          # read back: never trust the write code
            if isinstance(back, (int, float)) and abs(back - v) > max(0.05, abs(v) * 1e-6):
                log(f"[review] read-back {sh}!{co} shows {back} after writing {v}")
    except Exception as e:  # noqa: BLE001
        restore(loop, snap)
        return [], [f"{s_['sheet']}!{s_['coord']}" for s_ in sets], f"the write raised {type(e).__name__}: {str(e)[:90]}"
    if refused and applied:
        restore(loop, snap)
        return [], refused + applied, "the writer refused part of the change, so none of it was kept"
    return applied, refused, ""


# ── the loop ─────────────────────────────────────────────────────────────

def _one_call(loop, pre_wb, call, key_panel, panel_path, log, repair_round, gate_once, hold_zero, state):
    """Answer one tool call. -> (lines, result) where result is None, or the
    gate's tuple when the model changed."""
    tool = str(call.get("tool") or "").strip().lower()
    if tool == "show":
        return t_show(loop, pre_wb, call.get("ref", "")), None
    if tool == "find":
        return t_find(loop, call.get("q", "")), None
    if tool in ("try", "set"):
        sets, bad = [], []
        for s_ in (call.get("sets") or ([call] if call.get("ref") else [])):
            sh, co, err = _parse_ref(loop, s_.get("ref", ""))
            if err:
                bad.append(err)
                continue
            if _col_of(co) != _act_col(loop, sh):
                bad.append(f"{sh}!{co} is not in the actual column — the review changes the actual period only")
                continue
            v = _num(s_.get("value"))          # '46.3' is a number the brain typed as text; anything else is not a value
            if v is None:
                bad.append(f"{sh}!{co}: '{s_.get('value')}' is not a number — a model cell takes a figure")
                continue
            sets.append({"sheet": sh, "coord": co, "value": v})
        if bad:
            return [f"{tool}: " + "; ".join(bad)], None
        if not sets:
            return [f"{tool}: no cell named"], None
        m0 = _metrics(loop, key_panel, panel_path)
        if tool == "try":
            snap = snapshot(loop)
            try:
                _applied, refused, note = _apply_sets(loop, sets, True, True, "trial", log)
                repair_round("review try")
                lines = ["try " + _sets_text(sets) + ":"]
                if refused:
                    lines.append(f"    the writer refused {refused} — {note or 'this change cannot be written as it stands'}")
                lines += _metric_diff(m0, _metrics(loop, key_panel, panel_path))
            finally:
                restore(loop, snap)              # the repairs the trial ran are inside the snapshot and come back with it
            lines.append("    (the trial was unwound; the verdict and watch lines it wrote are the run's own record)")
            return lines, None
        plain, proven, why = _evidence_verdict(loop, sets, call.get("because"))
        note = ((f"Set by the review: {str(call.get('because'))[:200]}") if plain else
                (f"Set by the review WITHOUT tying evidence — please check: {str(call.get('because') or 'no reason given')[:200]}"))
        _applied, refused, said = _apply_sets(loop, sets, plain, proven, note, log)
        if hold_zero:
            hold_zero()
        repair_round("review")
        res = gate_once()
        lines = ["set " + _sets_text(sets) + " → " + ("plain, " + why if plain else "RED: " + why)]
        if refused:
            lines.append(f"    the writer refused {refused} — {said or 'a guard'} (said out loud, not silently dropped)")
        lines += _metric_diff(m0, _metrics(loop, key_panel, panel_path))
        return lines, res
    if tool == "restore":
        sh, co, err = _parse_ref(loop, call.get("ref", ""))
        if err:
            return [f"restore: {err}"], None
        if _col_of(co) != _act_col(loop, sh):
            return [f"restore {sh}!{co}: not in the actual column — the review changes the actual period only"], None
        from openpyxl.utils import column_index_from_string as _ci
        m0 = _metrics(loop, key_panel, panel_path)
        back = pre_wb[sh].cell(_row_of(co), _ci(_col_of(co))).value
        why = str(call.get("why") or "")[:200]
        ok = loop.writer.write(sh, co, back if back is not None else 0.0, trusted=True, force_lock=True, flag="red",
                               note=f"Reverted by the brain: {why}. This input was moved by the run and is put back to what you had.")
        if ok:
            loop.served.pop((sh, _row_of(co)), None)
        if hold_zero:
            hold_zero()
        repair_round("review restore")
        res = gate_once()
        return [f"restore {sh}!{co} → {_fmt(_num(back))} ({'written, red' if ok else 'the writer refused it'})"] \
            + _metric_diff(m0, _metrics(loop, key_panel, panel_path)), res
    if tool == "plug":
        sh, co, err = _parse_ref(loop, call.get("check", ""))
        if err:
            return [f"plug: {err}"], None
        m0 = _metrics(loop, key_panel, panel_path)
        loop.writer.plugs_allowed = True
        _plug_here(loop, sh, co, log)
        if hold_zero:
            hold_zero()
        repair_round("review plug")
        res = gate_once()
        state.setdefault("plugged", []).append(f"{sh}!{co}")
        return [f"plug {sh}!{co} ({str(call.get('why') or '')[:120]}) — the model's own residual site, orange, reported"] \
            + _metric_diff(m0, _metrics(loop, key_panel, panel_path)), res
    return [f"'{tool}' is not one of the tools: show, find, try, set, restore, plug, done"], None


def _finish(loop, pre_wb, log, res, statement):
    """The closing rows the report reads, whatever the exit: the headline lines
    as they stand NOW, the negatives beyond the horizon as a watch, the count."""
    from .sensecheck import _sense_row, headline_deltas, suspicious
    writer, lines = loop.writer, loop.writer.log.setdefault("ending", [])
    earlier = {s.get("name"): s for s in writer.log.get("sense_rows", []) if isinstance(s, dict) and s.get("stage") != "ending"}
    try:
        deltas = headline_deltas(loop.wb, pre_wb, loop.spec, int(loop.ty))
        bad = {d["name"] for d in suspicious(deltas)}
        for d in deltas:
            e = earlier.get(d.get("name")) or {}
            leaf = tuple(str(e["leaf"]).split("!", 1)) if e.get("leaf") and "!" in str(e["leaf"]) else None
            if d["name"] in bad:
                _sense_row(writer, d, "red", (e.get("text") or "") + " — still out of line after the review; your ruling", leaf, "ending")
            else:
                _sense_row(writer, d, e.get("verdict") if e.get("verdict") in ("fixed", "genuine", "unusual") else "inline",
                           e.get("text") or "", leaf, "ending")
    except Exception as e:  # noqa: BLE001
        _fault(loop, f"closing sense rows not written: {e!r}")
    try:
        for ref, (what, v, kind, _off) in _metrics(loop, _key_panel_of(loop), None).items():
            if kind != "watch" or v >= -0.5:
                continue
            w = (f"{what.split(' (')[0]} is negative ({_fmt(v)} at {ref}) — beyond the next two periods: "
                 "your assumptions, your call")
            lines.append("WATCH " + w)
            log("[review] watch: " + w)
            sh, co = ref.split("!", 1)
            writer.watch(sh, co, "Sense check: " + w)     # a forecast cell is watch-listed, never painted
    except Exception as e:  # noqa: BLE001
        _fault(loop, f"the watch was not written: {e!r}")
    for f in loop.__dict__.get("objective_faults", []):
        if f"objective not measured: {f}" not in lines:
            lines.append(f"objective not measured: {f}")
            log(f"[review] FAULT {f}")
    if statement:
        for k, v in statement.items():
            lines.append(f"{k}: {v}")
            log(f"[review] {k}: {v}")
    return res


def _inherited(res, spec):
    """The analyst's OWN pre-update breaks, named by the gate — reported, never
    worked: the review does not plug a standing imbalance it did not cause."""
    out = set()
    card = res[2] if isinstance(res, tuple) and len(res) > 2 and isinstance(res[2], dict) else {}
    for ln in card.get("inherited_breaks", []) or []:
        m = re.match(r"^(.*)!r(\d+) \((\d{4})\)", str(ln).split(": ")[0])
        col = year_columns(spec, m.group(1)).get(m.group(3)) if m else None
        if col:
            out.add((m.group(1), f"{col}{m.group(2)}"))
    return out


def _open(loop, keys_before, key_panel, panel_path, res):
    inh = _inherited(res, loop.spec)
    return [o for o in _broken(loop, keys_before, key_panel, panel_path) if (o[1], o[2]) not in inh]


def run_review(loop, pre_wb, log, ask_json, gate_once, repair_round, keys_before, key_panel, panel_path,
               deadline_s=720.0, max_turns=40, notes=(), brain=True, hold_zero=None):
    """THE REVIEW LOOP. Every turn: code measures and lays out the model, the
    brain calls tools, code applies and measures and logs the turn verbatim.
    `done` is accepted only with a statement per objective. At the clock the
    terminal ladder plugs any actual-year check still off — the only decision
    code makes. -> (ok, failures, card) from the last gate."""
    writer, lines = loop.writer, loop.writer.log.setdefault("ending", [])
    t0 = time.monotonic()
    if hold_zero:
        hold_zero()
    repair_round("first")
    result = gate_once()
    statement, state, answers = None, {}, []
    if ask_json is None or not brain:
        lines.append("no brain in this run: the review was not held")
        log("[review] no brain: straight to the last resort")
    else:
        for turn in range(max_turns):
            left = deadline_s - (time.monotonic() - t0)
            if left <= 0:
                lines.append("the clock ended the review; what remains is written up")
                log("[review] clock: what remains is written up")
                break
            try:
                ctx = build_context(loop, pre_wb, key_panel, panel_path, notes=notes, answers=answers)
            except Exception as e:  # noqa: BLE001
                # THE LAST RESORT IS NEVER SKIPPED (reviewer 2026-09-16: a raise in
                # the context walked out of the loop and the model shipped unbalanced)
                _fault(loop, f"the review context could not be built at turn {turn + 1}: {e!r}")
                lines.append(f"the review context could not be built ({type(e).__name__}) — what remains is written up")
                log(f"[review] turn {turn + 1}: the context could not be built ({e!r})")
                break
            log(f"[review] turn {turn + 1}: context {len(ctx):,} chars, {left / 60:.1f} min left")
            try:
                reply = ask_json(MANDATE, ctx)
            except Exception as e:  # noqa: BLE001
                lines.append(f"the brain gave no answer at turn {turn + 1} ({type(e).__name__}) — what remains is written up")
                log(f"[review] turn {turn + 1}: the brain answered nothing ({e!r})")
                break
            # VERBATIM, WHOLE (reviewer 2026-09-16: a truncated turn is not JSON, and
            # the replay that reads these lines then runs a shorter review in silence)
            log("[review] turn %d reply %s" % (turn + 1, json.dumps(reply, ensure_ascii=False)))
            calls = reply.get("calls") if isinstance(reply, dict) else None
            if not isinstance(calls, list) or not calls:
                answers = ["    your last reply carried no calls — answer with a JSON list of calls, or `done` with a statement per objective"]
                continue
            answers, finished = [], False
            for call in calls:
                if time.monotonic() - t0 > deadline_s:
                    answers.append("    the clock ended the review inside this turn; the rest of your calls were not run")
                    log("[review] clock: the rest of the turn's calls were not run")
                    finished = True
                    break
                if not isinstance(call, dict):
                    answers.append(f"    '{str(call)[:60]}' is not a call object")
                    continue
                if str(call.get("tool") or "").lower() == "done":
                    st = call.get("objectives") if isinstance(call.get("objectives"), dict) else {}
                    said = {k for k in ("balance", "keys", "rollforward") if str(st.get(k) or "").strip()}
                    if len(said) < 3 and _open(loop, keys_before, key_panel, panel_path, result):
                        answers.append("    done needs a reading of each objective — balance, keys, rollforward: "
                                       "'holds', or 'cannot be closed because ...'. What is still open:")
                        answers += _metrics_text(_metrics(loop, key_panel, panel_path))
                        continue
                    statement = {str(k): str(v)[:300] for k, v in st.items()}
                    log(f"[review] done: {json.dumps(statement, ensure_ascii=False)[:600]}")
                    finished = True
                    break
                try:
                    out, res = _one_call(loop, pre_wb, call, key_panel, panel_path, log, repair_round, gate_once, hold_zero, state)
                except Exception as e:  # noqa: BLE001
                    _fault(loop, f"the call {json.dumps(call, ensure_ascii=False)[:120]} failed: {e!r}")
                    out, res = [f"    that call failed: {type(e).__name__}: {str(e)[:120]}"], None
                answers += out
                for ln in out:
                    log(f"[review]   {ln.strip()[:300]}")
                if res is not None:
                    result = res
            if finished:
                break
    checks = [(o[1], o[2]) for o in _open(loop, keys_before, key_panel, panel_path, result) if o[0] == "check"]
    if checks:
        # THE ONLY DECISION CODE MAKES (owner 2026-09-15: "a balanced model with a
        # flagged plug beats an unbalanced model"): whatever was said, an actual-year
        # check still off at the end goes to the ladder, orange, reported.
        from .orchestrator import terminal_ladder
        log(f"[review] last resort: {len(checks)} actual-year check(s) still off — the model's own plug ladder, orange")
        lines.append(f"LAST RESORT: {len(checks)} actual-year check(s) went to the plug ladder")
        try:
            loop.writer.plugs_allowed = True
            terminal_ladder(loop, log)
            if hold_zero:
                hold_zero()
            repair_round("review last resort")
            result = gate_once()
        except Exception as e:  # noqa: BLE001
            _fault(loop, f"the last resort failed: {e!r}")
    for o in _open(loop, keys_before, key_panel, panel_path, result):
        why = (statement or {}).get("balance" if o[0] in ("check", "forecast-check") else ("keys" if o[0] == "key" else "rollforward"))
        writer.flag_ref(f"{o[1]}!{o[2]}", "red", f"OPEN: {o[4]}" + (f" — the review's reading: {why}" if why else " — the review did not close it"))
        lines.append(f"OPEN {o[1]}!{o[2]}: {o[4]}")
        log(f"[review] {lines[-1]}")
    left_n = len(_open(loop, keys_before, key_panel, panel_path, result))
    lines.append(f"ended: {left_n} objective(s) still broken" if left_n else "ended: every objective holds")
    log(f"[review] {lines[-1]}")
    return _finish(loop, pre_wb, log, result if result is not None else gate_once(), statement)
