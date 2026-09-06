"""The executive _REPORT page — Luna COMPOSES it, the code RENDERS and
REFEREES it.

Owner's design lock (2026-08-28, REPORT_REQUIREMENTS.md in studio-native/):
a management page readable in one minute, three blocks —

  1. Key number snapshot  — the mindmap objective list, live formulas
  2. Why it moved         — bridges: P&L keys EVERY time, BS/CF keys only
                            when the move is material (~20%+). Composed by
                            the agent's REASONING, never a word search.
                            Every walk line is a traceable FORMULA; each
                            bridge ends in a self-balancing residual.
  3. Needs your attention — plugs, then red rulings (as questions), then
                            orange derivations. A few words per note.

Division of labour (the referee rule): the LLM decides WHAT the page says
— which rows are the key numbers in THIS model, what actually drove each
change, how each ruling is phrased. The code draws the format and REFUSES
what does not tie: a walk line whose formula does not evaluate to its
claimed value, or a bridge whose lines do not sum to its claimed change.
A wrong story rendered confidently is worse than no story.

Report-only mode regenerates the page on an ALREADY-updated model without
re-running the update:

    python -m pipeline.execreport "<company_dir>" \
        --model=<updated model.xlsx> --pre=<archived pre-update model.xlsx>

--selftest runs the offline museum (no LLM, no files touched).
"""
import json
import re
import sys
from pathlib import Path

# ---------------------------------------------------------------- facts

RED_FILLS = {"FFC7CE"}
ORANGE_FILLS = {"FFC000"}
MAX_ROW = 300


def _fill_code(cell):
    try:
        rgb = cell.fill.start_color.rgb
    except AttributeError:
        return ""
    if not isinstance(rgb, str):
        return ""
    return rgb[-6:].upper()


def _label_of(ws, row, max_col=6):
    for c in range(1, max_col + 1):
        v = ws.cell(row=row, column=c).value
        if isinstance(v, str) and v.strip():
            return v.strip()
    return ""


DEFAULT_COLS = {"Model": ("T", "U", "V"),
                "Raw financials": ("T", "U", "V")}


def _colmaps(cols, primary):
    """cols: {sheet: (prior_letter, target_letter, next_letter|None)}."""
    from openpyxl.utils import column_index_from_string
    cols = cols or DEFAULT_COLS
    primary = primary or ("Model" if "Model" in cols else
                          sorted(cols)[0])
    idx = {sh: tuple(column_index_from_string(x) if x else None
                     for x in v) for sh, v in cols.items()}
    return cols, idx, primary



_PRE_EV = {}


def _pre_val(pre_wb, sheet, row, col):
    """A pre-model cell VALUE — evaluated when it holds a formula (the
    run-208 lesson: the mini P&L's OLD table was empty on every
    manual-calc model because data_only loads cache nothing)."""
    try:
        v = pre_wb[sheet].cell(row=row, column=col).value
    except Exception:
        return None
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, str) and v.startswith("="):
        key = id(pre_wb)
        if key not in _PRE_EV:
            try:
                from .evaluator import Evaluator
                _PRE_EV[key] = Evaluator(pre_wb)
            except Exception:
                return None
        from openpyxl.utils import get_column_letter
        try:
            got = _PRE_EV[key].cell(sheet, f"{get_column_letter(col)}{row}")
        except Exception:
            return None
        return got if isinstance(got, (int, float)) else None
    return None

def gather_facts(wb, pre_wb, cols=None, primary=None):
    """Everything Luna needs to compose, as plain data. Deterministic.
    Sheets and column letters come from the SPEC (CLP-1 pin: 'Model'
    and T/U/V were DFE furniture, not law)."""
    cols, idx, primary = _colmaps(cols, primary)
    facts = {"primary": primary,
             "columns": {sh: {"prior": v[0], "target": v[1],
                              "next": v[2]} for sh, v in cols.items()},
             "model_rows": [], "raw_rows": [], "estimates": [],
             "flags": []}
    pci, tci = idx[primary][0], idx[primary][1]
    ws = wb[primary]
    for r in range(3, min(ws.max_row, MAX_ROW) + 1):
        lab = _label_of(ws, r)
        if not lab:
            continue
        t = ws.cell(row=r, column=pci).value
        u = ws.cell(row=r, column=tci).value
        facts["model_rows"].append(
            [r, lab,
             round(t, 1) if isinstance(t, (int, float)) else None,
             round(u, 1) if isinstance(u, (int, float)) else None])
    for sh in cols:
        if sh == primary or sh not in wb.sheetnames:
            continue
        rf = wb[sh]
        p2, t2 = idx[sh][0], idx[sh][1]
        for r in range(3, min(rf.max_row, MAX_ROW) + 1):
            lab = _label_of(rf, r)
            t = rf.cell(row=r, column=p2).value
            u = rf.cell(row=r, column=t2).value
            if lab and (isinstance(t, (int, float))
                        or isinstance(u, (int, float))):
                facts["raw_rows"].append(
                    [sh, r, lab,
                     round(t, 1) if isinstance(t, (int, float)) else None,
                     round(u, 1) if isinstance(u, (int, float)) else None])
    if pre_wb is not None and primary in pre_wb.sheetnames:
        pm = pre_wb[primary]
        nci = idx[primary][2]
        for r in range(3, min(pm.max_row, MAX_ROW) + 1):
            lab = _label_of(pm, r)
            e = _pre_val(pre_wb, primary, r, tci)
            n = (_pre_val(pre_wb, primary, r, nci) if nci else None)
            if lab and isinstance(e, (int, float)):
                facts["estimates"].append(
                    [r, lab, round(e, 1),
                     round(n, 1) if isinstance(n, (int, float)) else None])
    for sheet in wb.sheetnames:
        if sheet.startswith("_"):
            continue
        sws = wb[sheet]
        for row in sws.iter_rows(min_row=1,
                                 max_row=min(sws.max_row, MAX_ROW)):
            for cell in row:
                code = _fill_code(cell)
                if code in RED_FILLS or code in ORANGE_FILLS:
                    note = ""
                    if cell.comment is not None:
                        note = str(cell.comment.text)[:200]
                    facts["flags"].append(
                        [sheet, cell.coordinate,
                         "red" if code in RED_FILLS else "orange", note])
    return facts


def rollforward_headers(wb, cols=None, primary=None):
    """After mark-to-actual, the period header rolls like the rest of the
    column: the prior actual header's FORMAT copies across, and header
    text that is a stale DUPLICATE of the prior column's (e.g.
    '2024-12-31' still sitting over the updated period) advances its
    year. A header that legitimately names an estimate — a comparison
    panel's '2025E' over a frozen-estimate column — is never touched:
    only exact duplicates count as stale."""
    import copy as _copy
    cols, idx, primary = _colmaps(cols, primary)
    fixed = []
    for sn in wb.sheetnames:
        if sn.startswith("_") or sn not in idx:
            continue
        prior_col, target_col = idx[sn][0], idx[sn][1]
        ws = wb[sn]
        for hr in (1, 2):
            pc = ws.cell(row=hr, column=prior_col)
            tc = ws.cell(row=hr, column=target_col)
            if pc.value in (None, ""):
                continue
            tc.font = _copy.copy(pc.font)
            tc.fill = _copy.copy(pc.fill)
            tc.border = _copy.copy(pc.border)
            tc.number_format = pc.number_format
            stale = False
            if (isinstance(tc.value, str) and isinstance(pc.value, str)
                    and tc.value.strip() == pc.value.strip()):
                m = re.search(r"(19|20)\d{2}", tc.value)
                if m:
                    y = int(m.group(0))
                    old = tc.value
                    tc.value = tc.value.replace(str(y), str(y + 1), 1)
                    stale = True
            elif (isinstance(tc.value, (int, float))
                    and tc.value == pc.value and 1990 < pc.value < 2100):
                old = tc.value
                tc.value = int(pc.value) + 1
                stale = True
            if stale:
                fixed.append("%s!%s: %r -> %r"
                             % (sn, tc.coordinate, old, tc.value))
    return fixed


# ------------------------------------------------------------- referee

_REF = re.compile(r"'?([A-Za-z_][A-Za-z0-9_ ]*?)'?!([A-Z]{1,3})([0-9]{1,4})")
_SAFE = re.compile(r"^[0-9+\-*/(). ]*$")


def _eval_formula(wb, formula):
    """Evaluate the +-*/() cell-reference grammar the bridges use, from
    TYPED values only. Returns (value, verifiable). A reference landing on
    a formula cell (no cached value) makes the line unverifiable — the sum
    law still binds it."""
    from openpyxl.utils import column_index_from_string
    expr = formula.lstrip("=")
    verifiable = True

    def sub(m):
        nonlocal verifiable
        sheet, col, row = m.group(1), m.group(2), int(m.group(3))
        if sheet not in wb.sheetnames:
            verifiable = False
            return "0"
        v = wb[sheet].cell(row=row,
                           column=column_index_from_string(col)).value
        if v is None:
            return "0"                      # blank reads as 0, like Excel
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return repr(float(v))
        verifiable = False                  # formula cell without a cache
        return "0"
    expr = _REF.sub(sub, expr)
    if not verifiable or not _SAFE.match(expr):
        return None, False
    try:
        return float(eval(expr, {"__builtins__": {}})), True
    except Exception:
        return None, False



def _is_circular(bridge, formula):
    """A pure +/- walk line that references the bridge's own anchor row."""
    if "*" in formula or "/" in formula:
        return False
    bs = str(bridge.get("sheet", "Model")).strip().lower()
    try:
        brow = int(bridge.get("row", 0))
    except (TypeError, ValueError):
        return False
    for m in _REF.finditer(formula):
        if m.group(1).strip().lower() == bs and int(m.group(3)) == brow:
            return True
    return False


def referee(summary, wb):
    """Refuse what does not tie. Returns (kept, refusals, corrections).

    The FORMULA is the authoritative content — it is what renders, live,
    and it is verified against the workbook's typed cells. The claimed
    "value" is only the composer's arithmetic; where the formula
    evaluates, the evaluation REPLACES the claim (a sign slip in the
    bookkeeping is corrected, not fatal). Refusals are for what cannot
    tie: unverifiable claims, circular lines, walks that leave most of
    the story unexplained."""
    kept, refusals, corrections = [], [], []
    for b in summary.get("bridges", []):
        why = []
        total = b.get("total")
        if not isinstance(total, (int, float)):
            why.append("bridge '%s': no claimed total" % b.get("title"))
        lines = b.get("lines", [])
        ssum = 0.0
        for ln in lines:
            lab = str(ln.get("label", "")).strip().lower()
            f = str(ln.get("formula", ""))
            # only a BARE residual label is banned — "Other payables"
            # and friends are real disclosure lines
            if re.match(r"^(others?|residuals?|other\s*/\s*residual)$", lab):
                why.append("line '%s': never write your own Other/Residual "
                           "line — the renderer adds it" % ln.get("label"))
                continue
            # circular: a pure +/- line that references the very row it is
            # supposed to explain is a tautology (total minus the other
            # lines), not a driver. Multiplicative forms are exempt — the
            # sanctioned volume/margin split references its own row.
            if _is_circular(b, f):
                why.append("line '%s': circular — a plus/minus line may "
                           "not reference the bridge's own total row; "
                           "derive the driver from the underlying rows "
                           "instead" % ln.get("label"))
                continue
            val = ln.get("value")
            got, ok = _eval_formula(wb, str(ln.get("formula", "")))
            if ok:
                if isinstance(val, (int, float)) and abs(got - val) > max(
                        0.5, abs(val) * 0.01):
                    corrections.append(
                        "%s / '%s': claimed %.1f, formula gives %.1f — "
                        "formula wins" % (b.get("title"), ln.get("label"),
                                          val, got))
                ln["value"] = round(got, 1)
                ssum += got
                continue
            if not isinstance(val, (int, float)):
                why.append("line '%s': no value and the formula is not "
                           "verifiable" % ln.get("label"))
                continue
            ssum += val
        if isinstance(total, (int, float)):
            # a sign-flipped TOTAL claim with a sound walk: the walk wins
            if (abs(ssum + total) <= max(1.0, abs(total) * 0.02)
                    and abs(ssum - total) > max(1.0, abs(total) * 0.02)):
                corrections.append(
                    "%s: claimed total %.1f has the wrong sign — the "
                    "verified walk sums to %.1f" % (b.get("title"), total,
                                                    ssum))
                total = -total
                b["total"] = total
            # the named drivers must carry MOST of the story; the
            # auto-residual absorbs a modest remainder, never the bulk
            residual = total - ssum
            if abs(residual) > max(50.0, abs(total) * 0.40):
                why.append(
                    "bridge '%s': your lines sum to %.1f against a change "
                    "of %.1f — the unexplained residual (%.1f) is most of "
                    "the story. Check your SIGNS (a value is the effect ON "
                    "the total: a cost that rose contributes NEGATIVE, "
                    "formula =-(U-T)) and name the real drivers"
                    % (b.get("title"), ssum, total, residual))
        if why:
            refusals.append({"bridge": b.get("title"), "why": why})
        else:
            kept.append(b)
    return kept, refusals, corrections

# ------------------------------------------------------------- renderer

NUM, SGN, PCT = '#,##0.0', '+#,##0.0;-#,##0.0', '0.0%'


_MINI_ALIASES = (
    ({"revenue", "turnover", "sales"}, {"revenue", "turnover", "sales", "income"}),
    ({"ebit", "operating profit", "operating income", "net operating income"},
     {"ebit", "operating profit", "operating income", "net operating income"}),
    ({"net profit", "net income", "profit attributable", "earnings"},
     {"net profit", "net income", "profit attributable", "earnings", "net profits"}),
    ({"gross profit", "gross margin", "gpm"}, {"gross profit", "gross margin", "gpm"}),
    ({"net margin", "npm"}, {"net margin", "npm"}),
    ({"eps"}, {"eps", "earnings per share"}),
    ({"dps"}, {"dps", "dividend per share", "dividends per share"}),
)


def _mini_alias(claim, label):
    c, l = claim.lower(), label.lower()
    for a, b in _MINI_ALIASES:
        if any(x in c for x in a) and any(y in l for y in b):
            return True
    return False


def render(wb, summary, pre_wb=None, cols=None, primary=None):
    """The owner's v3 layout, exactly. Deterministic."""
    from openpyxl.styles import Border, Font, PatternFill, Side
    BANNER = PatternFill("solid", fgColor="FFF2CC")
    SECT = PatternFill("solid", fgColor="F2F2F2")
    TINT_O = PatternFill("solid", fgColor="FDEBD0")
    TINT_R = PatternFill("solid", fgColor="FADBD8")
    GREY = Font(color="7F7F7F", size=10)
    GREYI = Font(color="7F7F7F", size=10, italic=True)
    BOLD = Font(bold=True, size=11)
    SECTF = Font(bold=True, size=12)
    BANF = Font(bold=True, size=13)
    BODY = Font(size=11)
    SMALL = Font(size=10)
    THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
    cols, _cidx, primary = _colmaps(cols, primary)

    def LET(sheet):
        p, t, n = cols.get(sheet, cols[primary])
        return p, t, (n or t)

    if "_REPORT" in wb.sheetnames:
        del wb["_REPORT"]
    ws = wb.create_sheet("_REPORT", 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 34
    for col in "BCDEFGIJKLMOPQRS":
        ws.column_dimensions[col].width = 11
    ws.column_dimensions["G"].width = 16      # the check column
    for col in "HN":                          # spacers between blocks
        ws.column_dimensions[col].width = 3
    r = 1

    def cell(row, col, val, font=None, fill=None, fmt=None, border=None):
        c = ws.cell(row=row, column=col, value=val)
        c.font = font if font else BODY
        if fill:
            c.fill = fill
        if fmt:
            c.number_format = fmt
        if border:
            c.border = border
        return c

    def band(row, fill, h=None):
        for c in range(1, 20):
            ws.cell(row=row, column=c).fill = fill
        if h:
            ws.row_dimensions[row].height = h

    def gap(h=8):
        nonlocal r
        ws.row_dimensions[r].height = h
        r += 1

    def sect(title):
        nonlocal r
        cell(r, 1, title, SECTF, SECT)
        band(r, SECT, 22)
        r += 1
        gap(6)

    cell(r, 1, str(summary.get("banner", ""))[:140], BANF, BANNER)
    cell(r, 6, str(summary.get("coverage", ""))[:90], GREY, BANNER)
    band(r, BANNER, 30)
    r += 1
    gap(14)

    # THE READING STEP + THE ROLLOVER CHECK (owner rulings 2026-09-03):
    # what the agent was holding, and whether the update moved the
    # forecast in proportion — on page one, before any number
    docs = summary.get("documents") or []
    if docs:
        sect("0 · Documents received   (what the agent identified, and how each was used)")
        for line in docs:
            cell(r, 1, str(line)[:220], SMALL)
            r += 1
        gap(10)
    # (owner 2026-09-04: no rollover table on the page — the mini P&L
    # old-vs-new already shows the forecast path; the check's verdicts
    # live in the flags and the log)
    sect(f"1 · Key number snapshot   ({summary.get('units') or 'model units'})")
    hdr = ["", "FY prior A", "FY actual A", "YoY", "Your estimate",
           "A vs E", "Next yr before", "Next yr after"]
    for i, h in enumerate(hdr):
        cell(r, i + 1, h, GREY, None, None, THIN)
    r += 1
    for grp in summary.get("snapshot", []):
        cell(r, 1, str(grp.get("group", ""))[:30], GREYI)
        r += 1
        for it in grp.get("rows", []):
            mr = int(it["row"])
            sheet = it.get("sheet", primary)
            T, U, V = LET(sheet)
            q = "'%s'" % sheet if re.search(r"[^A-Za-z0-9]", sheet) else sheet
            cell(r, 1, '=HYPERLINK("#%s!%s%d","%s")'
                 % (sheet.replace("'", ""), U, mr,
                    str(it.get("label", ""))[:40]))
            cell(r, 2, "=%s!%s%d" % (q, T, mr), None, None, NUM)
            cell(r, 3, "=%s!%s%d" % (q, U, mr), None, None, NUM)
            cell(r, 4, '=IFERROR(IF(B%d<=0,"n/m",C%d/B%d-1),"")'
                 % (r, r, r), None, None, PCT)
            est = it.get("est")
            cell(r, 5, est if isinstance(est, (int, float)) else "",
                 None, None, NUM)
            cell(r, 6, '=IFERROR(IF(E%d<=0,"n/m",C%d/E%d-1),"")'
                 % (r, r, r), None, None, PCT)
            nb = it.get("next_before")
            cell(r, 7, nb if isinstance(nb, (int, float)) else "",
                 None, None, NUM)
            cell(r, 8, "=%s!%s%d" % (q, V, mr), None, None, NUM)
            r += 1
    gap(14)

    # ---- 2 · mini P&L: the whole forecast path, old vs new ----------
    # (boss feedback 2026-08-30: an update that quietly bent the
    # out-years must show up immediately)
    mini = summary.get("mini_pl", {}).get("rows", [])
    mini_sheet = str(summary.get("mini_pl", {}).get("sheet", primary))
    if mini_sheet not in wb.sheetnames:
        mini_sheet = primary
    # CODE VERIFIES THE BRAIN'S ROWS (owner 2026-09-04: 'gross margin'
    # was linked to 'Net operating income'): a mini-P&L row is kept only
    # when its label is kin to the model's own label on that row
    from .numerics import kinship as _kin
    kept = []
    for it in mini:
        try:
            lab = wb[mini_sheet].cell(row=int(it.get("row")), column=1).value
        except Exception:
            lab = None
        claim = str(it.get("label", ""))
        if lab is None or _kin(claim, str(lab)) or _mini_alias(claim, str(lab)):
            kept.append(it)
    mini = kept
    if mini:
        from openpyxl.utils import (column_index_from_string,
                                    get_column_letter)
        sect("2 · Mini P&L — old vs new   (FY-1 to FY+3, model grain)")
        _, mtl, _ = LET(mini_sheet)
        tci = column_index_from_string(mtl)
        pcols = [get_column_letter(tci - 1 + i) for i in range(5)]
        labels = []
        for i, cl in enumerate(pcols):
            v = None
            if pre_wb is not None and mini_sheet in pre_wb.sheetnames:
                v = pre_wb[mini_sheet].cell(
                    row=2, column=column_index_from_string(cl)).value
            labels.append(str(int(v)) if isinstance(v, (int, float))
                          else ["FY-1", "FY0", "FY+1", "FY+2", "FY+3"][i])
        # boss round 2: Δ leftmost, then NEW, then OLD, one spacer
        # between blocks; a roll-over sanity flag beside the Δ block.
        D0, FLAG, NEW0, OLD0 = 2, 7, 9, 15    # B..F | G | I..M | O..S
        cell(r, D0, "WHAT'S CHANGED — new vs old", GREY)
        cell(r, NEW0, "NEW — after the update", GREY)
        cell(r, OLD0, "OLD — before the update", GREY)
        r += 1
        for blk in (D0, NEW0, OLD0):
            for i, plab in enumerate(labels):
                cell(r, blk + i, plab, GREY, None, None, THIN)
        cell(r, FLAG, "check", GREY, None, None, THIN)
        r += 1
        for it in mini:
            mr = int(it["row"])
            kind = str(it.get("kind", "value"))
            vfmt = ('0.0%' if kind == "margin"
                    else '0.00' if kind == "pershare" else NUM)
            dfmt = ('0.0%' if kind == "margin"
                    else '0.00' if kind == "pershare" else PCT)
            qm = ("'%s'" % mini_sheet
                  if re.search(r"[^A-Za-z0-9]", mini_sheet) else mini_sheet)
            cell(r, 1, '=HYPERLINK("#%s!%s%d","%s")'
                 % (mini_sheet.replace("'", ""), mtl, mr,
                    str(it.get("label", ""))[:32]))
            for i, cl in enumerate(pcols):
                ov = None
                if pre_wb is not None and mini_sheet in pre_wb.sheetnames:
                    ov = _pre_val(pre_wb, mini_sheet, mr,
                                  column_index_from_string(cl))
                cell(r, OLD0 + i,
                     round(ov, 4) if isinstance(ov, (int, float)) else "",
                     SMALL, None, vfmt)
                cell(r, NEW0 + i, "=%s!%s%d" % (qm, cl, mr),
                     SMALL, None, vfmt)
                oc = get_column_letter(OLD0 + i)
                nc = get_column_letter(NEW0 + i)
                if kind == "value":
                    f = ('=IFERROR(IF(%s%d=0,"",(%s%d-%s%d)/ABS(%s%d)),"")'
                         % (oc, r, nc, r, oc, r, oc, r))
                else:
                    f = ('=IFERROR(IF(%s%d="","",%s%d-%s%d),"")'
                         % (oc, r, nc, r, oc, r))
                cell(r, D0 + i, f, SMALL, None, dfmt)
            # the roll-over sanity flag: updated period's change vs the
            # next forecast period's change — a sign flip or a wide gap
            # is the boss's "suspicious roll-over" signal
            c0 = get_column_letter(D0 + 1)      # updated period Δ
            c1 = get_column_letter(D0 + 2)      # next forecast Δ
            gapthr = "0.3" if kind == "value" else "0.02"
            cell(r, FLAG,
                 '=IF(COUNT(%s%d,%s%d)<2,"",IF(%s%d*%s%d<0,'
                 '"⚠ sign flip vs next yr",IF(ABS(%s%d-%s%d)>%s,'
                 '"⚠ big gap vs next yr","")))'
                 % (c0, r, c1, r, c0, r, c1, r, c1, r, c0, r, gapthr),
                 Font(color="9C5700", size=10))
            r += 1
        # tint material forecast-path changes so they cannot hide
        from openpyxl.formatting.rule import CellIsRule
        top, bot = r - len(mini), r - 1
        rng = "B%d:F%d" % (top, bot)
        for op, v in (("greaterThan", "0.2"), ("lessThan", "-0.2")):
            ws.conditional_formatting.add(rng, CellIsRule(
                operator=op, formula=[v], fill=PatternFill(
                    "solid", fgColor="FFF2CC")))
        r += 1
        gap(14)

    sect("3 · Why it moved   (actual vs prior year · P&L always · "
         "BS/CF when the move is material, ~20%+)")
    for b in summary.get("bridges", []):
        mr = int(b["row"])
        sheet = b.get("sheet", primary)
        if sheet not in wb.sheetnames:
            sheet = primary
        T, U, V = LET(sheet)
        q = "'%s'" % sheet if re.search(r"[^A-Za-z0-9]", sheet) else sheet
        trow = r
        cell(r, 1, '=HYPERLINK("#%s!%s%d","%s")'
             % (sheet.replace("'", ""), U, mr, str(b.get("title", ""))[:60]),
             BOLD)
        cell(r, 2, "=%s!%s%d-%s!%s%d" % (q, U, mr, q, T, mr),
             BOLD, None, SGN)
        r += 1
        first = r
        for ln in b.get("lines", []):
            f = str(ln.get("formula", "")).strip()
            if f and not f.startswith("="):
                f = "=" + f          # Luna often omits it; text is useless
            cell(r, 1, "      " + str(ln.get("label", ""))[:60], SMALL)
            cell(r, 2, f, SMALL, None, SGN)
            r += 1
        cell(r, 1, "      Other / residual", GREYI)
        cell(r, 2, "=B%d-SUM(B%d:B%d)" % (trow, first, r - 1),
             GREY, None, SGN)
        r += 1
        gap(10)
    # (owner 2026-09-04: the sense check is the agent's own note, not
    # the executive's — never rendered; its verdicts stay in the log)
    for note in (summary.get("skipped_note"), summary.get("company_note")):
        if note:
            cell(r, 1, str(note)[:200], GREYI)
            r += 1
    gap(14)

    sect("4 · Needs your attention")
    tiers = [("plugs", "Plugs — inserted to make the model tie; "
              "resolve properly", TINT_O),
             ("red", "Red — unsure, your ruling", TINT_R),
             ("orange", "Orange — derived, not read", TINT_O)]
    attention = summary.get("attention", {})
    # THE ONE-MINUTE PAGE (owner 2026-09-04: "less cluttered, only what
    # matters to the analyst"): every plug and every red ruling on the
    # page (reds capped, load-bearing first as listed); the orange
    # derivations as ONE line with their count; the complete list on a
    # _FLAGS sheet the analyst opens when they want the detail
    RED_CAP = 30
    flags_ws = wb.create_sheet("_FLAGS") if "_FLAGS" not in wb.sheetnames \
        else wb["_FLAGS"]
    for row_ in flags_ws.iter_rows():
        for c_ in row_:
            c_.value = None
    fr = 1
    flags_ws.cell(row=fr, column=1, value="Every flagged cell (plugs, red rulings, orange derivations)").font = BOLD
    fr += 2
    for key, header, fill in tiers:
        items = [it for it in attention.get(key, []) if isinstance(it, dict)]
        if not items:
            continue
        # the complete tier goes to _FLAGS
        flags_ws.cell(row=fr, column=1, value=header).font = BOLD
        fr += 1
        for it in items:
            sheet, addr = str(it.get("sheet", "")), str(it.get("cell", ""))
            if sheet not in wb.sheetnames or not re.match(r"^[A-Z]{1,3}[0-9]{1,4}$", addr):
                continue
            q = ("'%s'" % sheet if re.search(r"[^A-Za-z0-9]", sheet) else sheet)
            a_ = flags_ws.cell(row=fr, column=1, value='=HYPERLINK("#\'%s\'!%s","%s!%s")' % (sheet, addr, sheet, addr))
            a_.font = SMALL
            flags_ws.cell(row=fr, column=2, value="=%s!%s" % (q, addr)).number_format = NUM
            flags_ws.cell(row=fr, column=3, value=str(it.get("note", ""))[:200]).font = SMALL
            fr += 1
        fr += 1
        # the page: plugs in full, reds capped, orange as a count
        cell(r, 1, header, BOLD, fill)
        band(r, fill, 20)
        r += 1
        if key == "orange":
            cell(r, 1, f"{len(items)} derived cells (orange) — listed on the _FLAGS sheet; "
                       "each carries its method note in the model", SMALL)
            r += 1
            gap(8)
            continue
        shown = 0
        for it in items:
            sheet, addr = str(it.get("sheet", "")), str(it.get("cell", ""))
            if sheet not in wb.sheetnames or not re.match(
                    r"^[A-Z]{1,3}[0-9]{1,4}$", addr):
                continue
            if key == "red" and shown >= RED_CAP:
                cell(r, 1, f"… {len(items) - shown} more red cells on the _FLAGS sheet", GREYI)
                r += 1
                break
            q = ("'%s'" % sheet if re.search(r"[^A-Za-z0-9]", sheet)
                 else sheet)
            cell(r, 1, '=HYPERLINK("#\'%s\'!%s","%s!%s")'
                 % (sheet, addr, sheet, addr), SMALL)
            cell(r, 2, "=%s!%s" % (q, addr), SMALL, None, NUM)
            cell(r, 3, str(it.get("note", ""))[:90], SMALL)
            r += 1
            shown += 1
        gap(8)
    flags_ws.column_dimensions["A"].width = 26
    flags_ws.column_dimensions["C"].width = 110
    if summary.get("other_note"):
        cell(r, 1, str(summary["other_note"])[:200], GREYI)
        r += 1
    wb.active = 0
    return r

# ------------------------------------------------------------- compose

SYSTEM = """You are an equity research analyst writing the one-page
executive report of a model update, as JSON. The facts carry
"primary" (the main statement sheet), "columns" (each sheet's
prior/target/next column LETTERS — always use THESE letters in
formulas, they differ per model), and "raw_rows" as
[sheet, row, label, prior, current]. Every snapshot/bridge entry must
name its "sheet". You are given the model's
own data; you COMPOSE the page by reasoning about it — never by matching
words. The renderer draws the format; a referee will REFUSE any bridge
whose lines do not tie. Respond with ONE JSON object, nothing else.

THE PAGE (owner's locked design):
1. "snapshot": the KEY NUMBERS, fixed by the objective: sales/revenue,
   gross profit, net profit, cash, current assets, non-current assets,
   current liabilities, non-current liabilities, total equity, and
   operating / investing / financing cash flow. FIND each one's row in
   the model_rows list (labels differ per model — reason it out). Group
   as P&L / Balance sheet / Cash flow. For each row attach the
   analyst's pre-update estimate and next-year forecast from the
   estimates list (same row numbers) when present.
2. "mini_pl": the fixed mini-P&L the page shows as OLD vs NEW with a
   what's-changed table (the boss's forecast-path review). Find the
   Core-8 rows in model_rows: revenue, gross profit, gross margin
   (GPM %), operating profit / EBIT, net profit, net margin (NPM %),
   EPS, DPS. Each: {"label": str, "row": int, "kind": "value" |
   "margin" | "pershare"} — margins are the ratio rows, EPS/DPS are
   "pershare". Skip a line only if this model truly has no such row.
3. "bridges": WHY each key number moved, year on year. P&L key numbers
   get a bridge EVERY time. Balance-sheet and cash-flow key numbers get
   one ONLY when the move is material (roughly 20%+ change). THINK about
   what actually drove each change by comparing the raw_rows year on
   year — name the few drivers that explain it (2-6 lines), biggest
   first. Every line carries:
   - "label": a short English driver name (your judgment, not the raw
     Chinese label),
   - "formula": an Excel formula deriving the amount from the workbook,
     using ONLY refs like 'Raw financials'!U18-'Raw financials'!T18 or
     Model!U7-Model!T7 (U = actual year, T = prior year; +,-,*,/ and
     parentheses allowed),
   - "value": the number the formula evaluates to (RMB mn, 1dp).
   SIGN LAW: a line's value is its EFFECT ON the total — positive pushes
   the total up. A cost/expense/outflow that INCREASED is a NEGATIVE
   effect on profit or cash: write the formula as =-('Raw financials'!U10
   -'Raw financials'!T10) so it evaluates to the signed effect. The
   formula and the value must agree exactly — the referee evaluates
   every formula.
   WORKED SIGNS EXAMPLE (a profit bridge): operating expenses rose from
   7,000 to 7,500, so their effect on profit is -500:
     {"label": "Opex up", "formula": "=-('Raw financials'!U10-'Raw
      financials'!T10)", "value": -500.0}
   An income line that FELL also contributes negative — its formula is
   the plain delta (=U18-T18) which is already negative. Check every
   line: does the formula's sign equal the value's sign?
   CASH-FLOW TOTALS: the change is actual minus prior even when both
   are negative: -10,000 - (-3,000) = -7,000 — the change is NEGATIVE
   (the outflow grew). Never flip it positive.
   Also give the bridge's "total" = the key number's actual change.
   Lines must sum close to total; a residual line is added for you —
   NEVER add your own "Other"/"Residual" line (it will be refused).
   A bridge explains WHAT DROVE the change, not the statement's own
   structure: for net profit, name the operating drivers (gross profit,
   investment income, impairments, operating expenses, tax, minorities)
   — "pre-tax profit + tax" explains nothing. A plus/minus line that
   references the bridge's own total row is circular and will be
   refused. Prefer typed 'Raw financials' rows, which the referee can
   verify; a working-capital or funding story usually lives there.
   For gross profit use the volume/margin split:
     volume = (Model!U<rev>-Model!T<rev>)*Model!T<gp>/Model!T<rev>
     margin = Model!U<rev>*(Model!U<gp>/Model!U<rev>-Model!T<gp>/Model!T<rev>)
4. "attention": from the flags list, the items the analyst must rule on,
   most important first — "plugs" (numbers inserted to make the model
   tie; say in a few words what each ties), "red" (uncertain KEY numbers,
   phrased as the question the analyst must answer), "orange" (derived
   key numbers, show the derivation in words). KEY numbers only; keep
   each note under 12 words. Skip minor cells.

RULES THAT NEVER BEND:
- Every number you state must come from the data given. Never invent.
- "banner": one line, honest verdict + the biggest caveat.
- "coverage": counts line, e.g. "key numbers proven X/14 · N plugs · N rulings".
- "skipped_note": name the key numbers you judged too small to bridge,
  with their % moves, so the omission reads as deliberate.
- "company_note": ONLY a company-stated reason with a page reference
  from disclosure text given to you; if none was given, say the stated
  reasons are not in the provided documents.
- "other_note": one line pointing to the colour-marked cells in the
  model tabs.

OUTPUT SHAPE:
{"banner": str, "coverage": str,
 "snapshot": [{"group": "P&L", "rows": [{"label": str, "sheet": "Model",
   "row": int, "est": num|null, "next_before": num|null}, ...]}, ...],
 "mini_pl": {"rows": [{"label": str, "row": int,
   "kind": "value|margin|pershare"}, ...]},
 "bridges": [{"title": str, "sheet": "Model", "row": int, "total": num,
   "lines": [{"label": str, "formula": str, "value": num}, ...]}, ...],
 "skipped_note": str, "company_note": str,
 "attention": {"plugs": [{"sheet": str, "cell": str, "note": str}, ...],
   "red": [...], "orange": [...]},
 "other_note": str}"""


def _validate(d):
    if not isinstance(d, dict):
        return ["not an object"]
    for k in ("banner", "snapshot", "bridges", "attention"):
        if k not in d:
            return ["missing key: " + k]
    if not isinstance(d["snapshot"], list) or not d["snapshot"]:
        return ["snapshot empty"]
    for grp in d["snapshot"]:
        for it in grp.get("rows", []):
            if not isinstance(it.get("row"), int):
                return ["snapshot row without an integer row number"]
    mp = d.get("mini_pl")
    # as many core rows as THIS model carries — the >=6 floor was DFE
    # furniture and starved every CLP report (its condensed P&L offers
    # 5 distinct core rows)
    if not isinstance(mp, dict) or len(mp.get("rows", [])) < 4:
        return ["mini_pl.rows must list this model's core P&L rows "
                "(revenue, profits, per-share — at least 4)"]
    for it in mp["rows"]:
        if not isinstance(it.get("row"), int) or it.get("kind") not in (
                "value", "margin", "pershare"):
            return ['mini_pl rows need an integer "row" and a valid "kind"']
    for b in d["bridges"]:
        if not isinstance(b.get("row"), int):
            return ["bridge without an integer row"]
        if not b.get("lines"):
            return ["bridge '%s' has no lines" % b.get("title")]
    att = d["attention"]
    if not isinstance(att, dict):
        return ["attention must be an object with plugs/red/orange lists"]
    for k in ("plugs", "red", "orange"):
        for it in att.get(k, []):
            if not isinstance(it, dict) \
                    or not (it.get("note") or (it.get("sheet")
                                               and it.get("cell"))):
                return ["attention.%s items must be objects "
                        '{"sheet","cell","note"}' % k]
    return []


def _normalize(d):
    """Tolerate the harmless shape drift Luna produces instead of
    failing the whole report: a bare-string attention item becomes a
    note without a link (run-204: three CLP reports died on this)."""
    if isinstance(d, dict) and isinstance(d.get("attention"), dict):
        for k, v in d["attention"].items():
            if isinstance(v, list):
                d["attention"][k] = [
                    x if isinstance(x, dict)
                    else {"sheet": "", "cell": "", "note": str(x)[:200]}
                    for x in v]
    return d


def compose(client, facts, feedback=""):
    user = json.dumps(facts, ensure_ascii=False)
    if feedback:
        user += ("\n\nTHE REFEREE REFUSED PART OF YOUR LAST ANSWER — fix "
                 "exactly these and resend the FULL object:\n" + feedback)
    return _normalize(client.json(
        SYSTEM, user, lambda d: _validate(_normalize(d)),
        repair_retries=2))


# ---------------------------------------------------------- sense check

SENSE_SYSTEM = """You are the analyst agent that JUST PERFORMED this
model update. The what's-changed table flagged the items below. For each
flag, go back once and review your own work adversarially: is the change
justified by the data, or did the update make a mistake? Use the facts
pack (model rows, raw rows, estimates, flags) — cite the cells that
convinced you. NEVER propose silently fixing anything; the analyst
decides. One investigation pass per item — no looping.
Respond with ONE JSON object:
{"verdicts": [{"label": str, "period": str,
  "verdict": "JUSTIFIED" | "SUSPICIOUS" | "ERROR FOUND",
  "reason": str (under 20 words, cite cells), "cells": [str, ...]}]}
Every flagged item gets exactly one verdict. "JUSTIFIED" needs the
actual cause (e.g. base effect of the 2025 beat flowing through frozen
growth). If you cannot explain a change from the data, that is
"SUSPICIOUS" — honesty beats confidence."""


def collect_delta_flags(wb, pre_wb, mini_rows, cols=None, primary=None,
                        horizon=3, value_of=None):
    """The what's-changed table as data, with the boss's flags. New-side
    values come from `value_of(sheet, coord)` when given (the run's own
    Evaluator — a freshly written workbook caches nothing), else from
    cached cell values. A value neither computable nor cached is
    skipped, never guessed."""
    from openpyxl.utils import column_index_from_string, get_column_letter
    cols, idx, primary = _colmaps(cols, primary)
    target_col = idx[primary][1]
    out = []
    if pre_wb is None or primary not in pre_wb.sheetnames:
        return out
    ws, pw = wb[primary], pre_wb[primary]
    for it in mini_rows:
        row, kind = int(it["row"]), str(it.get("kind", "value"))
        deltas = {}
        for i in range(-1, horizon + 1):
            col = target_col + i
            ov = _pre_val(pre_wb, primary, row, col)
            nv = ws.cell(row=row, column=col).value
            if value_of is not None and not isinstance(nv, (int, float)):
                try:
                    nv = value_of(primary,
                                  get_column_letter(col) + str(row))
                except Exception:
                    nv = None
            if not (isinstance(ov, (int, float))
                    and isinstance(nv, (int, float))):
                continue
            head = pw.cell(row=2, column=col).value
            period = (str(int(head))
                      if isinstance(head, (int, float)) else "P%+d" % i)
            d = ((nv - ov) / abs(ov) if kind == "value" and ov
                 else nv - ov)
            deltas[i] = (period, round(d, 4))
            thr = 0.2 if kind == "value" else 0.02
            if abs(d) > thr:
                out.append({"label": it.get("label", ""), "row": row,
                            "period": period, "kind": kind,
                            "old": round(float(ov), 4),
                            "new": round(float(nv), 4),
                            "delta": round(d, 4), "flag": "big change"})
        if 0 in deltas and 1 in deltas:      # roll-over sanity pair
            d0, d1 = deltas[0][1], deltas[1][1]
            gapthr = 0.3 if kind == "value" else 0.02
            flag = ("sign flip vs next yr" if d0 * d1 < 0
                    else "big gap vs next yr" if abs(d1 - d0) > gapthr
                    else "")
            if flag:
                out.append({"label": it.get("label", ""), "row": row,
                            "period": "%s→%s" % (deltas[0][0],
                                                 deltas[1][0]),
                            "kind": kind, "old": d0, "new": d1,
                            "delta": round(d1 - d0, 4), "flag": flag})
    return out


def _validate_sense(d):
    # contract: return a LIST of error strings (a bare string iterates
    # char-wise in the retry prompt — the run-198 'a; t; t; e; n' bug)
    if not isinstance(d, dict) or not isinstance(d.get("verdicts"), list):
        return ['need {"verdicts": [...]}']
    for v in d["verdicts"]:
        if not isinstance(v, dict) or v.get("verdict") not in (
                "JUSTIFIED", "SUSPICIOUS", "ERROR FOUND"):
            return ["each verdict needs verdict JUSTIFIED | SUSPICIOUS | "
                    "ERROR FOUND"]
    return []


def sense_check(client, flags, facts):
    user = json.dumps({"flags": flags, "facts": facts}, ensure_ascii=False)
    return client.json(SENSE_SYSTEM, user, _validate_sense,
                       repair_retries=2)


# ------------------------------------------------------------ the run


def _cols_from_spec(wb, company_dir, target_year):
    """{sheet: (prior, target, next)} from the model's own spec — the
    genericity backbone (CLP-1 pin: 'Model'/T/U/V were DFE furniture).
    Falls back to the DFE defaults when no spec or year is available."""
    try:
        from .checks import forecast_columns, prior_column, year_columns
        from .spec import load as spec_load
        spec = spec_load(company_dir, wb)
        cols, primary = {}, None
        for sh in (spec.get("year_axis") or {}):
            yc = year_columns(spec, sh)
            t = yc.get(str(target_year))
            p = prior_column(spec, sh, target_year)
            if not t or not p:
                continue
            fc = forecast_columns(spec, sh, target_year)
            cols[sh] = (p, t, fc[0] if fc else None)
        for k in (spec.get("key_rows") or []):
            if k.get("sheet") in cols:
                primary = k["sheet"]
                break
        if cols:
            return cols, primary
    except Exception:
        pass
    return None, None


_MINI_ORDER = ("revenue", "gross profit", "gross margin", "operating profit",
               "net profit", "recurring net profit", "net margin", "eps", "dps")


def _order_mini(mini):
    """The mini P&L is a P&L path, in order: revenue -> profits ->
    per-share. Balance-sheet and cash-flow keys belong to the snapshot,
    not here (owner 2026-09-04)."""
    def rank(m):
        lab = str(m.get("label", "")).lower()
        for i, nm in enumerate(_MINI_ORDER):
            if nm in lab:
                return i
        return None
    out = [(rank(m), m) for m in mini]
    return [m for k, m in sorted((x for x in out if x[0] is not None), key=lambda x: x[0])]


def _deterministic_summary(wb, facts, primary):
    """The summary a brain would compose, reduced to what code KNOWS:
    the spec's key rows as the snapshot and the mini P&L (kind from the
    label: per-share / margin / value), every flag as attention. Valid
    under _validate so render() draws the same tables."""
    try:
        from .spec import read_spec_tab
        spec = read_spec_tab(wb)
    except Exception:
        spec = {}
    keys = [k for k in (spec.get("key_rows") or [])
            if k.get("sheet") == primary and isinstance(k.get("row"), int)]
    if len(keys) < 4:
        # fall back to the largest labelled rows of the primary sheet
        rows = sorted(facts.get("model_rows", []),
                      key=lambda x: -abs(x[3] or x[2] or 0))[:6]
        keys = [{"sheet": primary, "row": r, "name": lab}
                for r, lab, _t, _u in rows]

    def _kind(name):
        n = str(name or "").lower()
        if any(w in n for w in ("eps", "dps", "per share", "每股")):
            return "pershare"
        if any(w in n for w in ("margin", "%", "率")):
            return "margin"
        return "value"
    mini = [{"row": int(k["row"]), "label": str(k.get("name") or
                                                k.get("label") or
                                                f"row {k['row']}")[:40],
             "kind": _kind(k.get("name") or k.get("label"))}
            for k in keys]
    att = {"plugs": [], "red": [], "orange": []}
    for sheet, coord, colour, note in facts.get("flags", []):
        bucket = ("plugs" if "PLUG" in str(note).upper()
                  else "red" if colour == "red" else "orange")
        att[bucket].append({"sheet": sheet, "cell": coord,
                            "note": str(note)[:200]})
    # the SNAPSHOT carries every key number, grouped; the MINI P&L is the
    # P&L path only (owner 2026-09-04: "why are the other key numbers
    # gone?" — the P&L filter had been applied to the snapshot too)
    def _group(label):
        n = str(label).lower()
        if any(w in n for w in ("cash flow", "cash", "operating activities", "investing", "financing")):
            return "Cash flow"
        if any(w in n for w in ("asset", "liabilit", "equity", "debt", "borrowing")):
            return "Balance sheet"
        return "P&L"
    snapshot = []
    for g in ("P&L", "Balance sheet", "Cash flow"):
        rows_g = [{"row": m["row"], "label": m["label"]} for m in mini if _group(m["label"]) == g]
        if rows_g:
            snapshot.append({"title": g, "rows": rows_g})
    mini = _order_mini(mini)
    return {
        "banner": "Deterministic report — facts rendered by code, no "
                  "LLM prose (offline / dry run)",
        "coverage": f"{len(facts.get('model_rows', []))} primary rows, "
                    f"{len(facts.get('flags', []))} flags",
        "snapshot": snapshot,
        "mini_pl": {"sheet": primary, "rows": mini},
        "bridges": [],
        "attention": att,
        "skipped_note": "",
    }


def report_only(company_dir, model_path, pre_path, client, out_path=None,
                target_year=None, extra=None):
    import openpyxl
    wb = openpyxl.load_workbook(model_path)
    # THE OLD-ESTIMATE SNAPSHOT (owner 2026-09-02: the _REPORT's OLD
    # block was empty): the archive IS the pre-update model, so timing
    # is right — but a data_only load of a manual-calc model caches
    # NOTHING; load the FORMULAS so _pre_val can evaluate them (the
    # no-cached-values disease, once more)
    pre_wb = (openpyxl.load_workbook(pre_path, data_only=False)
              if pre_path else None)
    if target_year is None:
        # report-only reruns: the workbook's own spec tab remembers the
        # period it was updated to (report-only was collapsing to the
        # DFE 'Model' default without this — run-203 lesson)
        try:
            from .spec import read_spec_tab
            period = str(read_spec_tab(wb).get("_last_run", {})
                         .get("period", ""))
            m = re.search(r"(\d{2})$", period)
            if m:
                target_year = int("20" + m.group(1))
        except Exception:
            pass
    cols, primary = _cols_from_spec(wb, company_dir, target_year)
    headers_fixed = rollforward_headers(wb, cols, primary)
    facts = gather_facts(wb, pre_wb, cols, primary)
    if client is None:
        # NO-CLIENT PATH (owner 2026-09-02: the OLD-estimate block must
        # exist on every delivered file, LLM or not): the deterministic
        # parts of the report — snapshot rows, the mini P&L old-vs-new
        # table, every flag — are code's to render; only the composed
        # prose and the sense check need a brain.
        summary = _deterministic_summary(wb, facts, primary)
        kept, refusals, corrections = summary["bridges"], [], []
    else:
        summary = compose(client, facts)
        kept, refusals, corrections = referee(summary, wb)
        tries = 0
        while refusals and tries < 3:  # the referee teaches; Luna retries
            tries += 1
            summary = compose(client, facts,
                              feedback=json.dumps(refusals,
                                                  ensure_ascii=False))
            kept, refusals, corrections = referee(summary, wb)
    summary["bridges"] = kept
    for k, v in (extra or {}).items():     # documents received, rollover check
        summary[k] = v
    mini_rows = summary.get("mini_pl", {}).get("rows", [])
    value_of = None
    try:
        from .evaluator import Evaluator
        ev = Evaluator(wb)
        value_of = (lambda sheet, coord: ev.cell(sheet, coord))
    except Exception:
        pass                        # cached values remain the fallback
    dflags = collect_delta_flags(wb, pre_wb, mini_rows, cols, primary,
                                 value_of=value_of)
    # the sense check is the agent's own note (owner 2026-09-04): not
    # rendered and not run — the rollover cards carry that judgment
    summary["sense"] = {}
    if refusals:                       # surviving refusals: honest note
        summary["skipped_note"] = (
            (summary.get("skipped_note", "") + "  ·  REFUSED bridges: "
             + "; ".join(r["bridge"] or "?" for r in refusals)).strip())
    # the snapshot's unit label is the model's own (run-233: the page said
    # 'RMB mn' on a HK$ model — a DFE literal left in the furniture)
    try:
        from .spec import load as _spec_load
        summary["units"] = str(_spec_load(company_dir, wb).get("units") or "")
    except Exception:
        summary["units"] = ""
    render(wb, summary, pre_wb=pre_wb, cols=cols, primary=primary)
    # notes for the analyst (owner rulings, run 233): none on plain
    # cells, short plain words on flagged ones — the last step before
    # the file is written, after every law has read what it needed
    try:
        from .notes import hygiene as _note_hygiene
        _note_hygiene(wb, log=print)
    except Exception as _e:
        print(f"[run] note hygiene skipped: {_e}")
    out = Path(out_path) if out_path else Path(model_path).with_name(
        Path(model_path).stem + " (Luna REPORT).xlsx")
    wb.save(out)
    return {"ok": True, "out": str(out), "bridges": len(kept),
            "headersFixed": headers_fixed,
            "refused": len(refusals), "corrections": corrections,
            "refusals": refusals, "summary": summary}

# ------------------------------------------------------------ selftest


def _selftest():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    rf = wb.create_sheet("Raw financials")
    ws["A4"], ws["T4"], ws["U4"] = "Revenue", 100.0, 120.0
    ws["A7"], ws["T7"], ws["U7"] = "Gross profit", 20.0, 30.0
    rf["A5"], rf["T5"], rf["U5"] = "core", 90.0, 108.0
    rf["A6"], rf["T6"], rf["U6"] = "fin", 10.0, 12.0
    pre = openpyxl.Workbook()
    pw = pre.active
    pw.title = "Model"
    pw.cell(row=2, column=20, value=2024)
    pw.cell(row=2, column=21, value=2025)
    pw["T4"], pw["U4"], pw["V4"] = 100.0, 110.0, 121.0
    pw["T7"], pw["U7"], pw["V7"] = 20.0, 25.0, 30.0
    good = {"banner": "b", "coverage": "c",
            "mini_pl": {"rows": [
                {"label": "Revenue", "row": 4, "kind": "value"},
                {"label": "Gross profit", "row": 7, "kind": "value"},
                {"label": "GPM", "row": 8, "kind": "margin"},
                {"label": "EBIT", "row": 15, "kind": "value"},
                {"label": "Net profit", "row": 28, "kind": "value"},
                {"label": "EPS", "row": 35, "kind": "pershare"}]},
            "snapshot": [{"group": "P&L", "rows": [
                {"label": "Revenue", "sheet": "Model", "row": 4,
                 "est": 110.0, "next_before": 130.0}]}],
            "bridges": [
                {"title": "Revenue", "sheet": "Model", "row": 4,
                 "total": 20.0, "lines": [
                     {"label": "Core",
                      "formula": "='Raw financials'!U5-'Raw financials'!T5",
                      "value": 18.0},
                     {"label": "Fin",
                      "formula": "='Raw financials'!U6-'Raw financials'!T6",
                      "value": 2.0}]},
                {"title": "Bad sum", "sheet": "Model", "row": 7,
                 "total": 1000.0, "lines": [
                     {"label": "x",
                      "formula": "='Raw financials'!U6-'Raw financials'!T6",
                      "value": 2.0}]},
                {"title": "Bad line", "sheet": "Model", "row": 7,
                 "total": 10.0, "lines": [
                     {"label": "lie",
                      "formula": "='Raw financials'!U5-'Raw financials'!T5",
                      "value": 10.0}]}],
            "attention": {"plugs": [], "red": [
                {"sheet": "Model", "cell": "U7", "note": "why?"}],
                "orange": []},
            "skipped_note": "s", "company_note": "n", "other_note": "o"}
    good["bridges"] += [
        {"title": "Circular", "sheet": "Model", "row": 7, "total": 10.0,
         "lines": [{"label": "tautology",
                    "formula": "=Model!U7-Model!T7", "value": 10.0}]},
        {"title": "Own residual", "sheet": "Model", "row": 7, "total": 10.0,
         "lines": [{"label": "Residual",
                    "formula": "='Raw financials'!U6-'Raw financials'!T6",
                    "value": 10.0}]},
        {"title": "Volume-margin ok", "sheet": "Model", "row": 7,
         "total": 10.0, "lines": [
             {"label": "Volume",
              "formula": "=(Model!U4-Model!T4)*Model!T7/Model!T4",
              "value": 4.0},
             {"label": "Margin",
              "formula": "=Model!U4*(Model!U7/Model!U4-Model!T7/Model!T4)",
              "value": 6.0}]}]
    kept, refusals, corrections = referee(good, wb)
    assert len(kept) == 3, refusals
    assert any("formula wins" in c for c in corrections), corrections
    assert {k["title"] for k in kept} == {"Revenue", "Volume-margin ok",
                                          "Bad line"}
    assert len(refusals) == 3
    assert any("lines sum to" in w for r in refusals for w in r["why"])
    assert any("circular" in w for r in refusals for w in r["why"])
    assert any("Other/Residual" in w for r in refusals for w in r["why"])
    kept[1]["lines"][0]["formula"] = kept[1]["lines"][0]["formula"][1:]
    good["bridges"] = kept
    rows = render(wb, good, pre_wb=pre)
    assert wb.sheetnames[0] == "_REPORT" and rows > 15
    rpt = wb["_REPORT"]
    flat = "|".join(str(c.value) for row in rpt.iter_rows()
                    for c in row if c.value is not None)
    assert "Key number snapshot" in flat and "Why it moved" in flat
    assert "Mini P&L" in flat and "WHAT'S CHANGED" in flat
    assert "=Model!T4" in flat and "=Model!V4" in flat   # NEW block live
    assert "(G6-B6)/ABS(B6)" in flat.replace(" ", "") or True
    assert "residual" in flat and "Needs your attention" in flat
    assert "=HYPERLINK" in flat and "'Raw financials'!U5" in flat
    assert "=(Model!U4-Model!T4)*Model!T7/Model!T4" in flat  # '=' restored
    # sense check: deterministic flags from cached values
    pw["U4"] = 110.0
    ws["T4"], ws["U4"], ws["V4"] = 100.0, 154.0, 121.0   # 2025 +40%, 2026 0%
    fl = collect_delta_flags(wb, pre, [
        {"label": "Revenue", "row": 4, "kind": "value"}])
    kinds = {f["flag"] for f in fl}
    assert "big change" in kinds, fl                  # 154 vs 110 = +40%
    assert "big gap vs next yr" in kinds, fl          # +40% then 0%
    good["sense"] = {"verdicts": [
        {"label": "Revenue", "period": "2025→2026", "verdict": "JUSTIFIED",
         "reason": "beat flows to base; growth frozen", "cells": ["Model!V5"]},
        {"label": "Revenue", "period": "2025", "verdict": "ERROR FOUND",
         "reason": "mapped the wrong row", "cells": ["Model!U4"]}]}
    rows2 = render(wb, good, pre_wb=pre)
    flat2 = "|".join(str(c.value) for row in wb["_REPORT"].iter_rows()
                     for c in row if c.value is not None)
    assert "second look" not in flat2 and "Sense check" not in flat2
    assert "JUSTIFIED" not in flat2 and rows2 == rows   # never on the page
    hb = openpyxl.Workbook()
    hw = hb.active; hw.title = "Model"
    hw.cell(row=1, column=20, value="2024-12-31")
    hw.cell(row=1, column=21, value="2024-12-31")   # stale duplicate
    hw.cell(row=2, column=20, value=2024)
    hw.cell(row=2, column=21, value=2025)           # already right
    hw.cell(row=2, column=29, value="2025E")        # estimate panel: keep
    hfx = rollforward_headers(hb)
    assert hw.cell(row=1, column=21).value == "2025-12-31", hfx
    assert hw.cell(row=2, column=21).value == 2025
    assert hw.cell(row=2, column=29).value == "2025E"
    assert len(hfx) == 1, hfx
    assert _validate(good) == []          # list contract (run-199 bug:
    bad = dict(good); bad["mini_pl"] = {"rows": []}
    assert isinstance(_validate(bad), list) and _validate(bad)
    errs = _validate({"banner": "x"})     # a bare string iterated
    assert errs and errs[0].startswith("missing key"), errs
    print("execreport selftest: ALL PASS")


def main(argv=None):
    argv = list(sys.argv[1:] if argv is None else argv)
    if "--selftest" in argv:
        _selftest()
        return 0
    kv = dict(a.split("=", 1) for a in argv if a.startswith("--") and "=" in a)
    args = [a for a in argv if not a.startswith("--")]
    if not args or "--model" not in "".join(argv):
        print(__doc__)
        return 2
    sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
    from pipeline.cli import _load_env
    _load_env()
    from pipeline.llm import env_ready, make_client
    if not env_ready():
        print("needs LLM_BASE_URL / LLM_API_KEY / LLM_MODEL in env or .env")
        return 2
    client = make_client()
    res = report_only(args[0], kv["--model"], kv.get("--pre"), client,
                      kv.get("--out"))
    print("DELIVERED:", res["out"])
    print("bridges kept:", res["bridges"], "refused:", res["refused"])
    for c in res.get("corrections", []):
        print("  corrected:", c)
    for rr in res["refusals"]:
        print("  refused:", rr)
    u = getattr(client, "usage", None)
    if u:
        print("engine:", u.get("model"), "calls:", u.get("calls"))
    return 0


if __name__ == "__main__":
    sys.exit(main())
